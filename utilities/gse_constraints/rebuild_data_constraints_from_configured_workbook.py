"""Build a review candidate CSV from the configured GSE source workbook.

The active CSV is not modified. By default this script validates the selected
workbook and reports the number of source rows. Use --write-candidate to create
data/data-constraints.from-gse.csv for review.
"""

from __future__ import annotations

import argparse
import csv
import os
from pathlib import Path
from tempfile import NamedTemporaryFile
from typing import Iterable

from openpyxl import load_workbook
from rdflib import Graph, Namespace, URIRef
from rdflib.namespace import RDFS


PROJECT_ROOT = Path(__file__).resolve().parents[2]
CONFIGURATION_FILE = PROJECT_ROOT / "config" / "configuration.ttl"
ACTIVE_CSV = PROJECT_ROOT / "data" / "data-constraints.csv"
CANDIDATE_CSV = PROJECT_ROOT / "data" / "data-constraints.from-gse.csv"

CFG = Namespace("https://dynamicontology.com/uad36/configuration#")
DEVELOPER_CONFIGURATION = URIRef(
    "https://dynamicontology.com/uad36/configuration/developer"
)

EXPECTED_HEADERS = (
    "Unique ID",
    "Primary Data Element",
    "Message ID",
    "Message Text",
    "Rule Logic",
    "Severity",
    "Property Affected",
    "Report Section",
    "Report Subsection",
    "Report Label / Value",
    "Data Point Name / Value",
    " xPath",
    "Related Value(s)",
    "Date Format",
    "Number Format",
    "Min Value",
    "Max Value",
)


def one_object(graph: Graph, subject: URIRef, predicate: URIRef, label: str):
    values = list(graph.objects(subject, predicate))
    if len(values) != 1:
        raise ValueError(
            f"Configuration must provide exactly one {label}; "
            f"found {len(values)}."
        )
    return values[0]


def configured_workbook() -> tuple[Path, str, str]:
    """Resolve the active source workbook and rule-data sheet from TTL config."""

    if not CONFIGURATION_FILE.is_file():
        raise FileNotFoundError(
            f"Configuration file not found: {CONFIGURATION_FILE}"
        )

    graph = Graph().parse(str(CONFIGURATION_FILE), format="turtle")

    active_source = one_object(
        graph,
        DEVELOPER_CONFIGURATION,
        CFG.activeComplianceSource,
        "active compliance source",
    )
    source_collection = one_object(
        graph,
        DEVELOPER_CONFIGURATION,
        CFG.hasSourceWorkbooks,
        "source workbook collection",
    )

    configured_sources = set(
        graph.objects(source_collection, CFG.containsWorkbook)
    )
    if active_source not in configured_sources:
        raise ValueError(
            "The active compliance source is not a member of the configured "
            "source workbook collection."
        )

    source_location = str(
        one_object(
            graph,
            active_source,
            CFG.sourceFileLocation,
            "source workbook file location",
        )
    )
    sheet_name = str(
        one_object(
            graph,
            active_source,
            CFG.ruleDataSheetName,
            "rule-data sheet name",
        )
    )
    source_label = str(graph.value(active_source, RDFS.label) or active_source)

    workbook_path = Path(source_location)
    if not workbook_path.is_file():
        raise FileNotFoundError(
            f"Configured source workbook not found: {workbook_path}"
        )

    return workbook_path, sheet_name, source_label


def csv_text(value: object) -> str:
    """Convert an Excel cell value to text without trimming source content."""

    if value is None:
        return ""
    if isinstance(value, bool):
        return "TRUE" if value else "FALSE"
    return str(value)


def read_source_rows(
    workbook_path: Path,
    sheet_name: str,
) -> tuple[tuple[str, ...], list[tuple[str, ...]]]:
    """Read and validate the configured rule sheet."""

    workbook = load_workbook(
        filename=workbook_path,
        read_only=True,
        data_only=False,
    )

    try:
        if sheet_name not in workbook.sheetnames:
            raise ValueError(
                f"Configured rule-data sheet {sheet_name!r} is not present "
                f"in {workbook_path.name}."
            )

        worksheet = workbook[sheet_name]
        if worksheet.max_column != len(EXPECTED_HEADERS):
            raise ValueError(
                f"Expected {len(EXPECTED_HEADERS)} columns in {sheet_name!r}; "
                f"found {worksheet.max_column}."
            )

        header_cells = next(
            worksheet.iter_rows(
                min_row=1,
                max_row=1,
                min_col=1,
                max_col=len(EXPECTED_HEADERS),
            )
        )
        headers = tuple(csv_text(cell.value) for cell in header_cells)
        if headers != EXPECTED_HEADERS:
            raise ValueError(
                "The source workbook headers do not match the governed "
                f"structure.\nExpected: {EXPECTED_HEADERS!r}\n"
                f"Found:    {headers!r}"
            )

        rows: list[tuple[str, ...]] = []
        seen_message_ids: set[str] = set()

        for row_number, cells in enumerate(
            worksheet.iter_rows(
                min_row=2,
                min_col=1,
                max_col=len(EXPECTED_HEADERS),
            ),
            start=2,
        ):
            if any(cell.data_type == "f" for cell in cells):
                raise ValueError(
                    f"Formula found in governed data row {row_number}; "
                    "source rule rows must be read as published values."
                )

            values = tuple(csv_text(cell.value) for cell in cells)
            if not any(value.strip() for value in values):
                continue

            message_id = values[2].strip()
            if not message_id:
                raise ValueError(
                    f"Governed data row {row_number} has no Message ID."
                )
            if message_id in seen_message_ids:
                raise ValueError(
                    f"Duplicate Message ID {message_id!r} in source row "
                    f"{row_number}."
                )

            seen_message_ids.add(message_id)
            rows.append(values)

        if not rows:
            raise ValueError("The configured rule-data sheet contains no rows.")

        return headers, rows
    finally:
        workbook.close()


def verify_csv(path: Path, headers: tuple[str, ...], rows: list[tuple[str, ...]]):
    """Verify that the written CSV round-trips every value exactly."""

    with path.open(encoding="utf-8-sig", newline="") as stream:
        actual = list(csv.reader(stream))

    expected = [list(headers), *(list(row) for row in rows)]
    if actual != expected:
        raise ValueError(
            "Generated CSV did not round-trip exactly to the workbook values."
        )


def write_candidate(
    headers: tuple[str, ...],
    rows: list[tuple[str, ...]],
    *,
    replace_existing: bool,
) -> None:
    """Write a separate candidate CSV without changing the active CSV."""

    if CANDIDATE_CSV.exists() and not replace_existing:
        raise FileExistsError(
            f"Candidate already exists: {CANDIDATE_CSV}\n"
            "Review or remove it first, or use --replace-candidate."
        )

    temporary_path: Path | None = None
    try:
        with NamedTemporaryFile(
            mode="w",
            encoding="utf-8-sig",
            newline="",
            prefix=".data-constraints-from-gse-",
            suffix=".tmp",
            dir=CANDIDATE_CSV.parent,
            delete=False,
        ) as stream:
            temporary_path = Path(stream.name)
            writer = csv.writer(stream, lineterminator="\r\n")
            writer.writerow(headers)
            writer.writerows(rows)

        verify_csv(temporary_path, headers, rows)
        os.replace(temporary_path, CANDIDATE_CSV)
        temporary_path = None
    finally:
        if temporary_path is not None and temporary_path.exists():
            temporary_path.unlink()


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(
        description=(
            "Validate or create a review candidate CSV from the active GSE "
            "workbook configured in config/configuration.ttl."
        )
    )
    mode = parser.add_mutually_exclusive_group()
    mode.add_argument(
        "--write-candidate",
        action="store_true",
        help="Create data/data-constraints.from-gse.csv; do not replace the active CSV.",
    )
    mode.add_argument(
        "--replace-candidate",
        action="store_true",
        help="Replace an existing generated candidate CSV.",
    )
    return parser.parse_args()


def main() -> int:
    args = parse_args()
    workbook_path, sheet_name, source_label = configured_workbook()
    headers, rows = read_source_rows(workbook_path, sheet_name)

    print(f"Validated {len(rows)} source rows.")
    print(f"Source: {source_label}")
    print(f"Sheet: {sheet_name}")
    print(f"Workbook: {workbook_path}")
    print(f"Active application CSV remains: {ACTIVE_CSV}")

    if args.write_candidate or args.replace_candidate:
        write_candidate(
            headers,
            rows,
            replace_existing=args.replace_candidate,
        )
        print(f"Wrote review candidate: {CANDIDATE_CSV}")
    else:
        print("No files written. Use --write-candidate to create a review copy.")

    return 0


if __name__ == "__main__":
    raise SystemExit(main())