"""Govern authoritative source evidence for GSE constraints."""

from __future__ import annotations

from dataclasses import dataclass
from datetime import date, datetime, time
from decimal import Decimal, InvalidOperation
from pathlib import Path
from typing import Mapping, Protocol, TypeVar

from rdflib import Graph, Literal, Namespace, RDF, URIRef


class AmbiguousGseAuthorityError(ValueError):
    """Raised when governed evidence cannot establish GSE authority."""


class AuthoritativeGseConstraintSourceEvidence(Protocol):
    source_document_name: str
    source_version: str
    issuing_authorities: frozenset[str]
    acquisition_location: str
    content_digest: str
    constraint_ids: tuple[str, ...]


class VersionedGseConstraintReleaseEvidence(
    AuthoritativeGseConstraintSourceEvidence,
    Protocol,
):
    constraint_definitions: tuple[tuple[str, str], ...]


class GseConstraintSourceRepository(Protocol):
    def save(self, registered_source: object) -> None: ...


class GseAuthorityGovernanceReview(Protocol):
    def record(self, **event: object) -> None: ...


class GseConstraintSourceCandidate(Protocol):
    acquisition_location: str


SourceCandidate = TypeVar(
    "SourceCandidate",
    bound=GseConstraintSourceCandidate,
)


class ClassifiedGseConstraintGovernanceRecord(Protocol):
    canonical_constraint_id: str
    gse_classification: str
    issuing_authorities: frozenset[str]
    source_document_name: str
    source_version: str
    content_digest: str
    source_sheet: str
    source_row: int


GovernanceRecord = TypeVar(
    "GovernanceRecord",
    bound=ClassifiedGseConstraintGovernanceRecord,
)


class GseConstraintGovernanceRepository(Protocol[GovernanceRecord]):
    def get_governance_record(
        self,
        canonical_constraint_id: str,
        source_version: str,
    ) -> GovernanceRecord | None: ...


@dataclass(frozen=True)
class RegisteredGseConstraintSource:
    """One governed source classified from its explicit issuer evidence."""

    source_document_name: str
    source_version: str
    issuing_authorities: frozenset[str]
    acquisition_location: str
    content_digest: str
    gse_classification: str
    constraint_classifications: tuple[tuple[str, str], ...]


@dataclass(frozen=True)
class RegisteredCanonicalGseConstraintSet:
    """One canonical release with all authoritative acquisition provenance."""

    source_document_name: str
    source_version: str
    issuing_authorities: frozenset[str]
    content_digest: str
    gse_classification: str
    canonical_constraint_ids: tuple[str, ...]
    acquisition_locations: tuple[str, ...]
    constraint_definitions: tuple[tuple[str, str], ...] = ()


def discover_authoritative_gse_constraint_sources(
    candidate_sources: tuple[SourceCandidate, ...],
) -> tuple[SourceCandidate, ...]:
    """Return only candidates acquired from governed GSE source roots."""

    authoritative_roots = (
        "/specs/fanniemae/",
        "/specs/freddiemac/",
    )
    return tuple(
        candidate
        for candidate in candidate_sources
        if any(
            root
            in (
                "/"
                + candidate.acquisition_location
                .replace("\\", "/")
                .strip("/")
                .casefold()
                + "/"
            )
            for root in authoritative_roots
        )
    )


def register_canonical_gse_constraint_set(
    source_evidence: tuple[AuthoritativeGseConstraintSourceEvidence, ...],
    repository: GseConstraintSourceRepository,
) -> RegisteredCanonicalGseConstraintSet:
    """Register identical authoritative copies as one shared GSE release."""

    if not source_evidence:
        raise ValueError("At least one authoritative source copy is required.")
    authoritative_evidence = (
        discover_authoritative_gse_constraint_sources(source_evidence)
    )
    if len(authoritative_evidence) != len(source_evidence):
        raise ValueError(
            "Every source copy must come from an authoritative GSE root."
        )

    signatures = {
        (
            evidence.source_document_name,
            evidence.source_version,
            evidence.issuing_authorities,
            evidence.content_digest,
            evidence.constraint_ids,
        )
        for evidence in authoritative_evidence
    }
    if len(signatures) != 1:
        raise ValueError(
            "Canonical copies must describe the same governed release."
        )

    first = authoritative_evidence[0]
    joint_issuers = frozenset({"fannie_mae", "freddie_mac"})
    if first.issuing_authorities != joint_issuers:
        raise ValueError(
            "A canonical joint release requires both GSEs as issuers."
        )

    registered_set = RegisteredCanonicalGseConstraintSet(
        source_document_name=first.source_document_name,
        source_version=first.source_version,
        issuing_authorities=first.issuing_authorities,
        content_digest=first.content_digest,
        gse_classification="shared",
        canonical_constraint_ids=tuple(dict.fromkeys(first.constraint_ids)),
        acquisition_locations=tuple(
            dict.fromkeys(
                evidence.acquisition_location
                for evidence in authoritative_evidence
            )
        ),
        constraint_definitions=getattr(
            first,
            "constraint_definitions",
            (),
        ),
    )
    repository.save(registered_set)
    return registered_set


def register_versioned_gse_constraint_sets(
    release_evidence: tuple[VersionedGseConstraintReleaseEvidence, ...],
    repository: GseConstraintSourceRepository,
) -> tuple[RegisteredCanonicalGseConstraintSet, ...]:
    """Register each governed release without rewriting another version."""

    registered_sets: list[RegisteredCanonicalGseConstraintSet] = []
    release_identities: set[tuple[str, str]] = set()
    for evidence in release_evidence:
        release_identity = (
            evidence.source_version,
            evidence.content_digest,
        )
        if release_identity in release_identities:
            raise ValueError(
                "Each versioned release identity may be registered once."
            )
        release_identities.add(release_identity)
        registered_sets.append(
            register_canonical_gse_constraint_set(
                (evidence,),
                repository,
            )
        )
    return tuple(registered_sets)


def get_classified_gse_constraint_governance_record(
    canonical_constraint_id: str,
    source_version: str,
    repository: GseConstraintGovernanceRepository[GovernanceRecord],
) -> GovernanceRecord:
    """Return one classified constraint's exact authoritative evidence."""

    governance_record = repository.get_governance_record(
        canonical_constraint_id,
        source_version,
    )
    if governance_record is None:
        raise LookupError(
            "No governance record exists for the requested constraint "
            "and source version."
        )
    if governance_record.canonical_constraint_id != canonical_constraint_id:
        raise ValueError("The repository returned another constraint.")
    if governance_record.source_version != source_version:
        raise ValueError("The repository returned another source version.")
    if (
        not governance_record.gse_classification
        or not governance_record.issuing_authorities
        or not governance_record.source_document_name
        or not governance_record.content_digest
        or not governance_record.source_sheet
        or governance_record.source_row < 1
    ):
        raise ValueError(
            "The classified constraint governance record is incomplete."
        )
    return governance_record


def register_authoritative_gse_constraint_source(
    source_evidence: AuthoritativeGseConstraintSourceEvidence,
    repository: GseConstraintSourceRepository,
    governance_review: GseAuthorityGovernanceReview | None = None,
) -> RegisteredGseConstraintSource:
    """Classify a source from explicit issuer evidence, never its folder."""

    classifications_by_issuers = {
        frozenset({"fannie_mae", "freddie_mac"}): "shared",
        frozenset({"fannie_mae"}): "fannie_mae_only",
        frozenset({"freddie_mac"}): "freddie_mac_only",
    }
    try:
        classification = classifications_by_issuers[
            source_evidence.issuing_authorities
        ]
    except KeyError as error:
        if governance_review is not None:
            governance_review.record(
                event_type="ambiguous-gse-authority",
                source_document_name=source_evidence.source_document_name,
                source_version=source_evidence.source_version,
                acquisition_location=source_evidence.acquisition_location,
                content_digest=source_evidence.content_digest,
            )
        raise AmbiguousGseAuthorityError(
            "Authoritative GSE source registration requires explicit "
            "evidence of Fannie Mae, Freddie Mac, or both as issuers."
        ) from error
    if not source_evidence.source_document_name.strip():
        raise ValueError("An authoritative source document name is required.")
    if not source_evidence.source_version.strip():
        raise ValueError("An authoritative source version is required.")
    if not source_evidence.content_digest.strip():
        raise ValueError("An authoritative source content digest is required.")

    registered_source = RegisteredGseConstraintSource(
        source_document_name=source_evidence.source_document_name,
        source_version=source_evidence.source_version,
        issuing_authorities=source_evidence.issuing_authorities,
        acquisition_location=source_evidence.acquisition_location,
        content_digest=source_evidence.content_digest,
        gse_classification=classification,
        constraint_classifications=tuple(
            (constraint_id, classification)
            for constraint_id in source_evidence.constraint_ids
        ),
    )
    repository.save(registered_source)
    return registered_source

# IT-60 governed workbook reading


CFG = Namespace("https://dynamicontology.com/uad36/configuration#")
APP_CONFIGURATION_TYPE = CFG.ApplicationConfiguration
REQUIRED_CONSTRAINT_HEADERS = frozenset(
    {
        "Unique ID",
        "Primary Data Element",
        "Message ID",
        "Rule Logic",
        " xPath",
    }
)


class GovernedWorkbookConfigurationError(ValueError):
    """Raised when configured governed workbook selection is invalid."""


@dataclass(frozen=True)
class ActiveComplianceWorkbookRows:
    """Rows read from the configured active GSE workbook worksheet."""

    workbook_path: Path
    worksheet_name: str
    rows: tuple[Mapping[str, str], ...]


def _one_object(
    graph: Graph,
    subject: URIRef,
    predicate: URIRef,
    *,
    setting_name: str,
) -> URIRef | Literal:
    values = tuple(graph.objects(subject, predicate))
    if len(values) != 1:
        raise ValueError(
            f"Configuration must provide exactly one {setting_name}; "
            f"found {len(values)}."
        )
    return values[0]


def _read_active_workbook_settings(
    configuration_path: str | Path,
) -> tuple[Path, str]:
    config_path = Path(configuration_path).expanduser().resolve()
    if not config_path.is_file():
        raise FileNotFoundError(
            f"UAD configuration file does not exist: {config_path}"
        )

    graph = Graph()
    graph.parse(config_path, format="turtle")
    applications = tuple(graph.subjects(RDF.type, APP_CONFIGURATION_TYPE))
    if len(applications) != 1:
        raise ValueError(
            "Configuration must identify exactly one application configuration; "
            f"found {len(applications)}."
        )
    application = applications[0]

    collection = _one_object(
        graph,
        application,
        CFG.hasSourceWorkbooks,
        setting_name="source workbook collection",
    )
    active_source = _one_object(
        graph,
        application,
        CFG.activeComplianceSource,
        setting_name="active compliance source",
    )
    if active_source not in set(graph.objects(collection, CFG.containsWorkbook)):
        raise ValueError(
            "The active compliance source must be a member of the configured "
            "source workbook collection."
        )

    workbook_location = _one_object(
        graph,
        active_source,
        CFG.sourceFileLocation,
        setting_name="active workbook file location",
    )
    worksheet_name = _one_object(
        graph,
        active_source,
        CFG.ruleDataSheetName,
        setting_name="active workbook rule data worksheet",
    )
    if not isinstance(workbook_location, Literal):
        raise ValueError("The active workbook location must be a literal path.")
    if not isinstance(worksheet_name, Literal):
        raise ValueError("The active workbook worksheet name must be a literal.")

    path = Path(str(workbook_location)).expanduser()
    if not path.is_absolute():
        path = (config_path.parent / path).resolve()
    else:
        path = path.resolve()
    sheet = str(worksheet_name).strip()
    if not sheet:
        raise ValueError("The active workbook worksheet name cannot be blank.")
    return path, sheet


def _cell_value_as_text(value: object, number_format: str) -> str:
    """Return stable text, respecting simple Excel zero-padding formats."""

    if value is None:
        return ""
    if isinstance(value, bool):
        return "TRUE" if value else "FALSE"
    if isinstance(value, (datetime, date, time)):
        return value.isoformat()
    if isinstance(value, (int, float, Decimal)):
        fmt = number_format.split(";", maxsplit=1)[0].strip()
        # Preserve leading zeroes for numeric identifiers when the source cell
        # uses an explicit Excel mask such as 0000.0000.
        if fmt and all(ch in "0#.?," for ch in fmt) and "." in fmt:
            integer_mask, fraction_mask = fmt.split(".", maxsplit=1)
            decimals = len(fraction_mask)
            try:
                number = Decimal(str(value))
                rendered = f"{number:.{decimals}f}"
                integer_part, fraction_part = rendered.split(".", maxsplit=1)
                required_integer_digits = sum(ch == "0" for ch in integer_mask)
                sign = "-" if integer_part.startswith("-") else ""
                digits = integer_part.lstrip("-")
                integer_part = sign + digits.zfill(required_integer_digits)
                if decimals:
                    return f"{integer_part}.{fraction_part}"
                return integer_part
            except (InvalidOperation, ValueError):
                pass
        if isinstance(value, float) and value.is_integer():
            return str(int(value))
        return str(value)
    return str(value)


def read_active_compliance_workbook_rows(
    configuration_path: str | Path,
) -> ActiveComplianceWorkbookRows:
    """Read all governed constraint rows from the configured active workbook.

    The RDF configuration selects one active workbook and its rule-data sheet.
    The worksheet header row is located by its governed column names so title
    or note rows above the table do not shift interpretation. Every cell value
    is returned as text; no CSV conversion or numeric ID coercion is performed.
    """

    workbook_path, worksheet_name = _read_active_workbook_settings(
        configuration_path
    )
    if workbook_path.suffix.casefold() not in {".xlsx", ".xlsm"}:
        raise ValueError(
            "The configured compliance source must be an .xlsx or .xlsm workbook."
        )
    if not workbook_path.is_file():
        raise FileNotFoundError(
            f"Configured compliance workbook does not exist: {workbook_path}"
        )

    from openpyxl import load_workbook

    workbook = load_workbook(
        workbook_path,
        read_only=True,
        data_only=True,
        keep_vba=workbook_path.suffix.casefold() == ".xlsm",
    )
    try:
        if worksheet_name not in workbook.sheetnames:
            raise GovernedWorkbookConfigurationError(
                f"Configured worksheet {worksheet_name!r} is not present in "
                f"{workbook_path.name}; available worksheets: "
                f"{', '.join(workbook.sheetnames)}"
            )
        worksheet = workbook[worksheet_name]
        header_row_number: int | None = None
        headers: tuple[str, ...] = ()
        for row_number, cells in enumerate(
            worksheet.iter_rows(),
            start=1,
        ):
            candidate = tuple(
                str(cell.value) if cell.value is not None else ""
                for cell in cells
            )
            if REQUIRED_CONSTRAINT_HEADERS.issubset(candidate):
                header_row_number = row_number
                headers = candidate
                break
        if header_row_number is None:
            raise ValueError(
                "Configured worksheet does not contain the required governed "
                "constraint headers: "
                + ", ".join(sorted(REQUIRED_CONSTRAINT_HEADERS))
            )

        rows: list[Mapping[str, str]] = []
        for cells in worksheet.iter_rows(min_row=header_row_number + 1):
            values = tuple(
                _cell_value_as_text(cell.value, cell.number_format)
                for cell in cells
            )
            if not any(value.strip() for value in values):
                continue
            record = {
                header: values[index] if index < len(values) else ""
                for index, header in enumerate(headers)
                if header
            }
            rows.append(record)
        return ActiveComplianceWorkbookRows(
            workbook_path=workbook_path,
            worksheet_name=worksheet_name,
            rows=tuple(rows),
        )
    finally:
        workbook.close()
