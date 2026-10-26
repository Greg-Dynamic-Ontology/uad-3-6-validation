"""Run workbook-scoped constraint construction without changing the registry."""

from __future__ import annotations

from dataclasses import dataclass
from datetime import datetime, timezone
from pathlib import Path
from uuid import uuid4

from openpyxl import load_workbook
from rdflib import Graph, Literal, RDF, RDFS, URIRef
from rdflib.namespace import PROV

from app.services import constraint_production_pipeline
from app.services.constraint_set_build_workbench import (
    CFG, CV, _index_preserved_stages, _location, _one, _text,
)
from app.services.governed_constraint_source_reader import _cell_text
from app.services.construction_exchange import UADEX, read_construction_stage_result


@dataclass(frozen=True)
class ConstraintPipelineRunResult:
    run_id: str
    workbook_id: str
    workbook_label: str
    worksheet: str
    unique_id: str
    status: str
    manifest_path: Path
    batch_results_path: Path
    run_directory: Path
    unique_ids: tuple[str, ...]
    completed_row_count: int
    exception_row_count: int = 0


def _workbook_input(configuration_path: str | Path, source_workbook_id: str):
    if not isinstance(source_workbook_id, str) or not source_workbook_id.strip():
        raise ValueError("Select a configured source workbook.")
    configuration_path = Path(configuration_path).resolve(strict=True)
    registry = Graph().parse(configuration_path, format="turtle")
    roots = set(registry.subjects(RDF.type, CFG.ApplicationConfiguration))
    if len(roots) != 1:
        raise ValueError("Construction requires one application configuration.")
    configuration = next(iter(roots))
    collection = _one(registry, configuration, CFG.hasSourceWorkbooks)
    workbook = URIRef(source_workbook_id)
    if (collection, RDF.type, CFG.GovernedSourceWorkbookCollection) not in registry or (
        collection, CFG.containsWorkbook, workbook
    ) not in registry or (workbook, RDF.type, CFG.GovernedSourceArtifact) not in registry:
        raise ValueError("Select a workbook from the configured source collection.")
    label = _text(registry, workbook, RDFS.label)
    worksheet = _text(registry, workbook, CFG.ruleDataSheetName)
    source_path = _location(registry, workbook, CFG.sourceFileLocation, configuration_path.parent, existing=True)
    formats = set(registry.objects(workbook, CFG.sourceFormat))
    source_format = (
        _text(registry, workbook, CFG.sourceFormat).lower().strip()
        if formats else source_path.suffix.lower().lstrip(".")
    )
    if source_format not in {"xlsx", "xlsm"}:
        raise ValueError("The selected source must be a supported Excel workbook.")
    return configuration_path, registry, configuration, workbook, label, worksheet, source_path, source_format


def list_workbook_row_identifiers(
    *, configuration_path: str | Path, source_workbook_id: str,
) -> tuple[str, ...]:
    """Offer exact textual IDs from the selected worksheet without constructing rows."""
    selected = _workbook_input(configuration_path, source_workbook_id)
    worksheet, source_path = selected[5], selected[6]
    workbook = load_workbook(source_path, read_only=True, data_only=True)
    try:
        if worksheet not in workbook.sheetnames:
            raise ValueError("The configured rule-data worksheet is missing.")
        iterator = workbook[worksheet].iter_rows(values_only=True)
        headers = tuple(_cell_text(value) for value in next(iterator, ()))
        # Selection reads identifiers only. The production reader validates the
        # complete governed row contract when construction is requested.
        if headers.count("Unique ID") != 1:
            raise ValueError("The rule-data worksheet must identify its Unique ID column.")
        column = headers.index("Unique ID")
        identifiers = []
        seen = set()
        for values in iterator:
            unique_id = _cell_text(values[column]) if column < len(values) else ""
            if not unique_id:
                continue
            if unique_id in seen:
                raise ValueError("A source-row identifier matches more than one workbook row.")
            identifiers.append(unique_id)
            seen.add(unique_id)
        return tuple(identifiers)
    finally:
        workbook.close()


def process_one_workbook_row(
    *, configuration_path: str | Path, source_workbook_id: str, unique_id: str,
) -> ConstraintPipelineRunResult:
    return _process_workbook_rows(
        configuration_path=configuration_path, source_workbook_id=source_workbook_id,
        unique_ids=(unique_id,), row_scope="one",
    )


def process_selected_workbook_rows(
    *, configuration_path: str | Path, source_workbook_id: str,
    unique_ids: list[str] | tuple[str, ...],
) -> ConstraintPipelineRunResult:
    if isinstance(unique_ids, str):
        raise ValueError("Select a group of governed row identifiers.")
    return _process_workbook_rows(
        configuration_path=configuration_path, source_workbook_id=source_workbook_id,
        unique_ids=tuple(unique_ids), row_scope="group",
    )


def process_all_workbook_rows(
    *, configuration_path: str | Path, source_workbook_id: str,
) -> ConstraintPipelineRunResult:
    unique_ids = list_workbook_row_identifiers(
        configuration_path=configuration_path, source_workbook_id=source_workbook_id,
    )
    return _process_workbook_rows(
        configuration_path=configuration_path, source_workbook_id=source_workbook_id,
        unique_ids=unique_ids, row_scope="all",
    )


def _process_workbook_rows(
    *, configuration_path: str | Path, source_workbook_id: str,
    unique_ids: tuple[str, ...], row_scope: str,
) -> ConstraintPipelineRunResult:
    if (not unique_ids and row_scope != "all") or any(not isinstance(value, str) or not value.strip() for value in unique_ids):
        raise ValueError("Select or enter the exact governed row identifiers.")
    if any(value != value.strip() for value in unique_ids):
        raise ValueError("Row identifiers must not contain surrounding whitespace.")
    # Deduplicate by exact string, preserving the operator's selection order.
    unique_ids = tuple(dict.fromkeys(unique_ids))
    (
        configuration_path, registry, configuration, workbook, label,
        worksheet, source_path, source_format,
    ) = _workbook_input(configuration_path, source_workbook_id)
    if row_scope == "group":
        available = set(list_workbook_row_identifiers(
            configuration_path=configuration_path, source_workbook_id=source_workbook_id,
        ))
        if any(unique_id not in available for unique_id in unique_ids):
            raise ValueError("A selected identifier does not belong to this workbook worksheet.")
    base = configuration_path.parent
    profile = _one(registry, configuration, CFG.usesConstraintBuildProfile)
    if (profile, RDF.type, CFG.ConstraintBuildProfile) not in registry:
        raise ValueError("A governed construction build profile is required.")
    models = {}
    model_paths = []
    for argument, predicate in (
        ("logical_schema_graph", CFG.logicalSchemaFileLocation),
        ("projection_model_graph", CFG.projectionModelFileLocation),
        ("established_behaviors_graph", CFG.establishedBehaviorsFileLocation),
        ("implementation_patterns_graph", CFG.shaclPatternsFileLocation),
    ):
        path = _location(registry, profile, predicate, base, existing=True)
        models[argument] = Graph().parse(path, format="turtle")
        model_paths.append(path)
    output_root = _location(registry, profile, CFG.buildOutputDirectory, base, existing=False)
    token = uuid4().hex
    directory = output_root / token
    directory.mkdir(parents=True, exist_ok=False)
    run = URIRef(f"https://dynamicontology.com/uad36/construction/run/{token}")
    manifest_path = directory / "run.ttl"
    batch_path = directory / "batch-results.ttl"
    run_configuration = directory / "source-configuration.ttl"
    registry.set((configuration, CFG.activeComplianceSource, workbook))
    registry.set((workbook, CFG.sourceFileLocation, Literal(str(source_path))))
    registry.set((workbook, CFG.sourceFormat, Literal(source_format)))
    registry.serialize(destination=run_configuration, format="turtle", encoding="utf-8")
    manifest = Graph()
    manifest.bind("cv", CV)
    manifest.bind("prov", PROV)
    for triple in (
        (run, RDF.type, CV.ConstraintPipelineRun),
        (run, RDF.type, PROV.Activity),
        (run, CV.sourceWorkbook, workbook),
        (run, CV.sourceWorksheet, Literal(worksheet)),
        (run, CV.rowScope, Literal(row_scope)),
        (run, CV.totalRowCount, Literal(len(unique_ids))),
        (run, CV.completedRowCount, Literal(0)),
        (run, CV.status, Literal("RUNNING")),
        (run, CV.sourceWorkbookArtifact, URIRef(source_path.as_uri())),
        (run, CV.configurationArtifact, URIRef(run_configuration.resolve().as_uri())),
        (run, CV.batchResultsArtifact, URIRef(batch_path.resolve().as_uri())),
        (run, PROV.used, workbook),
        (run, PROV.used, URIRef(source_path.as_uri())),
        (run, PROV.startedAtTime, Literal(datetime.now(timezone.utc))),
    ):
        manifest.add(triple)
    for unique_id in unique_ids:
        manifest.add((run, CV.governedUniqueId, Literal(unique_id)))
    for path in model_paths:
        manifest.add((run, PROV.used, URIRef(path.as_uri())))
    manifest.serialize(destination=manifest_path, format="turtle", encoding="utf-8")
    completed = []
    exceptions = ()
    try:
        if row_scope == "all":
            # IT-58 owns all-row enumeration and independent row exceptions.
            all_rows = constraint_production_pipeline.process_all_governed_source_rows(
                configuration_path=run_configuration,
                output_directory=directory / "constraints", batch_results_path=batch_path,
                **models,
            )
            produced_rows = tuple(all_rows.completed)
            exceptions = tuple(all_rows.exceptions)
            if not batch_path.exists():
                Graph().serialize(destination=batch_path, format="turtle", encoding="utf-8")
            accounted_ids = [row.unique_id for row in produced_rows] + [row.unique_id for row in exceptions]
            if len(accounted_ids) != len(unique_ids) or set(accounted_ids) != set(unique_ids):
                raise ValueError("All-row construction did not account for exactly the selected worksheet rows.")
            for exception in exceptions:
                manifest.add((run, CV.productionException, exception.result_id))
        else:
            # A generator keeps single/group execution and persistence row by row.
            produced_rows = (
                constraint_production_pipeline.process_governed_source_row(
                    configuration_path=run_configuration, unique_id=unique_id,
                    output_directory=directory / "constraints", batch_results_path=batch_path,
                    **models,
                )
                for unique_id in unique_ids
            )
        for produced in produced_rows:
            if not produced.is_complete or produced.unique_id not in unique_ids:
                raise ValueError("Construction did not complete the exact requested governed row.")
            batch = Graph().parse(batch_path, format="turtle")
            constraint = URIRef(produced.constraint_id)
            states = {
                state for state in batch.subjects(RDF.type, CV.ConstraintProductionState)
                if (state, CV.governedConstraint, constraint) in batch
                and (state, CV.status, Literal("COMPLETE")) in batch
            }
            if len(states) != 1:
                raise ValueError("The completed row must have exactly one persistent production state.")
            state = next(iter(states))
            batch.set((state, CV.sourceWorkbook, workbook))
            batch.set((state, CV.sourceWorksheet, Literal(worksheet)))
            batch.serialize(destination=batch_path, format="turtle", encoding="utf-8")
            completed.append(produced)
            manifest.add((run, CV.producedConstraint, constraint))
            manifest.set((run, CV.completedRowCount, Literal(len(completed))))
            manifest.serialize(destination=manifest_path, format="turtle", encoding="utf-8")
        _index_preserved_stages(manifest, run, batch_path, directory)
    except constraint_production_pipeline._RowProductionError as error:
        # IT-58 has already persisted the worker's original exception RDF and
        # stopped dependent workers. Retain that identity and its diagnostic.
        batch = Graph().parse(batch_path, format="turtle")
        candidates = {
            result for result in error.graph.subjects(RDF.type, UADEX.ConstructionStageResult)
            if (result, UADEX.constructionStage, Literal(error.stage)) in error.graph
            and (result, UADEX.constructionStatus, Literal("EXCEPTION")) in error.graph
        }
        if len(candidates) != 1:
            raise ValueError("The failed worker did not return one contracted exception result.") from error
        result_id = next(iter(candidates))
        record = read_construction_stage_result(exchange_graph=batch, result_id=str(result_id))
        if record.construction_status != "EXCEPTION" or record.available_for_downstream:
            raise ValueError("The failed worker result is not a contracted exception.") from error
        constraint = _one(batch, result_id, UADEX.forConstraint)
        failed_id = _text(batch, constraint, CV.governedUniqueId)
        if failed_id not in unique_ids:
            raise ValueError("The worker exception identifies an unselected source row.") from error
        source = _one(batch, constraint, PROV.wasDerivedFrom)
        batch.set((result_id, CV.governedUniqueId, Literal(failed_id)))
        batch.add((result_id, PROV.wasDerivedFrom, source))
        batch.set((result_id, CV.sourceWorkbook, workbook))
        batch.set((result_id, CV.sourceWorksheet, Literal(worksheet)))
        batch.serialize(destination=batch_path, format="turtle", encoding="utf-8")
        # Index the original verified snapshots; never promote the failed
        # worker's output or rerun any earlier production stage.
        _index_preserved_stages(manifest, run, batch_path, directory)
        manifest.set((run, CV.status, Literal("EXCEPTION")))
        manifest.add((run, CV.productionException, result_id))
        manifest.set((run, CV.completedRowCount, Literal(len(completed))))
        manifest.set((run, CV.exceptionRowCount, Literal(len(exceptions) + 1)))
        manifest.set((run, CV.exceptionReason, Literal(record.exception_reason or str(error))))
        manifest.add((run, PROV.endedAtTime, Literal(datetime.now(timezone.utc))))
        manifest.serialize(destination=manifest_path, format="turtle", encoding="utf-8")
        return ConstraintPipelineRunResult(
            run_id=str(run), workbook_id=str(workbook), workbook_label=label,
            worksheet=worksheet, unique_id=unique_ids[0] if len(unique_ids) == 1 else "",
            unique_ids=unique_ids, completed_row_count=len(completed), status="EXCEPTION",
            exception_row_count=len(exceptions) + 1,
            manifest_path=manifest_path, batch_results_path=batch_path, run_directory=directory,
        )
    except Exception as error:
        manifest.set((run, CV.status, Literal("EXCEPTION")))
        manifest.add((run, CV.exceptionReason, Literal(str(error) or type(error).__name__)))
        manifest.add((run, PROV.endedAtTime, Literal(datetime.now(timezone.utc))))
        manifest.serialize(destination=manifest_path, format="turtle", encoding="utf-8")
        raise
    status = "COMPLETE_WITH_EXCEPTIONS" if exceptions else "COMPLETE"
    manifest.set((run, CV.status, Literal(status)))
    manifest.set((run, CV.completedRowCount, Literal(len(completed))))
    manifest.add((run, CV.exceptionRowCount, Literal(len(exceptions))))
    manifest.add((run, PROV.endedAtTime, Literal(datetime.now(timezone.utc))))
    manifest.serialize(destination=manifest_path, format="turtle", encoding="utf-8")
    return ConstraintPipelineRunResult(
        run_id=str(run), workbook_id=str(workbook), workbook_label=label,
        worksheet=worksheet, unique_id=unique_ids[0] if len(unique_ids) == 1 else "",
        unique_ids=unique_ids, completed_row_count=len(completed), status=status,
        exception_row_count=len(exceptions),
        manifest_path=manifest_path, batch_results_path=batch_path, run_directory=directory,
    )
