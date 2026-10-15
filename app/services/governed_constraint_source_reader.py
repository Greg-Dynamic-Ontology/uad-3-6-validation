"""Read governed constraint rows from the authoritative CSV or workbook."""

from __future__ import annotations

import csv
from pathlib import Path
from typing import Any

from openpyxl import load_workbook
from openpyxl.utils.cell import range_boundaries


EXPECTED_COLUMNS = (
    "Unique ID",
    "Primary Data Element",
    "Message ID",
    "Message Text",
    "Rule Logic",
    "Severity",
    "Property Affected",
    "Report Section",
    "Report Subsection",
    "Report Label / Value",
    "Data Point Name / Value",
    " xPath",
    "Related Value(s)",
    "Date Format",
    "Number Format",
    "Min Value",
    "Max Value",
)

REQUIRED_GOVERNED_KNOWLEDGE_COLUMNS = (
    "Primary Data Element",
    "Message ID",
    "Message Text",
    "Rule Logic",
    "Severity",
)


def _validate_headers(headers: tuple[str, ...] | list[str]) -> None:
    actual = tuple("" if header is None else str(header) for header in headers)
    if actual != EXPECTED_COLUMNS:
        missing = [column for column in EXPECTED_COLUMNS if column not in actual]
        unexpected = [column for column in actual if column not in EXPECTED_COLUMNS]

        details = []
        if missing:
            details.append(f"missing columns: {missing}")
        if unexpected:
            details.append(f"unexpected columns: {unexpected}")
        if not details:
            details.append("columns are not in the governed order")

        raise ValueError(
            "Governed constraint source header does not match the expected "
            f"column structure ({'; '.join(details)})."
        )


def _cell_text(value: Any) -> str:
    if value is None:
        return ""
    return str(value).strip()


def _make_governed_row(
    values: tuple[Any, ...] | list[Any],
    *,
    source_locator: str,
    source_constraint_uri: str,
) -> dict[str, Any]:
    if len(values) != len(EXPECTED_COLUMNS):
        raise ValueError(
            "Governed constraint source row does not match the expected "
            f"column structure: expected {len(EXPECTED_COLUMNS)} values, "
            f"received {len(values)}."
        )

    row = {
        column: _cell_text(value)
        for column, value in zip(EXPECTED_COLUMNS, values)
    }

    for column in REQUIRED_GOVERNED_KNOWLEDGE_COLUMNS:
        if not row[column]:
            raise ValueError(
                f"Required governed knowledge field {column!r} is blank "
                f"in source row {source_locator!r}."
            )

    return {
        "unique_id": row["Unique ID"],
        "primary_data_element": row["Primary Data Element"],
        "source_rule_id": row["Message ID"],
        "requirement_text": row["Message Text"],
        "violation_condition": row["Rule Logic"],
        "severity": row["Severity"],
        "property_affected": row["Property Affected"],
        "report_section": row["Report Section"],
        "report_subsection": row["Report Subsection"],
        "report_label_or_value": row["Report Label / Value"],
        "data_point_name_or_value": row["Data Point Name / Value"],
        "xpath": row[" xPath"],
        "related_values": row["Related Value(s)"],
        "date_format": row["Date Format"],
        "number_format": row["Number Format"],
        "min_value": row["Min Value"],
        "max_value": row["Max Value"],
        "context_fields": tuple(row[column] for column in EXPECTED_COLUMNS[6:11]),
        "source_locator": source_locator,
        "source_constraint_uri": source_constraint_uri,
    }


def _source_locator(
    workbook_path: str | Path,
    worksheet_name: str,
    cell_range: str,
) -> str:
    return f"{Path(workbook_path).name}#{worksheet_name}!{cell_range}"


def read_governed_constraint_row(
    workbook_path: str | Path,
    worksheet_name: str,
    cell_range: str,
) -> dict[str, Any]:
    """Read one governed row from a workbook using an explicit worksheet range."""
    workbook = load_workbook(
        filename=workbook_path,
        read_only=True,
        data_only=True,
    )
    try:
        if worksheet_name not in workbook.sheetnames:
            raise ValueError(
                f"Governed worksheet {worksheet_name!r} was not found."
            )

        worksheet = workbook[worksheet_name]
        min_col, min_row, max_col, max_row = range_boundaries(cell_range)

        if min_row != max_row:
            raise ValueError(
                "The governed workbook cell range must identify exactly one row."
            )

        header_values = next(
            worksheet.iter_rows(
                min_row=1,
                max_row=1,
                min_col=min_col,
                max_col=max_col,
                values_only=True,
            )
        )
        _validate_headers(header_values)

        values = next(
            worksheet.iter_rows(
                min_row=min_row,
                max_row=max_row,
                min_col=min_col,
                max_col=max_col,
                values_only=True,
            )
        )

        locator = _source_locator(workbook_path, worksheet_name, cell_range)
        return _make_governed_row(
            values,
            source_locator=locator,
            source_constraint_uri=f"urn:uad36:governed-source:{locator}",
        )
    finally:
        workbook.close()


def read_governed_constraint_row_from_csv(
    csv_path: str | Path,
    unique_id: str,
) -> dict[str, Any]:
    """Read one uniquely selected governed row from the authoritative CSV."""
    requested_id = _cell_text(unique_id)
    if not requested_id:
        raise ValueError("A non-empty governed Unique ID is required.")

    path = Path(csv_path)
    matches: list[tuple[int, list[str]]] = []

    with path.open("r", encoding="utf-8-sig", newline="") as source_file:
        reader = csv.reader(source_file)
        try:
            headers = next(reader)
        except StopIteration as error:
            raise ValueError("Governed constraint CSV is empty.") from error

        _validate_headers(headers)

        for row_number, values in enumerate(reader, start=2):
            if len(values) != len(EXPECTED_COLUMNS):
                raise ValueError(
                    f"Governed constraint CSV row {row_number} does not match "
                    "the expected column structure."
                )
            if _cell_text(values[0]) == requested_id:
                matches.append((row_number, values))

    if not matches:
        raise ValueError(
            f"No governed constraint row was found for Unique ID "
            f"{requested_id!r}."
        )
    if len(matches) > 1:
        raise ValueError(
            f"More than one governed constraint row was found for Unique ID "
            f"{requested_id!r}."
        )

    row_number, values = matches[0]
    locator = f"{path.name}#row={row_number}"
    return _make_governed_row(
        values,
        source_locator=locator,
        source_constraint_uri=f"urn:uad36:governed-source:{locator}",
    )


__all__ = [
    "EXPECTED_COLUMNS",
    "REQUIRED_GOVERNED_KNOWLEDGE_COLUMNS",
    "read_governed_constraint_row",
    "read_governed_constraint_row_from_csv",
]