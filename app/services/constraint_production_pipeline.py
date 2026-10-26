"""Coordinate configured governed rows through constraint construction.

Construction consumes source and model knowledge, never appraisal instances.
Each run preserves a Turtle snapshot after every verified worker handoff.
"""

from __future__ import annotations

import csv
from dataclasses import dataclass
from datetime import datetime, timezone
from hashlib import sha256
from pathlib import Path
from uuid import uuid4

from openpyxl import load_workbook
from openpyxl.utils import get_column_letter
from rdflib import BNode, Graph, Literal, Namespace, RDF, URIRef
from rdflib.compare import isomorphic
from rdflib.namespace import PROV, SH

from app.services import (
    constraint_behavior_classification,
    constraint_normalization_agent,
    context_resolution_agent,
    instance_rdf_binding_agent,
    logical_schema_binding_agent,
    shacl_generation_agent,
)
from app.services.construction_exchange import (
    UADEX,
    exchange_construction_stage_exception,
    exchange_successful_construction_stage_result,
    read_construction_stage_result,
)
from app.services.governed_constraint_source_reader import (
    EXPECTED_COLUMNS,
    _make_governed_row,
    _validate_headers,
    read_governed_constraint_row,
    read_governed_constraint_row_from_csv,
)


CFG = Namespace("https://dynamicontology.com/uad36/configuration#")
CV = Namespace("https://dynamicontology.com/uad36/constraint-vocabulary#")
STAGES = (
    "normalized-constraint-representation",
    "logical-schema-binding",
    "context-resolution",
    "validation-behavior",
    "instance-rdf-binding",
    "shacl-representation",
)


@dataclass(frozen=True)
class ConstraintProductionResult:
    """Saved reusable artifacts and the ordered snapshots for one run."""

    constraint_id: str
    unique_id: str
    is_complete: bool
    rdf_path: Path
    shacl_path: Path
    batch_results_path: Path
    stage_paths: tuple[Path, ...]


@dataclass(frozen=True)
class ConstraintProductionException:
    """One persistent row exception, including failure before normalization."""

    unique_id: str
    result_id: URIRef
    exception_reason: str
    source_constraint_uri: str
    construction_stage: str


@dataclass(frozen=True)
class ConstraintBatchProductionResult:
    """Independent successful artifacts and exceptions for an all-row request."""

    completed: tuple[ConstraintProductionResult, ...]
    exceptions: tuple[ConstraintProductionException, ...]
    batch_results_path: Path


class _RowProductionError(ValueError):
    """Carry unsuccessful worker RDF to the batch without advancing the row."""

    def __init__(self, reason: str, stage: str, graph: Graph) -> None:
        super().__init__(reason)
        self.stage = stage
        self.graph = graph


@dataclass(frozen=True)
class _ResumeState:
    constraint_id: str
    current: Graph
    batch: Graph
    first_stage_index: int
    previous_stage_paths: tuple[Path, ...]
    exception_id: URIRef
    resolution_id: URIRef


def _one(graph: Graph, subject, predicate):
    values = set(graph.objects(subject, predicate))
    if len(values) != 1:
        raise ValueError(f"{subject} must identify exactly one {predicate}.")
    return next(iter(values))


def _copy(source: Graph, destination: Graph) -> None:
    for prefix, namespace in source.namespaces():
        destination.bind(prefix, namespace, replace=False)
    for triple in source:
        destination.add(triple)


def _configured_source(configuration_path: Path) -> tuple[Path, str, str | None]:
    configuration = Graph().parse(configuration_path, format="turtle")
    roots = set(configuration.subjects(RDF.type, CFG.ApplicationConfiguration))
    if len(roots) != 1:
        raise ValueError("Configuration must identify one ApplicationConfiguration.")
    root = next(iter(roots))
    source = _one(configuration, root, CFG.activeComplianceSource)
    collection = _one(configuration, root, CFG.hasSourceWorkbooks)
    if not isinstance(source, URIRef) or (
        collection, CFG.containsWorkbook, source
    ) not in configuration:
        raise ValueError("Active compliance source must belong to the configured collection.")
    location = _one(configuration, source, CFG.sourceFileLocation)
    source_format = _one(configuration, source, CFG.sourceFormat)
    if not all(isinstance(value, Literal) and str(value).strip()
               for value in (location, source_format)):
        raise ValueError("Source location and format must be nonblank literals.")
    path = Path(str(location))
    if not path.is_absolute():
        path = configuration_path.parent / path
    path = path.resolve(strict=True)
    format_name = str(source_format).strip().lower()
    if format_name not in {"csv", "xlsx", "xlsm"}:
        raise ValueError(f"Unsupported governed source format: {format_name!r}.")
    sheet = None
    if format_name != "csv":
        sheet_value = _one(configuration, source, CFG.ruleDataSheetName)
        if not isinstance(sheet_value, Literal) or not str(sheet_value).strip():
            raise ValueError("Excel source must identify a nonblank ruleDataSheetName.")
        sheet = str(sheet_value)
    return path, format_name, sheet


def _read_selected_row(configuration_path: Path, unique_id: str) -> dict:
    path, source_format, worksheet_name = _configured_source(configuration_path)
    if source_format == "csv":
        row = read_governed_constraint_row_from_csv(path, unique_id)
    else:
        # Inspect identifiers only. Insufficient knowledge in an unselected
        # row must not prevent construction of the requested row.
        workbook = load_workbook(path, read_only=True, data_only=True)
        try:
            if worksheet_name not in workbook.sheetnames:
                raise ValueError(f"Configured worksheet {worksheet_name!r} does not exist.")
            sheet = workbook[worksheet_name]
            headers = tuple(next(sheet.iter_rows(min_row=1, max_row=1, values_only=True)))
            if headers != EXPECTED_COLUMNS:
                raise ValueError("Configured worksheet headers do not match the governed structure.")
            matches = [
                number
                for number, values in enumerate(
                    sheet.iter_rows(min_row=2, min_col=1, max_col=1, values_only=True),
                    start=2,
                )
                if values[0] is not None and str(values[0]).strip() == unique_id
            ]
        finally:
            workbook.close()
        if len(matches) != 1:
            raise ValueError(
                f"Unique ID {unique_id!r} must select exactly one governed row; "
                f"found {len(matches)}."
            )
        number = matches[0]
        last_column = get_column_letter(len(EXPECTED_COLUMNS))
        row = read_governed_constraint_row(
            workbook_path=path,
            worksheet_name=worksheet_name,
            cell_range=f"A{number}:{last_column}{number}",
        )
    # The reader's explicit source field and IT-50's normalized input field
    # describe the same data point. Preserve all original reader knowledge.
    source = dict(row)
    source["data_point_name"] = row["primary_data_element"]
    return source


def _verify_stage(graph: Graph, constraint_id: str, stage: str):
    constraint = URIRef(constraint_id)
    results = {
        result
        for result in graph.subjects(UADEX.forConstraint, constraint)
        if (result, UADEX.constructionStage, Literal(stage)) in graph
    }
    if len(results) != 1:
        raise ValueError(f"Stage {stage!r} must publish exactly one result for this constraint.")
    record = read_construction_stage_result(
        exchange_graph=graph, result_id=str(next(iter(results)))
    )
    if record.construction_status != "GREEN" or not record.available_for_downstream:
        raise ValueError(
            f"Stage {stage!r} is not verified for downstream use: "
            f"{record.exception_reason or record.construction_status}"
        )
    return record


def process_governed_source_row(
    *,
    configuration_path: str | Path,
    unique_id: str,
    logical_schema_graph: Graph,
    projection_model_graph: Graph,
    established_behaviors_graph: Graph,
    implementation_patterns_graph: Graph,
    output_directory: str | Path,
    batch_results_path: str | Path,
) -> ConstraintProductionResult:
    """Select exactly one row and persist verified construction artifacts.

    Relative source locations are resolved beside the configuration file.
    Supported formats are CSV and modern Excel workbooks (xlsx and xlsm).
    This operation neither validates an appraisal nor projects findings.
    """
    if not isinstance(unique_id, str) or not unique_id.strip():
        raise ValueError("unique_id must be a nonblank string, preserving leading zeros.")
    unique_id = unique_id.strip()
    source = _read_selected_row(Path(configuration_path).resolve(strict=True), unique_id)
    return _construct_governed_source(
        source=source,
        logical_schema_graph=logical_schema_graph,
        projection_model_graph=projection_model_graph,
        established_behaviors_graph=established_behaviors_graph,
        implementation_patterns_graph=implementation_patterns_graph,
        output_directory=output_directory,
        batch_results_path=batch_results_path,
    )


def _construct_governed_source(
    *,
    source: dict,
    logical_schema_graph: Graph,
    projection_model_graph: Graph,
    established_behaviors_graph: Graph,
    implementation_patterns_graph: Graph,
    output_directory: str | Path,
    batch_results_path: str | Path,
    resume_state: _ResumeState | None = None,
) -> ConstraintProductionResult:
    """Run the shared worker sequence for an already selected source row."""
    unique_id = source["unique_id"]
    run_directory = Path(output_directory) / uuid4().hex
    run_directory.mkdir(parents=True, exist_ok=False)
    batch_path = Path(batch_results_path)
    source_iri = URIRef(source["source_constraint_uri"])
    if resume_state is None:
        batch = Graph()
        if batch_path.exists():
            batch.parse(batch_path, format="turtle")
        stage_paths: list[Path] = []
        normalized_path = run_directory / "01-normalized-constraint.ttl"
        stage_started_at = datetime.now(timezone.utc)
        normalized = constraint_normalization_agent.normalize_governed_constraint_to_file(
            governed_source=source, output_path=normalized_path,
        )
        constraint_id = normalized.constraint_id
        constraint = URIRef(constraint_id)
        current = Graph().parse(normalized_path, format="turtle")
        if (constraint, RDF.type, CV.NormalizedConstraint) not in current or set(
            current.objects(constraint, CV.governedUniqueId)
        ) != {Literal(unique_id)} or (constraint, PROV.wasDerivedFrom, source_iri) not in current:
            raise ValueError("Normalization output does not preserve selected source identity.")
        normalization_result = exchange_successful_construction_stage_result(
            exchange_graph=current, constraint_id=constraint_id, construction_stage=STAGES[0],
        )
        current.add((normalization_result.result_id, UADEX.stageKnowledge, constraint))
        first_stage_index = 1
    else:
        batch = resume_state.batch
        current = resume_state.current
        constraint_id = resume_state.constraint_id
        constraint = URIRef(constraint_id)
        stage_paths = list(resume_state.previous_stage_paths)
        first_stage_index = resume_state.first_stage_index

    def persist_stage(graph: Graph, index: int) -> None:
        record = _verify_stage(graph, constraint_id, STAGES[index])
        graph.add((record.result_id, PROV.wasDerivedFrom, source_iri))
        # The coordinator observes this actual invocation. Capture its provenance
        # before serializing the verified snapshot; never reconstruct it on replay.
        activities = set(graph.objects(record.result_id, PROV.wasGeneratedBy))
        if not activities:
            activity = URIRef(f"{record.result_id}/execution")
            agent = URIRef(f"https://dynamicontology.com/uad36/construction/worker/{STAGES[index]}")
            graph.add((record.result_id, PROV.wasGeneratedBy, activity))
            graph.add((activity, RDF.type, PROV.Activity))
            graph.add((activity, PROV.wasAssociatedWith, agent))
            graph.add((agent, RDF.type, PROV.Agent))
            graph.add((activity, PROV.startedAtTime, Literal(stage_started_at)))
            graph.add((activity, PROV.endedAtTime, Literal(datetime.now(timezone.utc))))
            graph.add((activity, PROV.used, source_iri))
            if index:
                upstream = _verify_stage(graph, constraint_id, STAGES[index - 1])
                graph.add((activity, PROV.used, upstream.result_id))
        snapshot = run_directory / f"{index + 1:02d}-{STAGES[index]}.ttl"
        if index == 0:
            snapshot = run_directory / "01-normalized-constraint.ttl"
        graph.add((record.result_id, UADEX.rdfSnapshotPath, Literal(str(snapshot.resolve()))))
        if resume_state is not None and index == first_stage_index:
            graph.add((record.result_id, PROV.wasRevisionOf, resume_state.exception_id))
            graph.add((record.result_id, PROV.wasDerivedFrom, resume_state.resolution_id))
        graph.serialize(destination=snapshot, format="turtle", encoding="utf-8")
        stage_paths.append(snapshot)
        _copy(graph, batch)
        batch.set((record.result_id, UADEX.snapshotSha256,
                   Literal(sha256(snapshot.read_bytes()).hexdigest())))
        batch_path.parent.mkdir(parents=True, exist_ok=True)
        batch.serialize(destination=batch_path, format="turtle", encoding="utf-8")

    if resume_state is None:
        persist_stage(current, 0)
    workers = (
        (logical_schema_binding_agent.bind_normalized_constraint_to_logical_schema,
         "normalized_graph", {"logical_schema_graph": logical_schema_graph}),
        (context_resolution_agent.resolve_bound_constraint_context,
         "bound_graph", {"logical_schema_graph": logical_schema_graph}),
        (constraint_behavior_classification.classify_governed_constraint_behavior,
         "context_graph", {"established_behaviors_graph": established_behaviors_graph}),
        (instance_rdf_binding_agent.bind_governed_constraint_to_instance_rdf,
         "classified_graph", {"projection_model_graph": projection_model_graph}),
        (shacl_generation_agent.generate_governed_constraint_shacl,
         "bound_graph", {
             "implementation_patterns_graph": implementation_patterns_graph,
             "output_path": run_directory / "shapes.ttl",
         }),
    )
    for index, (worker, input_name, arguments) in enumerate(workers, start=1):
        if index < first_stage_index:
            continue
        _verify_stage(current, constraint_id, STAGES[index - 1])
        output = Graph()
        stage_started_at = datetime.now(timezone.utc)
        try:
            worker(**{
                input_name: current,
                "constraint_id": constraint_id,
                "output_graph": output,
                **arguments,
            })
            # Downstream output must preserve the verified upstream RDF verbatim.
            if not set(current).issubset(set(output)):
                raise ValueError(f"Stage {STAGES[index]!r} discarded upstream knowledge.")
            persist_stage(output, index)
        except ValueError as error:
            # Retain the actual worker result before stopping this row. Its
            # original identity, status, availability and knowledge are audit
            # state; persisting them does not authorize downstream consumption.
            snapshot = run_directory / f"{index + 1:02d}-{STAGES[index]}-unverified.ttl"
            output.serialize(destination=snapshot, format="turtle", encoding="utf-8")
            _copy(output, batch)
            batch_path.parent.mkdir(parents=True, exist_ok=True)
            batch.serialize(destination=batch_path, format="turtle", encoding="utf-8")
            raise _RowProductionError(str(error), STAGES[index], output) from error
        current = output

    shacl_path = run_directory / "shapes.ttl"
    shapes = Graph().parse(shacl_path, format="turtle")
    shape_ids = set(shapes.subjects(RDF.type, SH.NodeShape))
    if len(shape_ids) != 1:
        raise ValueError("SHACL artifact must contain one governed NodeShape.")
    shape = next(iter(shape_ids))
    if (shape, CV.implementsConstraint, constraint) not in shapes or (
        shape, PROV.wasDerivedFrom, source_iri
    ) not in shapes:
        raise ValueError("SHACL artifact must retain selected constraint and source traceability.")
    for stage in STAGES:
        _verify_stage(current, constraint_id, stage)
    rdf_path = run_directory / "constraint.ttl"
    current.serialize(destination=rdf_path, format="turtle", encoding="utf-8")
    _persist_completed_production(
        source=source,
        current=current,
        batch=batch,
        rdf_path=rdf_path,
        shacl_path=shacl_path,
        batch_path=batch_path,
        constraint_id=constraint_id,
    )
    return ConstraintProductionResult(
        constraint_id=constraint_id,
        unique_id=unique_id,
        is_complete=True,
        rdf_path=rdf_path,
        shacl_path=shacl_path,
        batch_results_path=batch_path,
        stage_paths=tuple(stage_paths),
    )


def _source_entries(configuration_path: Path):
    """Read the configured source once, retaining physical row locations.

    Validate file-wide structure here; required row knowledge is checked
    separately inside the row exception boundary. Blank Excel padding is
    ignored, but a populated row with no identifier is still considered.
    """
    path, source_format, worksheet_name = _configured_source(configuration_path)
    if source_format == "csv":
        with path.open("r", encoding="utf-8-sig", newline="") as stream:
            reader = csv.reader(stream)
            try:
                headers = next(reader)
            except StopIteration as error:
                raise ValueError("Governed constraint CSV is empty.") from error
            _validate_headers(headers)
            for number, values in enumerate(reader, start=2):
                locator = f"{path.name}#row={number}"
                yield tuple(values), locator
    else:
        workbook = load_workbook(path, read_only=True, data_only=True)
        try:
            if worksheet_name not in workbook.sheetnames:
                raise ValueError(f"Configured worksheet {worksheet_name!r} does not exist.")
            sheet = workbook[worksheet_name]
            iterator = sheet.iter_rows(values_only=True)
            try:
                headers = next(iterator)
            except StopIteration as error:
                raise ValueError("Configured governed worksheet is empty.") from error
            _validate_headers(headers)
            last_column = get_column_letter(len(EXPECTED_COLUMNS))
            for number, values in enumerate(iterator, start=2):
                if all(value is None or not str(value).strip() for value in values):
                    continue
                locator = f"{path.name}#{worksheet_name}!A{number}:{last_column}{number}"
                yield tuple(values), locator
        finally:
            workbook.close()


def _record_row_exception(
    *,
    unique_id: str,
    source_uri: str,
    error: ValueError,
    stage: str,
    batch_path: Path,
) -> ConstraintProductionException:
    batch = Graph()
    if batch_path.exists():
        batch.parse(batch_path, format="turtle")
    failed_graph = error.graph if isinstance(error, _RowProductionError) else Graph()
    # _construct_governed_source saved this worker output before raising
    # _RowProductionError. Read its result by governed IRI from that saved
    # batch; do not merge the in-memory graph again. Turtle parsing renames
    # blank nodes, so a second merge would duplicate RDF list branches.
    constraints = {
        subject for subject in batch.subjects(CV.governedUniqueId, Literal(unique_id))
        if (subject, RDF.type, CV.NormalizedConstraint) in batch
        and (subject, PROV.wasDerivedFrom, URIRef(source_uri)) in batch
    }
    if len(constraints) > 1:
        raise ValueError("Failed row has more than one normalized constraint identity.")
    # Before normalization, the governed source resource is the only known
    # constraint identity. Do not manufacture a normalized constraint.
    constraint = next(iter(constraints), URIRef(source_uri))
    candidates = {
        result for result in failed_graph.subjects(UADEX.forConstraint, constraint)
        if (result, UADEX.constructionStage, Literal(stage)) in failed_graph
        and (result, UADEX.constructionStatus, Literal("EXCEPTION")) in failed_graph
    }
    if len(candidates) > 1:
        raise ValueError("Worker produced conflicting exception result identities.")
    if candidates:
        record = read_construction_stage_result(
            exchange_graph=batch, result_id=str(next(iter(candidates))),
        )
    else:
        record = exchange_construction_stage_exception(
            exchange_graph=batch,
            constraint_id=str(constraint),
            construction_stage=stage,
            exception_reason=str(error) or type(error).__name__,
            requires_review=True,
        )
    batch.add((record.result_id, CV.governedUniqueId, Literal(unique_id)))
    batch.add((record.result_id, PROV.wasDerivedFrom, URIRef(source_uri)))
    batch_path.parent.mkdir(parents=True, exist_ok=True)
    batch.serialize(destination=batch_path, format="turtle", encoding="utf-8")
    return ConstraintProductionException(
        unique_id=unique_id,
        result_id=record.result_id,
        exception_reason=record.exception_reason or str(error),
        source_constraint_uri=source_uri,
        construction_stage=stage,
    )


def process_all_governed_source_rows(
    *,
    configuration_path: str | Path,
    logical_schema_graph: Graph,
    projection_model_graph: Graph,
    established_behaviors_graph: Graph,
    implementation_patterns_graph: Graph,
    output_directory: str | Path,
    batch_results_path: str | Path,
) -> ConstraintBatchProductionResult:
    """Consider every source row and continue after row knowledge exceptions.

    Source/configuration errors and filesystem errors remain request-level
    failures. Row knowledge errors are persisted independently. Neither
    appraisal XML nor validation execution is part of this operation.
    """
    configuration_path = Path(configuration_path).resolve(strict=True)
    # Finish source enumeration before starting construction, so a bad file
    # structure cannot leave an ambiguously partial all-row request.
    entries = list(_source_entries(configuration_path))
    completed: list[ConstraintProductionResult] = []
    exceptions: list[ConstraintProductionException] = []
    batch_path = Path(batch_results_path)
    for values, locator in entries:
        source_uri = f"urn:uad36:governed-source:{locator}"
        unique_id = str(values[0]).strip() if values and values[0] is not None else ""
        stage = "source-reading"
        try:
            if not unique_id:
                raise ValueError("A governed source row must have a nonblank Unique ID.")
            source = _make_governed_row(
                values, source_locator=locator, source_constraint_uri=source_uri,
            )
            source["data_point_name"] = source["primary_data_element"]
            stage = STAGES[0]
            result = _construct_governed_source(
                source=source,
                logical_schema_graph=logical_schema_graph,
                projection_model_graph=projection_model_graph,
                established_behaviors_graph=established_behaviors_graph,
                implementation_patterns_graph=implementation_patterns_graph,
                output_directory=output_directory,
                batch_results_path=batch_path,
            )
        except ValueError as error:
            if isinstance(error, _RowProductionError):
                stage = error.stage
            exceptions.append(_record_row_exception(
                unique_id=unique_id,
                source_uri=source_uri,
                error=error,
                stage=stage,
                batch_path=batch_path,
            ))
        else:
            completed.append(result)
    return ConstraintBatchProductionResult(
        completed=tuple(completed),
        exceptions=tuple(exceptions),
        batch_results_path=batch_path,
    )


def _persist_completed_production(
    *,
    source: dict,
    current: Graph,
    batch: Graph,
    rdf_path: Path,
    shacl_path: Path,
    batch_path: Path,
    constraint_id: str,
) -> URIRef:
    """Save completion evidence only after verifying artifacts and history.

    Completion is orchestration knowledge in the batch graph. It is not
    added to worker-produced domain artifacts and requires no appraisal.
    """
    constraint = URIRef(constraint_id)
    source_iri = URIRef(source["source_constraint_uri"])
    if not rdf_path.is_file() or not shacl_path.is_file():
        raise ValueError("Completion requires both saved RDF and SHACL artifacts.")
    saved_rdf = Graph().parse(rdf_path, format="turtle")
    saved_shapes = Graph().parse(shacl_path, format="turtle")
    if not isomorphic(current, saved_rdf):
        raise ValueError("Saved constraint RDF differs from verified worker output.")
    if set(saved_rdf.subjects(RDF.type, CV.NormalizedConstraint)) != {constraint} or (
        constraint, PROV.wasDerivedFrom, source_iri
    ) not in saved_rdf or set(saved_rdf.objects(constraint, CV.governedUniqueId)) != {
        Literal(source["unique_id"])
    }:
        raise ValueError("Saved RDF does not preserve the completed source row identity.")
    shape_ids = set(saved_shapes.subjects(RDF.type, SH.NodeShape))
    if len(shape_ids) != 1:
        raise ValueError("Completion requires one saved governed SHACL representation.")
    shape = next(iter(shape_ids))
    if (shape, CV.implementsConstraint, constraint) not in saved_shapes or (
        shape, PROV.wasDerivedFrom, source_iri
    ) not in saved_shapes:
        raise ValueError("Saved SHACL does not preserve governed source traceability.")
    verified = [_verify_stage(saved_rdf, constraint_id, stage) for stage in STAGES]
    for record in verified:
        persisted = read_construction_stage_result(
            exchange_graph=batch, result_id=str(record.result_id),
        )
        if persisted != record:
            raise ValueError("Batch stage state differs from the verified artifact evidence.")
    related_results = set(batch.subjects(UADEX.forConstraint, constraint))
    related_results.update(batch.subjects(UADEX.forConstraint, source_iri))
    for result_id in related_results:
        record = read_construction_stage_result(exchange_graph=batch, result_id=str(result_id))
        if record.construction_status == "GREEN" and record.available_for_downstream:
            continue
        if record.construction_status != "EXCEPTION":
            raise ValueError("Completion is blocked by an unverified production state.")
        resolutions = set(batch.objects(result_id, CV.reviewResolution))
        if len(resolutions) != 1:
            raise ValueError("Completion is blocked by an unresolved production exception.")
        resolution = next(iter(resolutions))
        if not isinstance(resolution, URIRef):
            raise ValueError("An exception resolution must retain its governed RDF identity.")
        resolved = _reviewed_exception(
            batch, resolution, URIRef(record.constraint_id), require_checkpoint=False,
        )
        if resolved.result_id != result_id:
            raise ValueError("Governed review resolved a different production exception.")
    state = URIRef(f"{constraint}/production-state/{rdf_path.parent.name}")
    for triple in (
        (state, RDF.type, CV.ConstraintProductionState),
        (state, CV.governedConstraint, constraint),
        (state, CV.governedUniqueId, Literal(source["unique_id"])),
        (state, CV.status, Literal("COMPLETE")),
        (state, CV.rdfArtifact, URIRef(rdf_path.resolve().as_uri())),
        (state, CV.shaclArtifact, URIRef(shacl_path.resolve().as_uri())),
        (state, PROV.wasDerivedFrom, source_iri),
    ):
        batch.add(triple)
    for record in verified:
        batch.add((state, CV.verifiedStageResult, record.result_id))
    batch.serialize(destination=batch_path, format="turtle", encoding="utf-8")
    return state


def _reviewed_exception(
    batch: Graph,
    resolution: URIRef,
    constraint: URIRef,
    *,
    require_checkpoint: bool = True,
):
    if (resolution, RDF.type, CV.GovernedReviewResolution) not in batch:
        raise ValueError("Resume requires a recorded GovernedReviewResolution.")
    if _one(batch, resolution, CV.reviewStatus) != Literal("RESOLVED"):
        raise ValueError("Governed review must explicitly resolve the exception.")
    exception_id = _one(batch, resolution, CV.resolvesException)
    if not isinstance(exception_id, URIRef) or (
        exception_id, CV.reviewResolution, resolution
    ) not in batch:
        raise ValueError("Review resolution must identify its recorded exception.")
    activity = _one(batch, resolution, PROV.wasGeneratedBy)
    if (activity, RDF.type, PROV.Activity) not in batch or (
        activity, PROV.used, exception_id
    ) not in batch:
        raise ValueError("Resolution must preserve its governed review activity.")
    reviewers = set(batch.objects(activity, PROV.wasAssociatedWith))
    if not reviewers or any((reviewer, RDF.type, PROV.Agent) not in batch for reviewer in reviewers):
        raise ValueError("Governed review must identify its responsible agents.")
    record = read_construction_stage_result(exchange_graph=batch, result_id=str(exception_id))
    if record.constraint_id != str(constraint) or record.construction_status != "EXCEPTION":
        raise ValueError("Resolution must refer to an exception for the selected constraint.")
    if require_checkpoint and record.construction_stage not in STAGES[1:]:
        raise ValueError("Resume requires an exception after a verified normalization checkpoint.")
    return record


def _load_verified_checkpoint(batch: Graph, record, output_root: Path) -> tuple[Path, Graph]:
    locations = set(batch.objects(record.result_id, UADEX.rdfSnapshotPath))
    if len(locations) > 1:
        raise ValueError("A verified stage has conflicting checkpoint locations.")
    if locations:
        candidates = [Path(str(next(iter(locations)))).resolve(strict=True)]
    else:
        # Older runs predate explicit checkpoint-location metadata. Identify
        # their checkpoint by exact saved result IRI, never by filename alone.
        number = STAGES.index(record.construction_stage) + 1
        candidates = [
            path.resolve() for path in output_root.rglob(f"{number:02d}-*.ttl")
            if "-unverified" not in path.name
        ]
    matches = []
    for path in candidates:
        if not path.is_relative_to(output_root) or not path.is_file():
            raise ValueError("Verified checkpoint must be a file inside the configured artifact directory.")
        graph = Graph().parse(path, format="turtle")
        if (record.result_id, RDF.type, UADEX.ConstructionStageResult) in graph:
            matches.append((path, graph))
    if len(matches) != 1:
        raise ValueError("The verified stage must identify exactly one saved checkpoint.")
    path, graph = matches[0]
    checksums = set(batch.objects(record.result_id, UADEX.snapshotSha256))
    if len(checksums) > 1 or (checksums and str(next(iter(checksums))) != sha256(path.read_bytes()).hexdigest()):
        raise ValueError("Verified checkpoint contents no longer match the saved checksum.")
    saved = _verify_stage(graph, record.constraint_id, record.construction_stage)
    if saved.result_id != record.result_id:
        raise ValueError("Checkpoint does not preserve the verified stage identity.")
    return path, graph


def _checkpoint_with_batch_blank_nodes(checkpoint: Graph, batch: Graph) -> Graph:
    """Recover checkpoint knowledge using the already loaded batch nodes.

    Blank-node labels differ between Turtle documents. Align their anchored
    structures before adding new knowledge, so merging does not duplicate
    existing RDF lists. All named IRIs and literal values remain exact.
    """
    named = {
        triple for triple in checkpoint
        if not any(isinstance(term, BNode) for term in triple)
    }
    if not named.issubset(set(batch)):
        raise ValueError("Verified checkpoint knowledge is missing from the Batch Results Graph.")
    roots = {
        (subject, predicate) for subject, predicate, value in checkpoint
        if not isinstance(subject, BNode) and isinstance(value, BNode)
    }

    def components(graph: Graph) -> Graph:
        result = Graph()
        pending = []
        visited = set()
        for subject, predicate in roots:
            for value in graph.objects(subject, predicate):
                result.add((subject, predicate, value))
                if isinstance(value, BNode):
                    pending.append(value)
        while pending:
            node = pending.pop()
            if node in visited:
                continue
            visited.add(node)
            for triple in graph.triples((node, None, None)):
                result.add(triple)
                if isinstance(triple[2], BNode):
                    pending.append(triple[2])
        return result

    expected = components(checkpoint)
    actual = components(batch)
    if any(triple not in expected for triple in checkpoint
           if any(isinstance(term, BNode) for term in triple)):
        raise ValueError("Checkpoint contains unanchored blank-node knowledge.")
    if not isomorphic(expected, actual):
        raise ValueError("Batch RDF structure does not preserve the verified checkpoint.")
    current = Graph()
    for prefix, namespace in checkpoint.namespaces():
        current.bind(prefix, namespace)
    for triple in named:
        current.add(triple)
    _copy(actual, current)
    return current


def resume_governed_source_row(
    *,
    configuration_path: str | Path,
    unique_id: str,
    review_resolution_id: str,
    logical_schema_graph: Graph,
    projection_model_graph: Graph,
    established_behaviors_graph: Graph,
    implementation_patterns_graph: Graph,
    output_directory: str | Path,
    batch_results_path: str | Path,
) -> ConstraintProductionResult:
    """Restart a reviewed exception from its last verified saved checkpoint.

    The batch retains historical exception and review knowledge. Only the
    verified checkpoint is supplied to workers; no unverified stage output
    is promoted. Earlier workers and snapshot files are left untouched.
    """
    if not isinstance(unique_id, str) or not unique_id.strip():
        raise ValueError("Resume requires a nonblank textual Unique ID.")
    if not isinstance(review_resolution_id, str) or not review_resolution_id.strip():
        raise ValueError("Resume requires the governed review-resolution IRI.")
    source = _read_selected_row(Path(configuration_path).resolve(strict=True), unique_id.strip())
    batch_path = Path(batch_results_path)
    batch = Graph().parse(batch_path, format="turtle")
    constraints = {
        subject for subject in batch.subjects(CV.governedUniqueId, Literal(source["unique_id"]))
        if (subject, RDF.type, CV.NormalizedConstraint) in batch
        and (subject, PROV.wasDerivedFrom, URIRef(source["source_constraint_uri"])) in batch
    }
    if len(constraints) != 1:
        raise ValueError("Resume must identify exactly one previously normalized constraint.")
    constraint = next(iter(constraints))
    if not isinstance(constraint, URIRef):
        raise ValueError("The normalized constraint must retain its governed IRI.")
    resolution = URIRef(review_resolution_id)
    exception = _reviewed_exception(batch, resolution, constraint)
    first_index = STAGES.index(exception.construction_stage)
    # Refuse stale/repeated resume requests rather than silently rerunning
    # already completed stages or choosing among incompatible histories.
    for result_id in batch.subjects(UADEX.forConstraint, constraint):
        record = read_construction_stage_result(exchange_graph=batch, result_id=str(result_id))
        if record.construction_stage in STAGES[first_index:] and record.construction_status == "GREEN":
            raise ValueError("Selected exception has already advanced; review the current production state.")
    output_root = Path(output_directory).resolve(strict=True)
    paths = []
    last_checkpoint = None
    verified = []
    for stage in STAGES[:first_index]:
        record = _verify_stage(batch, str(constraint), stage)
        path, checkpoint = _load_verified_checkpoint(batch, record, output_root)
        paths.append(path)
        verified.append(record)
        last_checkpoint = checkpoint
    if last_checkpoint is None:
        raise ValueError("Resume requires a preceding verified checkpoint.")
    current = _checkpoint_with_batch_blank_nodes(last_checkpoint, batch)
    for record in verified:
        saved = _verify_stage(current, str(constraint), record.construction_stage)
        if saved.result_id != record.result_id:
            raise ValueError("Checkpoint changed an earlier verified stage identity.")
    # Changed governed source knowledge invalidates automatic reuse. Review
    # resolution does not authorize retaining stale normalized meaning.
    for predicate, key in (
        (CV.governedUniqueId, "unique_id"),
        (CV.dataPointName, "data_point_name"),
        (CV.sourceRuleId, "source_rule_id"),
        (CV.requirementText, "requirement_text"),
        (CV.violationCondition, "violation_condition"),
        (CV.severity, "severity"),
        (CV.sourceLocator, "source_locator"),
    ):
        if _one(current, constraint, predicate) != Literal(source[key]):
            raise ValueError("Governed source meaning changed; earlier normalization cannot be reused.")
    contexts = sorted(
        (
            int(str(_one(current, node, CV.contextOrder))),
            str(_one(current, node, CV.contextValue)),
        )
        for node in current.objects(constraint, CV.hasSourceContext)
    )
    if tuple(value for _, value in contexts) != tuple(source["context_fields"]):
        raise ValueError("Governed source context changed; verified stages cannot be reused.")
    state = _ResumeState(
        constraint_id=str(constraint),
        current=current,
        batch=batch,
        first_stage_index=first_index,
        previous_stage_paths=tuple(paths),
        exception_id=exception.result_id,
        resolution_id=resolution,
    )
    try:
        return _construct_governed_source(
            source=source,
            logical_schema_graph=logical_schema_graph,
            projection_model_graph=projection_model_graph,
            established_behaviors_graph=established_behaviors_graph,
            implementation_patterns_graph=implementation_patterns_graph,
            output_directory=output_directory,
            batch_results_path=batch_path,
            resume_state=state,
        )
    except _RowProductionError as error:
        _record_row_exception(
            unique_id=source["unique_id"],
            source_uri=source["source_constraint_uri"],
            error=error,
            stage=error.stage,
            batch_path=batch_path,
        )
        raise


__all__ = [
    "ConstraintProductionResult",
    "ConstraintProductionException",
    "ConstraintBatchProductionResult",
    "process_governed_source_row",
    "process_all_governed_source_rows",
    "resume_governed_source_row",
]
