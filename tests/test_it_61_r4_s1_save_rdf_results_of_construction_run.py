"""IT-61R4S1: save actual stage RDF with source selection, scope and replay order.

Construct through the real UI for each supported row scope. Inspect the saved
manifest/batch index and every original completed-stage snapshot, then use a
fresh client to open the saved run and its actual stage-result links. Review
must display the same exact Turtle and equivalent RDF without modifying files.
No successful production outcome or artifact is mocked. Helpers are local.
"""
from hashlib import sha256
from html.parser import HTMLParser
from importlib import import_module
from pathlib import Path
from urllib.parse import parse_qs, urlsplit
from urllib.request import url2pathname
from uuid import uuid4

import pytest
from fastapi.testclient import TestClient
from openpyxl import Workbook
from openpyxl.worksheet.worksheet import Worksheet
from rdflib import Graph, Literal, Namespace, RDF, RDFS, URIRef
from rdflib.namespace import OWL, PROV, SH, XSD
from rdflib.compare import isomorphic
from rdflib.util import from_n3

from app.main import app
from app.models.schema_model import (
    ComplexTypeDefinition, ElementDeclaration, ModelGroup, ModelGroupKind,
    QName, SchemaModel,
)
from app.services.construction_exchange import UADEX, read_construction_stage_result
from app.services.governed_constraint_source_reader import EXPECTED_COLUMNS
from app.services.logical_schema_serializer import logical_schema_model_to_graph
from app.services.rdf_projection import UAD_SCHEMA
CV = Namespace("https://dynamicontology.com/uad36/constraint-vocabulary#")
CFG = Namespace("https://dynamicontology.com/uad36/configuration#")
LSM = Namespace("https://dynamicontology.com/uad36/logical-schema#")
WORKERS = (
    ("constraint_normalization_agent", "normalize_governed_constraint_to_file",
     "normalized-constraint-representation", None),
    ("logical_schema_binding_agent", "bind_normalized_constraint_to_logical_schema",
     "logical-schema-binding", "normalized_graph"),
    ("context_resolution_agent", "resolve_bound_constraint_context",
     "context-resolution", "bound_graph"),
    ("constraint_behavior_classification", "classify_governed_constraint_behavior",
     "validation-behavior", "context_graph"),
    ("instance_rdf_binding_agent", "bind_governed_constraint_to_instance_rdf",
     "instance-rdf-binding", "classified_graph"),
    ("shacl_generation_agent", "generate_governed_constraint_shacl",
     "shacl-representation", "bound_graph"),
)


def _models(ns: Namespace) -> tuple[Graph, Graph]:
    xml_namespace = "urn:uad36:test:all-rows-pipeline:xml"

    def qname(name: str) -> QName:
        return QName(xml_namespace, name)

    value = ElementDeclaration(name=qname("AddressLineText"))
    address = ElementDeclaration(name=qname("ADDRESS"), type_name=qname("ADDRESS_TYPE"))
    root = ElementDeclaration(name=qname("PROPERTY"), type_name=qname("PROPERTY_TYPE"))
    address_type = ComplexTypeDefinition(
        name=qname("ADDRESS_TYPE"),
        content=ModelGroup(kind=ModelGroupKind.SEQUENCE, elements=(value,)),
    )
    property_type = ComplexTypeDefinition(
        name=qname("PROPERTY_TYPE"),
        content=ModelGroup(kind=ModelGroupKind.SEQUENCE, elements=(address,)),
    )
    schema = logical_schema_model_to_graph(SchemaModel(
        target_namespace=xml_namespace,
        elements={root.name: root},
        complex_types={item.name: item for item in (address_type, property_type)},
    ))
    projection = Graph()
    for triple in schema:
        projection.add(triple)
    for local_name, predicate, range_iri, property_type_iri in (
        ("PROPERTY", ns.property_predicate, ns.property_class, OWL.ObjectProperty),
        ("ADDRESS", ns.address_predicate, ns.address_class, OWL.ObjectProperty),
        ("AddressLineText", ns.value_predicate, XSD.string, OWL.DatatypeProperty),
    ):
        names = {
            subject for subject in schema.subjects(LSM.local_name, Literal(local_name))
            if (subject, RDF.type, LSM.QName) in schema
        }
        assert len(names) == 1
        resource = next(iter(names))
        assert isinstance(resource, URIRef)
        projection.add((resource, UAD_SCHEMA.sourceQName,
                        Literal(f"{{{xml_namespace}}}{local_name}")))
        projection.add((resource, UAD_SCHEMA.projectsTo, predicate))
        projection.add((predicate, RDF.type, property_type_iri))
        projection.add((predicate, RDFS.range, range_iri))
    projection.add((ns.property_class, RDF.type, OWL.Class))
    projection.add((ns.address_class, RDF.type, OWL.Class))
    return schema, projection


class _BuildPage(HTMLParser):
    def __init__(self) -> None:
        super().__init__(convert_charrefs=True)
        self.forms: list[dict] = []
        self.run_ids: list[str] = []
        self._form: dict | None = None
        self._button_parts: list[str] | None = None

    def handle_starttag(self, tag: str, attrs) -> None:
        attributes = dict(attrs)
        if tag == "form":
            self._form = {"attributes": attributes, "inputs": [], "buttons": []}
            self.forms.append(self._form)
        if tag == "input" and self._form is not None:
            self._form["inputs"].append(attributes)
        if tag == "button" and self._form is not None and attributes.get("type", "submit") == "submit":
            self._button_parts = []
        if attributes.get("data-build-run-id"):
            self.run_ids.append(attributes["data-build-run-id"])

    def handle_data(self, data: str) -> None:
        if self._button_parts is not None:
            self._button_parts.append(data)

    def handle_endtag(self, tag: str) -> None:
        if tag == "button" and self._button_parts is not None:
            assert self._form is not None
            self._form["buttons"].append(" ".join("".join(self._button_parts).split()))
            self._button_parts = None
        if tag == "form":
            self._form = None


def _file_artifact(graph: Graph, subject, predicate, output_root: Path) -> Path:
    values = set(graph.objects(subject, predicate))
    assert len(values) == 1
    uri = next(iter(values))
    assert isinstance(uri, URIRef)
    parsed = urlsplit(str(uri))
    assert parsed.scheme == "file" and parsed.netloc in ("", "localhost")
    path = Path(url2pathname(parsed.path)).resolve()
    assert path.is_file()
    assert path.is_relative_to(output_root.resolve())
    return path


class _SavedRunPage(_BuildPage):
    def __init__(self):
        super().__init__()
        self.links = []
        self.stage_links = []
        self.selected = []
        self.rdf_blocks = []
        self.edges = []
        self._rdf_parts = None
        self._svg_depth = 0

    def handle_starttag(self, tag, attrs):
        super().handle_starttag(tag, attrs)
        attributes = dict(attrs)
        if tag == "a" and attributes.get("href"):
            self.links.append(attributes["href"])
            if attributes.get("data-visualization-action") == "select":
                self.stage_links.append(attributes["href"])
        if "data-selected-stage-result-id" in attributes:
            self.selected.append(attributes["data-selected-stage-result-id"])
        if tag == "pre" and attributes.get("id") == "stage-rdf":
            self._rdf_parts = []
        if tag == "svg":
            self._svg_depth += 1
        keys = ("data-rdf-subject-n3", "data-rdf-predicate-n3", "data-rdf-object-n3")
        if self._svg_depth and all(key in attributes for key in keys):
            self.edges.append(tuple(attributes[key] for key in keys))

    def handle_data(self, data):
        super().handle_data(data)
        if self._rdf_parts is not None:
            self._rdf_parts.append(data)

    def handle_endtag(self, tag):
        super().handle_endtag(tag)
        if tag == "pre" and self._rdf_parts is not None:
            self.rdf_blocks.append(''.join(self._rdf_parts))
            self._rdf_parts = None
        if tag == "svg":
            self._svg_depth -= 1


def _files(directory):
    return {path.relative_to(directory): path.read_bytes()
            for path in directory.rglob("*") if path.is_file()}


@pytest.mark.parametrize("row_scope", ["one", "group", "all"], ids=["one-row", "selected-group", "all-rows"])
def test_it_61_r4_s1_saves_every_completed_stage_and_reopens_its_preserved_rdf(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch, row_scope: str,
) -> None:
    selected_index = 0
    ns = Namespace(f"urn:uad36:test:it61r4s1:{uuid4().hex}:")
    requirement = "Provide the address line for each property."
    violation = "The property address line is absent."
    source_directory = tmp_path / "sources"
    source_directory.mkdir()
    registrations = []
    for number, version in enumerate(("1.5", "1.6")):
        digits = f"{uuid4().int % (10 ** 8):08d}"
        # These two textual IDs become the same number if coerced numerically.
        ids = ("0" + digits, digits.lstrip("0") or "0")
        workbook_path = source_directory / f"registered-source-{number}.xlsx"
        workbook = Workbook()
        wrong_sheet = workbook.active
        assert isinstance(wrong_sheet, Worksheet)
        wrong_sheet.title = "DoNotUse"
        wrong_sheet.append(["The configured worksheet is GovernedRules."])
        sheet = workbook.create_sheet("GovernedRules")
        sheet.append(list(EXPECTED_COLUMNS))
        for unique_id in ids:
            row = dict.fromkeys(EXPECTED_COLUMNS, "")
            row.update({
                "Unique ID": unique_id,
                "Primary Data Element": "AddressLineText",
                "Message ID": "test-rule-" + uuid4().hex,
                "Message Text": requirement,
                "Rule Logic": violation,
                "Severity": "Fatal",
                "Property Affected": "/PROPERTY/ADDRESS/AddressLineText",
                "Data Point Name / Value": "AddressLineText",
                " xPath": "/PROPERTY/ADDRESS/AddressLineText",
            })
            sheet.append([row[column] for column in EXPECTED_COLUMNS])
        for cells in sheet.iter_rows():
            for cell in cells:
                cell.number_format = "@"
        workbook.save(workbook_path)
        workbook.close()
        registrations.append({
            "id": URIRef(f"{ns}set-{number}"),
            "workbook": URIRef(f"{ns}workbook-{number}"),
            "path": workbook_path,
            "version": version,
            "row_ids": ids,
        })
    selected = registrations[selected_index]
    other = registrations[1 - selected_index]
    schema, projection = _models(ns)
    behaviors = Graph()
    for triple in (
        (ns.behavior, RDF.type, CV.ValidationBehavior),
        (ns.behavior, CV.requirementText, Literal(requirement)),
        (ns.behavior, CV.violationCondition, Literal(violation)),
    ):
        behaviors.add(triple)
    patterns = Graph()
    for triple in (
        (ns.behavior, CV.shaclImplementationPattern, ns.pattern),
        (ns.pattern, RDF.type, CV.ShaclImplementationPattern),
        (ns.pattern, CV.propertyShapeTemplate, ns.template),
        (ns.template, RDF.type, SH.PropertyShape),
        (ns.template, SH.minCount, Literal(1)),
    ):
        patterns.add(triple)
    model_directory = tmp_path / "models"
    model_directory.mkdir()
    model_files = {}
    for name, graph in (
        ("logical-schema", schema), ("projection-model", projection),
        ("established-behaviors", behaviors), ("shacl-patterns", patterns),
    ):
        path = model_directory / f"{name}.ttl"
        graph.serialize(destination=path, format="turtle")
        model_files[name] = path
    output_root = tmp_path / "build-runs"
    configuration = Graph()
    for triple in (
        (ns.configuration, RDF.type, CFG.ApplicationConfiguration),
        (ns.configuration, CFG.hasSourceWorkbooks, ns.workbooks),
        (ns.configuration, CFG.hasConstraintSets, ns.constraint_sets),
        (ns.configuration, CFG.activeComplianceSource, other["workbook"]),
        (ns.workbooks, RDF.type, CFG.GovernedSourceWorkbookCollection),
        (ns.constraint_sets, RDF.type, CFG.GovernedConstraintSetCollection),
        (ns.authority, RDF.type, CFG.GoverningAuthority),
        (ns.authority, RDFS.label, Literal("Registered governing authority")),
        # Proposed build-profile contract: dependencies and destinations are
        # explicit governed configuration, never guessed from directories.
        (ns.configuration, CFG.usesConstraintBuildProfile, ns.build_profile),
        (ns.build_profile, RDF.type, CFG.ConstraintBuildProfile),
        (ns.build_profile, CFG.logicalSchemaFileLocation, Literal(str(model_files["logical-schema"]))),
        (ns.build_profile, CFG.projectionModelFileLocation, Literal(str(model_files["projection-model"]))),
        (ns.build_profile, CFG.establishedBehaviorsFileLocation, Literal(str(model_files["established-behaviors"]))),
        (ns.build_profile, CFG.shaclPatternsFileLocation, Literal(str(model_files["shacl-patterns"]))),
        (ns.build_profile, CFG.buildOutputDirectory, Literal(str(output_root))),
    ):
        configuration.add(triple)
    for registration in registrations:
        for triple in (
            (ns.constraint_sets, CFG.containsConstraintSet, registration["id"]),
            (registration["id"], RDF.type, CFG.GovernedConstraintSet),
            (registration["id"], RDFS.label, Literal(f"Governed release {registration['version']}")),
            (registration["id"], CFG.governingAuthority, ns.authority),
            (registration["id"], CFG.sourceWorkbook, registration["workbook"]),
            (registration["id"], CFG.governedVersion, Literal(registration["version"])),
            (ns.workbooks, CFG.containsWorkbook, registration["workbook"]),
            (registration["workbook"], RDF.type, CFG.GovernedSourceArtifact),
            (registration["workbook"], CFG.sourceFileLocation, Literal(str(registration["path"]))),
            (registration["workbook"], CFG.ruleDataSheetName, Literal("GovernedRules")),
            (registration["workbook"], CFG.sourceFormat, Literal("xlsx")),
        ):
            configuration.add(triple)
    # Workbook-scope construction must not require an IT-59 constraint-set registry.
    configuration.remove((ns.configuration, CFG.hasConstraintSets, None))
    configuration.remove((ns.constraint_sets, None, None))
    for registration in registrations:
        configuration.remove((registration["id"], None, None))
        configuration.add((registration["workbook"], RDFS.label,
                           Literal(f"Governed source workbook {registration['version']}")))
    configuration_path = tmp_path / "configuration.ttl"
    configuration.serialize(destination=configuration_path, format="turtle")
    immutable_files = [configuration_path, *model_files.values(), *[item["path"] for item in registrations]]
    originals = {path: path.read_bytes() for path in immutable_files}
    monkeypatch.setattr(app.state, "configuration_file", configuration_path)
    requested_ids = tuple(selected["row_ids"][:1] if row_scope == "one" else selected["row_ids"])
    data = {"source_workbook_id": str(selected["workbook"]), "row_scope": row_scope}
    if row_scope == "one":
        data["unique_id"] = requested_ids[0]
    elif row_scope == "group":
        data["unique_ids"] = list(requested_ids)
    with TestClient(app, follow_redirects=False) as client:
        response = client.post("/constraint-pipeline/runs", data=data)
    assert response.status_code == 200
    completion = _SavedRunPage()
    completion.feed(response.text)
    assert len(completion.run_ids) == 1
    run = URIRef(completion.run_ids[0])
    manifests = list(output_root.rglob("run.ttl"))
    assert len(manifests) == 1
    directory = manifests[0].parent.resolve()
    manifest = Graph().parse(manifests[0], format="turtle")
    for triple in (
        (run, RDF.type, CV.ConstraintPipelineRun),
        (run, CV.status, Literal("COMPLETE")),
        (run, CV.sourceWorkbook, selected["workbook"]),
        (run, CV.sourceWorksheet, Literal("GovernedRules")),
        (run, CV.rowScope, Literal(row_scope)),
        (run, CV.sourceWorkbookArtifact, URIRef(selected["path"].resolve().as_uri())),
    ):
        assert triple in manifest
    assert set(manifest.objects(run, CV.governedUniqueId)) == {Literal(value) for value in requested_ids}
    configuration_copy = _file_artifact(manifest, run, CV.configurationArtifact, directory)
    saved_configuration = Graph().parse(configuration_copy, format="turtle")
    assert set(saved_configuration.objects(ns.configuration, CFG.activeComplianceSource)) == {selected["workbook"]}
    assert (selected["workbook"], CFG.ruleDataSheetName, Literal("GovernedRules")) in saved_configuration
    batch_path = _file_artifact(manifest, run, CV.batchResultsArtifact, directory)
    batch = Graph().parse(batch_path, format="turtle")
    constraints = set(batch.subjects(RDF.type, CV.NormalizedConstraint))
    assert len(constraints) == len(requested_ids)
    expected_results = {}
    stages = [worker[2] for worker in WORKERS]
    for constraint in constraints:
        records = [read_construction_stage_result(exchange_graph=batch, result_id=str(node))
                   for node in batch.subjects(UADEX.forConstraint, constraint)]
        assert len(records) == len(stages)
        assert {record.construction_stage for record in records} == set(stages)
        for record in records:
            assert record.construction_status == "GREEN" and record.available_for_downstream
            result = record.result_id
            assert (run, CV.preservedStageResult, result) in manifest
            ordinal = stages.index(record.construction_stage) + 1
            assert set(manifest.objects(result, CV.stageOrder)) == {Literal(ordinal)}
            locations = set(batch.objects(result, UADEX.rdfSnapshotPath))
            assert len(locations) == 1
            assert set(manifest.objects(result, UADEX.rdfSnapshotPath)) == locations
            path = Path(str(next(iter(locations)))).resolve()
            assert path.is_file() and path.is_relative_to(directory)
            saved_bytes = path.read_bytes()
            digest = Literal(sha256(saved_bytes).hexdigest())
            assert set(manifest.objects(result, UADEX.snapshotSha256)) == {digest}
            assert set(batch.objects(result, UADEX.snapshotSha256)) == {digest}
            graph = Graph().parse(data=saved_bytes, format="turtle")
            assert (result, UADEX.forConstraint, constraint) in graph
            assert (result, UADEX.constructionStage, Literal(record.construction_stage)) in graph
            assert (constraint, RDF.type, CV.NormalizedConstraint) in graph
            expected_results[str(result)] = (ordinal, path, saved_bytes, graph)
    assert set(manifest.objects(run, CV.preservedStageResult)) == {URIRef(value) for value in expected_results}
    review_url = f"/constraint-workbench/runs/{directory.name}"
    assert any(urlsplit(href).path == review_url for href in completion.links), (
        "The completed UI run must provide access to its saved stage results for later review."
    )
    originals_after_save = _files(directory)
    with TestClient(app, follow_redirects=False) as later_client:
        response = later_client.get(review_url, params={"playback": "paused"})
        assert response.status_code == 200
        review = _SavedRunPage()
        review.feed(response.text)
        assert review.run_ids == [str(run)]
        listed_results = []
        for href in review.stage_links:
            parsed = urlsplit(href)
            assert not parsed.scheme and not parsed.netloc and parsed.path == review_url
            values = parse_qs(parsed.query).get("result_id", [])
            assert len(values) == 1
            listed_results.append(values[0])
        assert set(listed_results) == set(expected_results)
        assert len(listed_results) == len(expected_results)
        ordinals = [expected_results[value][0] for value in listed_results]
        assert ordinals == sorted(ordinals), "Saved results must remain accessible in governed stage order."
        for href, result_id in zip(review.stage_links, listed_results):
            response = later_client.get(href)
            assert response.status_code == 200
            displayed = _SavedRunPage()
            displayed.feed(response.text)
            _, path, saved_bytes, graph = expected_results[result_id]
            assert displayed.run_ids == [str(run)]
            assert displayed.selected == [result_id]
            assert displayed.rdf_blocks == [saved_bytes.decode("utf-8")]
            rendered = Graph()
            assert displayed.edges
            for edge in displayed.edges:
                rendered.add(tuple(from_n3(term) for term in edge))
            assert isomorphic(rendered, graph)
            assert path.read_bytes() == saved_bytes
    assert _files(directory) == originals_after_save, "Review changed the saved construction run."
    assert {path: path.read_bytes() for path in immutable_files} == originals
