"""IT-59R4S1: save a complete, ordered RDF replay index for a real build.

Proposed run-bundle index: run CV.preservedStageResult result; each result
records CV.stageOrder (1..6 per constraint), its original UADEX stage identity,
UADEX.rdfSnapshotPath, and UADEX.snapshotSha256. The original snapshot and
execution provenance remain preserved in the same run bundle's batch graph.
The index must refer to actual worker results, not newly minted substitutes.

All fixtures and helpers are local; no other test module is imported.
"""

from hashlib import sha256
from html.parser import HTMLParser
from pathlib import Path
from urllib.parse import urlsplit
from urllib.request import url2pathname
from uuid import uuid4

import pytest
from fastapi.testclient import TestClient
from openpyxl import Workbook
from openpyxl.worksheet.worksheet import Worksheet
from rdflib import Graph, Literal, Namespace, RDF, RDFS, URIRef
from rdflib.namespace import OWL, PROV, SH, XSD

from app.main import app
from app.models.schema_model import (
    ComplexTypeDefinition, ElementDeclaration, ModelGroup, ModelGroupKind,
    QName, SchemaModel,
)
from app.services.construction_exchange import UADEX, read_construction_stage_result
from app.services.governed_constraint_source_reader import EXPECTED_COLUMNS
from app.services.logical_schema_serializer import logical_schema_model_to_graph
from app.services.rdf_projection import UAD_SCHEMA

CV = Namespace("https://dynamicontology.com/uad36/constraint-vocabulary#")
CFG = Namespace("https://dynamicontology.com/uad36/configuration#")
LSM = Namespace("https://dynamicontology.com/uad36/logical-schema#")
WORKERS = (
    ("constraint_normalization_agent", "normalize_governed_constraint_to_file",
     "normalized-constraint-representation", None),
    ("logical_schema_binding_agent", "bind_normalized_constraint_to_logical_schema",
     "logical-schema-binding", "normalized_graph"),
    ("context_resolution_agent", "resolve_bound_constraint_context",
     "context-resolution", "bound_graph"),
    ("constraint_behavior_classification", "classify_governed_constraint_behavior",
     "validation-behavior", "context_graph"),
    ("instance_rdf_binding_agent", "bind_governed_constraint_to_instance_rdf",
     "instance-rdf-binding", "classified_graph"),
    ("shacl_generation_agent", "generate_governed_constraint_shacl",
     "shacl-representation", "bound_graph"),
)


def _models(ns: Namespace) -> tuple[Graph, Graph]:
    xml_namespace = "urn:uad36:test:all-rows-pipeline:xml"

    def qname(name: str) -> QName:
        return QName(xml_namespace, name)

    value = ElementDeclaration(name=qname("AddressLineText"))
    address = ElementDeclaration(name=qname("ADDRESS"), type_name=qname("ADDRESS_TYPE"))
    root = ElementDeclaration(name=qname("PROPERTY"), type_name=qname("PROPERTY_TYPE"))
    address_type = ComplexTypeDefinition(
        name=qname("ADDRESS_TYPE"),
        content=ModelGroup(kind=ModelGroupKind.SEQUENCE, elements=(value,)),
    )
    property_type = ComplexTypeDefinition(
        name=qname("PROPERTY_TYPE"),
        content=ModelGroup(kind=ModelGroupKind.SEQUENCE, elements=(address,)),
    )
    schema = logical_schema_model_to_graph(SchemaModel(
        target_namespace=xml_namespace,
        elements={root.name: root},
        complex_types={item.name: item for item in (address_type, property_type)},
    ))
    projection = Graph()
    for triple in schema:
        projection.add(triple)
    for local_name, predicate, range_iri, property_type_iri in (
        ("PROPERTY", ns.property_predicate, ns.property_class, OWL.ObjectProperty),
        ("ADDRESS", ns.address_predicate, ns.address_class, OWL.ObjectProperty),
        ("AddressLineText", ns.value_predicate, XSD.string, OWL.DatatypeProperty),
    ):
        names = {
            subject for subject in schema.subjects(LSM.local_name, Literal(local_name))
            if (subject, RDF.type, LSM.QName) in schema
        }
        assert len(names) == 1
        resource = next(iter(names))
        assert isinstance(resource, URIRef)
        projection.add((resource, UAD_SCHEMA.sourceQName,
                        Literal(f"{{{xml_namespace}}}{local_name}")))
        projection.add((resource, UAD_SCHEMA.projectsTo, predicate))
        projection.add((predicate, RDF.type, property_type_iri))
        projection.add((predicate, RDFS.range, range_iri))
    projection.add((ns.property_class, RDF.type, OWL.Class))
    projection.add((ns.address_class, RDF.type, OWL.Class))
    return schema, projection


class _BuildPage(HTMLParser):
    def __init__(self) -> None:
        super().__init__(convert_charrefs=True)
        self.forms: list[dict] = []
        self.run_ids: list[str] = []
        self._form: dict | None = None
        self._button_parts: list[str] | None = None

    def handle_starttag(self, tag: str, attrs) -> None:
        attributes = dict(attrs)
        if tag == "form":
            self._form = {"attributes": attributes, "inputs": [], "buttons": []}
            self.forms.append(self._form)
        if tag == "input" and self._form is not None:
            self._form["inputs"].append(attributes)
        if tag == "button" and self._form is not None and attributes.get("type", "submit") == "submit":
            self._button_parts = []
        if attributes.get("data-build-run-id"):
            self.run_ids.append(attributes["data-build-run-id"])

    def handle_data(self, data: str) -> None:
        if self._button_parts is not None:
            self._button_parts.append(data)

    def handle_endtag(self, tag: str) -> None:
        if tag == "button" and self._button_parts is not None:
            assert self._form is not None
            self._form["buttons"].append(" ".join("".join(self._button_parts).split()))
            self._button_parts = None
        if tag == "form":
            self._form = None


def _file_artifact(graph: Graph, subject, predicate, output_root: Path) -> Path:
    values = set(graph.objects(subject, predicate))
    assert len(values) == 1
    uri = next(iter(values))
    assert isinstance(uri, URIRef)
    parsed = urlsplit(str(uri))
    assert parsed.scheme == "file" and parsed.netloc in ("", "localhost")
    path = Path(url2pathname(parsed.path)).resolve()
    assert path.is_file()
    assert path.is_relative_to(output_root.resolve())
    return path


def test_it_59_r4_s1_preserves_every_completed_stage_as_ordered_replay_knowledge(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch,
) -> None:
    selected_index = 0
    ns = Namespace(f"urn:uad36:test:it59r4s1:{uuid4().hex}:")
    requirement = "Provide the address line for each property."
    violation = "The property address line is absent."
    source_directory = tmp_path / "sources"
    source_directory.mkdir()
    registrations = []
    for number, version in enumerate(("1.5", "1.6")):
        ids = tuple("0" + uuid4().hex for _ in range(2))
        workbook_path = source_directory / f"registered-source-{number}.xlsx"
        workbook = Workbook()
        wrong_sheet = workbook.active
        assert isinstance(wrong_sheet, Worksheet)
        wrong_sheet.title = "DoNotUse"
        wrong_sheet.append(["The configured worksheet is GovernedRules."])
        sheet = workbook.create_sheet("GovernedRules")
        sheet.append(list(EXPECTED_COLUMNS))
        for unique_id in ids:
            row = dict.fromkeys(EXPECTED_COLUMNS, "")
            row.update({
                "Unique ID": unique_id,
                "Primary Data Element": "AddressLineText",
                "Message ID": "test-rule-" + uuid4().hex,
                "Message Text": requirement,
                "Rule Logic": violation,
                "Severity": "Fatal",
                "Property Affected": "/PROPERTY/ADDRESS/AddressLineText",
                "Data Point Name / Value": "AddressLineText",
                " xPath": "/PROPERTY/ADDRESS/AddressLineText",
            })
            sheet.append([row[column] for column in EXPECTED_COLUMNS])
        for cells in sheet.iter_rows():
            for cell in cells:
                cell.number_format = "@"
        workbook.save(workbook_path)
        workbook.close()
        registrations.append({
            "id": URIRef(f"{ns}set-{number}"),
            "workbook": URIRef(f"{ns}workbook-{number}"),
            "path": workbook_path,
            "version": version,
            "row_ids": ids,
        })
    selected = registrations[selected_index]
    other = registrations[1 - selected_index]
    schema, projection = _models(ns)
    behaviors = Graph()
    for triple in (
        (ns.behavior, RDF.type, CV.ValidationBehavior),
        (ns.behavior, CV.requirementText, Literal(requirement)),
        (ns.behavior, CV.violationCondition, Literal(violation)),
    ):
        behaviors.add(triple)
    patterns = Graph()
    for triple in (
        (ns.behavior, CV.shaclImplementationPattern, ns.pattern),
        (ns.pattern, RDF.type, CV.ShaclImplementationPattern),
        (ns.pattern, CV.propertyShapeTemplate, ns.template),
        (ns.template, RDF.type, SH.PropertyShape),
        (ns.template, SH.minCount, Literal(1)),
    ):
        patterns.add(triple)
    model_directory = tmp_path / "models"
    model_directory.mkdir()
    model_files = {}
    for name, graph in (
        ("logical-schema", schema), ("projection-model", projection),
        ("established-behaviors", behaviors), ("shacl-patterns", patterns),
    ):
        path = model_directory / f"{name}.ttl"
        graph.serialize(destination=path, format="turtle")
        model_files[name] = path
    output_root = tmp_path / "build-runs"
    configuration = Graph()
    for triple in (
        (ns.configuration, RDF.type, CFG.ApplicationConfiguration),
        (ns.configuration, CFG.hasSourceWorkbooks, ns.workbooks),
        (ns.configuration, CFG.hasConstraintSets, ns.constraint_sets),
        (ns.configuration, CFG.activeComplianceSource, other["workbook"]),
        (ns.workbooks, RDF.type, CFG.GovernedSourceWorkbookCollection),
        (ns.constraint_sets, RDF.type, CFG.GovernedConstraintSetCollection),
        (ns.authority, RDF.type, CFG.GoverningAuthority),
        (ns.authority, RDFS.label, Literal("Registered governing authority")),
        # Proposed build-profile contract: dependencies and destinations are
        # explicit governed configuration, never guessed from directories.
        (ns.configuration, CFG.usesConstraintBuildProfile, ns.build_profile),
        (ns.build_profile, RDF.type, CFG.ConstraintBuildProfile),
        (ns.build_profile, CFG.logicalSchemaFileLocation, Literal(str(model_files["logical-schema"]))),
        (ns.build_profile, CFG.projectionModelFileLocation, Literal(str(model_files["projection-model"]))),
        (ns.build_profile, CFG.establishedBehaviorsFileLocation, Literal(str(model_files["established-behaviors"]))),
        (ns.build_profile, CFG.shaclPatternsFileLocation, Literal(str(model_files["shacl-patterns"]))),
        (ns.build_profile, CFG.buildOutputDirectory, Literal(str(output_root))),
    ):
        configuration.add(triple)
    for registration in registrations:
        for triple in (
            (ns.constraint_sets, CFG.containsConstraintSet, registration["id"]),
            (registration["id"], RDF.type, CFG.GovernedConstraintSet),
            (registration["id"], RDFS.label, Literal(f"Governed release {registration['version']}")),
            (registration["id"], CFG.governingAuthority, ns.authority),
            (registration["id"], CFG.sourceWorkbook, registration["workbook"]),
            (registration["id"], CFG.governedVersion, Literal(registration["version"])),
            (ns.workbooks, CFG.containsWorkbook, registration["workbook"]),
            (registration["workbook"], RDF.type, CFG.GovernedSourceArtifact),
            (registration["workbook"], CFG.sourceFileLocation, Literal(str(registration["path"]))),
            (registration["workbook"], CFG.ruleDataSheetName, Literal("GovernedRules")),
            (registration["workbook"], CFG.sourceFormat, Literal("xlsx")),
        ):
            configuration.add(triple)
    configuration_path = tmp_path / "configuration.ttl"
    configuration.serialize(destination=configuration_path, format="turtle")
    immutable_files = [configuration_path, *model_files.values(), *[item["path"] for item in registrations]]
    originals = {path: path.read_bytes() for path in immutable_files}
    monkeypatch.setattr(app.state, "configuration_file", configuration_path)
    with TestClient(app, follow_redirects=False) as client:
        response = client.post("/constraint-workbench/build", data={
            "authority": str(ns.authority), "constraint_set_id": str(selected["id"]),
        })
    assert response.status_code == 200
    page = _BuildPage()
    page.feed(response.text)
    assert len(page.run_ids) == 1
    run = URIRef(page.run_ids[0])
    manifests = list(output_root.rglob("run.ttl"))
    assert len(manifests) == 1
    manifest_path = manifests[0]
    directory = manifest_path.parent.resolve()
    manifest = Graph().parse(manifest_path, format="turtle")
    assert (run, RDF.type, CV.ConstraintSetBuildRun) in manifest
    assert (run, CV.status, Literal("COMPLETE")) in manifest
    assert set(manifest.objects(run, CV.selectedConstraintSet)) == {selected["id"]}
    assert set(manifest.objects(run, CV.sourceWorkbook)) == {selected["workbook"]}
    assert set(manifest.objects(run, CV.sourceVersion)) == {Literal(selected["version"])}
    assert (run, PROV.used, URIRef(selected["path"].resolve().as_uri())) in manifest
    batch_path = _file_artifact(manifest, run, CV.batchResultsArtifact, directory)
    batch = Graph().parse(batch_path, format="turtle")
    constraints = set(batch.subjects(RDF.type, CV.NormalizedConstraint))
    assert len(constraints) == len(selected["row_ids"])
    assert {str(value) for constraint in constraints
            for value in batch.objects(constraint, CV.governedUniqueId)} == set(selected["row_ids"])
    expected_results = {
        result for constraint in constraints
        for result in batch.subjects(UADEX.forConstraint, constraint)
        if (result, RDF.type, UADEX.ConstructionStageResult) in batch
        and (result, UADEX.constructionStatus, Literal("GREEN")) in batch
    }
    assert len(expected_results) == len(WORKERS) * len(constraints)
    indexed_results = set(manifest.objects(run, CV.preservedStageResult))
    assert indexed_results == expected_results, (
        "IT-59R4S1 requires the saved run manifest to index every completed worker's "
        "original RDF result for replay, without omissions or replacement identities."
    )
    saved_paths = set()
    for constraint in constraints:
        records = {result for result in expected_results
                   if (result, UADEX.forConstraint, constraint) in batch}
        ordinals = set()
        for result in records:
            record = read_construction_stage_result(exchange_graph=batch, result_id=str(result))
            stage_names = [worker[2] for worker in WORKERS]
            assert record.construction_stage in stage_names
            ordinal = stage_names.index(record.construction_stage) + 1
            assert set(manifest.objects(result, CV.stageOrder)) == {Literal(ordinal)}, (
                "Replay order must be explicit saved knowledge, not inferred from filenames."
            )
            ordinals.add(ordinal)
            for predicate, expected in (
                (RDF.type, UADEX.ConstructionStageResult),
                (UADEX.forConstraint, constraint),
                (UADEX.constructionStage, Literal(record.construction_stage)),
                (UADEX.constructionStatus, Literal("GREEN")),
            ):
                assert (result, predicate, expected) in manifest
            locations = set(batch.objects(result, UADEX.rdfSnapshotPath))
            assert len(locations) == 1
            assert set(manifest.objects(result, UADEX.rdfSnapshotPath)) == locations
            snapshot_path = Path(str(next(iter(locations)))).resolve()
            assert snapshot_path.is_file() and snapshot_path.is_relative_to(directory)
            assert snapshot_path not in saved_paths, "Each stage must preserve its own serialized RDF."
            saved_paths.add(snapshot_path)
            saved_bytes = snapshot_path.read_bytes()
            digest = Literal(sha256(saved_bytes).hexdigest())
            assert set(manifest.objects(result, UADEX.snapshotSha256)) == {digest}
            assert set(batch.objects(result, UADEX.snapshotSha256)) == {digest}
            snapshot = Graph().parse(data=saved_bytes, format="turtle")
            assert (result, RDF.type, UADEX.ConstructionStageResult) in snapshot
            assert (result, UADEX.forConstraint, constraint) in snapshot
            assert (constraint, RDF.type, CV.NormalizedConstraint) in snapshot
            assert list(snapshot.objects(constraint, PROV.wasDerivedFrom))
            # Preserve all recorded execution-provenance edges, including the
            # generating activity and its agent/input/timing knowledge.
            activities = set(batch.objects(result, PROV.wasGeneratedBy))
            assert activities, "Every saved worker result needs execution provenance."
            assert set(snapshot.objects(result, PROV.wasGeneratedBy)) == activities
            for activity in activities:
                assert (activity, RDF.type, PROV.Activity) in snapshot
                for predicate, obj in batch.predicate_objects(activity):
                    if str(predicate).startswith(str(PROV)):
                        assert (activity, predicate, obj) in snapshot
        assert ordinals == set(range(1, len(WORKERS) + 1))
    assert len(saved_paths) == len(expected_results)
    assert {path: path.read_bytes() for path in immutable_files} == originals
