"""IT-59R2S2: repeated builds preserve independent, persistent run histories."""

from pathlib import Path
from uuid import uuid4

import pytest
from fastapi.testclient import TestClient
from openpyxl import Workbook
from openpyxl.worksheet.worksheet import Worksheet
from rdflib import Graph, Literal, Namespace, RDF, RDFS, URIRef
from rdflib.namespace import OWL, PROV, SH, XSD

from app.main import app
from app.services.construction_exchange import UADEX, read_construction_stage_result
from app.services.governed_constraint_source_reader import EXPECTED_COLUMNS
from html.parser import HTMLParser
from urllib.parse import urlsplit
from urllib.request import url2pathname

from app.models.schema_model import (
    ComplexTypeDefinition, ElementDeclaration, ModelGroup, ModelGroupKind,
    QName, SchemaModel,
)
from app.services.logical_schema_serializer import logical_schema_model_to_graph
from app.services.rdf_projection import UAD_SCHEMA


CV = Namespace("https://dynamicontology.com/uad36/constraint-vocabulary#")
CFG = Namespace("https://dynamicontology.com/uad36/configuration#")
LSM = Namespace("https://dynamicontology.com/uad36/logical-schema#")
WORKERS = (
    ("constraint_normalization_agent", "normalize_governed_constraint_to_file",
     "normalized-constraint-representation", None),
    ("logical_schema_binding_agent", "bind_normalized_constraint_to_logical_schema",
     "logical-schema-binding", "normalized_graph"),
    ("context_resolution_agent", "resolve_bound_constraint_context",
     "context-resolution", "bound_graph"),
    ("constraint_behavior_classification", "classify_governed_constraint_behavior",
     "validation-behavior", "context_graph"),
    ("instance_rdf_binding_agent", "bind_governed_constraint_to_instance_rdf",
     "instance-rdf-binding", "classified_graph"),
    ("shacl_generation_agent", "generate_governed_constraint_shacl",
     "shacl-representation", "bound_graph"),
)


def _models(ns: Namespace) -> tuple[Graph, Graph]:
    xml_namespace = "urn:uad36:test:all-rows-pipeline:xml"

    def qname(name: str) -> QName:
        return QName(xml_namespace, name)

    value = ElementDeclaration(name=qname("AddressLineText"))
    address = ElementDeclaration(name=qname("ADDRESS"), type_name=qname("ADDRESS_TYPE"))
    root = ElementDeclaration(name=qname("PROPERTY"), type_name=qname("PROPERTY_TYPE"))
    address_type = ComplexTypeDefinition(
        name=qname("ADDRESS_TYPE"),
        content=ModelGroup(kind=ModelGroupKind.SEQUENCE, elements=(value,)),
    )
    property_type = ComplexTypeDefinition(
        name=qname("PROPERTY_TYPE"),
        content=ModelGroup(kind=ModelGroupKind.SEQUENCE, elements=(address,)),
    )
    schema = logical_schema_model_to_graph(SchemaModel(
        target_namespace=xml_namespace,
        elements={root.name: root},
        complex_types={item.name: item for item in (address_type, property_type)},
    ))
    projection = Graph()
    for triple in schema:
        projection.add(triple)
    for local_name, predicate, range_iri, property_type_iri in (
        ("PROPERTY", ns.property_predicate, ns.property_class, OWL.ObjectProperty),
        ("ADDRESS", ns.address_predicate, ns.address_class, OWL.ObjectProperty),
        ("AddressLineText", ns.value_predicate, XSD.string, OWL.DatatypeProperty),
    ):
        names = {
            subject for subject in schema.subjects(LSM.local_name, Literal(local_name))
            if (subject, RDF.type, LSM.QName) in schema
        }
        assert len(names) == 1
        resource = next(iter(names))
        assert isinstance(resource, URIRef)
        projection.add((resource, UAD_SCHEMA.sourceQName,
                        Literal(f"{{{xml_namespace}}}{local_name}")))
        projection.add((resource, UAD_SCHEMA.projectsTo, predicate))
        projection.add((predicate, RDF.type, property_type_iri))
        projection.add((predicate, RDFS.range, range_iri))
    projection.add((ns.property_class, RDF.type, OWL.Class))
    projection.add((ns.address_class, RDF.type, OWL.Class))
    return schema, projection


class _BuildPage(HTMLParser):
    def __init__(self) -> None:
        super().__init__(convert_charrefs=True)
        self.forms: list[dict] = []
        self.run_ids: list[str] = []
        self._form: dict | None = None
        self._button_parts: list[str] | None = None

    def handle_starttag(self, tag: str, attrs) -> None:
        attributes = dict(attrs)
        if tag == "form":
            self._form = {"attributes": attributes, "inputs": [], "buttons": []}
            self.forms.append(self._form)
        if tag == "input" and self._form is not None:
            self._form["inputs"].append(attributes)
        if tag == "button" and self._form is not None and attributes.get("type", "submit") == "submit":
            self._button_parts = []
        if attributes.get("data-build-run-id"):
            self.run_ids.append(attributes["data-build-run-id"])

    def handle_data(self, data: str) -> None:
        if self._button_parts is not None:
            self._button_parts.append(data)

    def handle_endtag(self, tag: str) -> None:
        if tag == "button" and self._button_parts is not None:
            assert self._form is not None
            self._form["buttons"].append(" ".join("".join(self._button_parts).split()))
            self._button_parts = None
        if tag == "form":
            self._form = None


def _file_artifact(graph: Graph, subject, predicate, output_root: Path) -> Path:
    values = set(graph.objects(subject, predicate))
    assert len(values) == 1
    uri = next(iter(values))
    assert isinstance(uri, URIRef)
    parsed = urlsplit(str(uri))
    assert parsed.scheme == "file" and parsed.netloc in ("", "localhost")
    path = Path(url2pathname(parsed.path)).resolve()
    assert path.is_file()
    assert path.is_relative_to(output_root.resolve())
    return path


def _saved_run(response, output_root: Path):
    assert response.status_code == 200, (
        f"The build must finish successfully; got HTTP {response.status_code}."
    )
    page = _BuildPage()
    page.feed(response.text)
    assert len(page.run_ids) == 1, "Each build response must identify exactly one run."
    run = URIRef(page.run_ids[0])
    matches = []
    for path in output_root.rglob("run.ttl"):
        graph = Graph().parse(path, format="turtle")
        if (run, RDF.type, CV.ConstraintSetBuildRun) in graph:
            matches.append((path, graph))
    assert len(matches) == 1, "The returned run must have exactly one saved manifest."
    path, graph = matches[0]
    assert (run, CV.status, Literal("COMPLETE")) in graph
    batch_path = _file_artifact(graph, run, CV.batchResultsArtifact, output_root)
    assert batch_path.is_relative_to(path.parent.resolve())
    batch = Graph().parse(batch_path, format="turtle")
    return run, path.parent, graph, batch


def _verified_results(batch: Graph, expected_ids: set[str], run_directory: Path):
    constraints = set(batch.subjects(RDF.type, CV.NormalizedConstraint))
    assert {str(value) for constraint in constraints
            for value in batch.objects(constraint, CV.governedUniqueId)} == expected_ids
    assert len(constraints) == len(expected_ids)
    result_ids = set()
    artifact_paths = set()
    for constraint in constraints:
        nodes = set(batch.subjects(UADEX.forConstraint, constraint))
        records = [read_construction_stage_result(
            exchange_graph=batch, result_id=str(node),
        ) for node in nodes]
        assert len(records) == len(WORKERS)
        assert {record.construction_stage for record in records} == {item[2] for item in WORKERS}
        assert all(record.construction_status == "GREEN" and record.available_for_downstream
                   for record in records)
        result_ids.update(nodes)
        states = {node for node in batch.subjects(RDF.type, CV.ConstraintProductionState)
                  if (node, CV.governedConstraint, constraint) in batch}
        assert len(states) == 1
        state = next(iter(states))
        assert (state, CV.status, Literal("COMPLETE")) in batch
        rdf_path = _file_artifact(batch, state, CV.rdfArtifact, run_directory)
        shapes_path = _file_artifact(batch, state, CV.shaclArtifact, run_directory)
        artifact_paths.update((rdf_path, shapes_path))
        assert (constraint, RDF.type, CV.NormalizedConstraint) in Graph().parse(
            rdf_path, format="turtle",
        )
        assert list(Graph().parse(shapes_path, format="turtle").subjects(
            CV.implementsConstraint, constraint,
        ))
    return constraints, result_ids, artifact_paths


def test_it_59_r2_s2_keeps_repeated_build_runs_and_previous_stage_results_distinct(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch,
) -> None:
    # Build temporary governed inputs and execute two real builds.
    selected_index = 0
    ns = Namespace(f"urn:uad36:test:it59r2s2:{uuid4().hex}:")
    requirement = "Provide the address line for each property."
    violation = "The property address line is absent."
    source_directory = tmp_path / "sources"
    source_directory.mkdir()
    registrations = []
    for number, version in enumerate(("1.5", "1.6")):
        ids = tuple("0" + uuid4().hex for _ in range(2))
        workbook_path = source_directory / f"registered-source-{number}.xlsx"
        workbook = Workbook()
        wrong_sheet = workbook.active
        assert isinstance(wrong_sheet, Worksheet)
        wrong_sheet.title = "DoNotUse"
        wrong_sheet.append(["The configured worksheet is GovernedRules."])
        sheet = workbook.create_sheet("GovernedRules")
        sheet.append(list(EXPECTED_COLUMNS))
        for unique_id in ids:
            row = dict.fromkeys(EXPECTED_COLUMNS, "")
            row.update({
                "Unique ID": unique_id,
                "Primary Data Element": "AddressLineText",
                "Message ID": "test-rule-" + uuid4().hex,
                "Message Text": requirement,
                "Rule Logic": violation,
                "Severity": "Fatal",
                "Property Affected": "/PROPERTY/ADDRESS/AddressLineText",
                "Data Point Name / Value": "AddressLineText",
                " xPath": "/PROPERTY/ADDRESS/AddressLineText",
            })
            sheet.append([row[column] for column in EXPECTED_COLUMNS])
        for cells in sheet.iter_rows():
            for cell in cells:
                cell.number_format = "@"
        workbook.save(workbook_path)
        workbook.close()
        registrations.append({
            "id": URIRef(f"{ns}set-{number}"),
            "workbook": URIRef(f"{ns}workbook-{number}"),
            "path": workbook_path,
            "version": version,
            "row_ids": ids,
        })
    selected = registrations[selected_index]
    other = registrations[1 - selected_index]
    schema, projection = _models(ns)
    behaviors = Graph()
    for triple in (
        (ns.behavior, RDF.type, CV.ValidationBehavior),
        (ns.behavior, CV.requirementText, Literal(requirement)),
        (ns.behavior, CV.violationCondition, Literal(violation)),
    ):
        behaviors.add(triple)
    patterns = Graph()
    for triple in (
        (ns.behavior, CV.shaclImplementationPattern, ns.pattern),
        (ns.pattern, RDF.type, CV.ShaclImplementationPattern),
        (ns.pattern, CV.propertyShapeTemplate, ns.template),
        (ns.template, RDF.type, SH.PropertyShape),
        (ns.template, SH.minCount, Literal(1)),
    ):
        patterns.add(triple)
    model_directory = tmp_path / "models"
    model_directory.mkdir()
    model_files = {}
    for name, graph in (
        ("logical-schema", schema), ("projection-model", projection),
        ("established-behaviors", behaviors), ("shacl-patterns", patterns),
    ):
        path = model_directory / f"{name}.ttl"
        graph.serialize(destination=path, format="turtle")
        model_files[name] = path
    output_root = tmp_path / "build-runs"
    configuration = Graph()
    for triple in (
        (ns.configuration, RDF.type, CFG.ApplicationConfiguration),
        (ns.configuration, CFG.hasSourceWorkbooks, ns.workbooks),
        (ns.configuration, CFG.hasConstraintSets, ns.constraint_sets),
        (ns.configuration, CFG.activeComplianceSource, other["workbook"]),
        (ns.workbooks, RDF.type, CFG.GovernedSourceWorkbookCollection),
        (ns.constraint_sets, RDF.type, CFG.GovernedConstraintSetCollection),
        (ns.authority, RDF.type, CFG.GoverningAuthority),
        (ns.authority, RDFS.label, Literal("Registered governing authority")),
        # Proposed build-profile contract: dependencies and destinations are
        # explicit governed configuration, never guessed from directories.
        (ns.configuration, CFG.usesConstraintBuildProfile, ns.build_profile),
        (ns.build_profile, RDF.type, CFG.ConstraintBuildProfile),
        (ns.build_profile, CFG.logicalSchemaFileLocation, Literal(str(model_files["logical-schema"]))),
        (ns.build_profile, CFG.projectionModelFileLocation, Literal(str(model_files["projection-model"]))),
        (ns.build_profile, CFG.establishedBehaviorsFileLocation, Literal(str(model_files["established-behaviors"]))),
        (ns.build_profile, CFG.shaclPatternsFileLocation, Literal(str(model_files["shacl-patterns"]))),
        (ns.build_profile, CFG.buildOutputDirectory, Literal(str(output_root))),
    ):
        configuration.add(triple)
    for registration in registrations:
        for triple in (
            (ns.constraint_sets, CFG.containsConstraintSet, registration["id"]),
            (registration["id"], RDF.type, CFG.GovernedConstraintSet),
            (registration["id"], RDFS.label, Literal(f"Governed release {registration['version']}")),
            (registration["id"], CFG.governingAuthority, ns.authority),
            (registration["id"], CFG.sourceWorkbook, registration["workbook"]),
            (registration["id"], CFG.governedVersion, Literal(registration["version"])),
            (ns.workbooks, CFG.containsWorkbook, registration["workbook"]),
            (registration["workbook"], RDF.type, CFG.GovernedSourceArtifact),
            (registration["workbook"], CFG.sourceFileLocation, Literal(str(registration["path"]))),
            (registration["workbook"], CFG.ruleDataSheetName, Literal("GovernedRules")),
            (registration["workbook"], CFG.sourceFormat, Literal("xlsx")),
        ):
            configuration.add(triple)
    configuration_path = tmp_path / "configuration.ttl"
    configuration.serialize(destination=configuration_path, format="turtle")
    immutable_files = [configuration_path, *model_files.values(), *[item["path"] for item in registrations]]
    originals = {path: path.read_bytes() for path in immutable_files}
    monkeypatch.setattr(app.state, "configuration_file", configuration_path)
    request_data = {
        "authority": str(ns.authority),
        "constraint_set_id": str(selected["id"]),
    }
    with TestClient(app, follow_redirects=False) as client:
        first_response = client.post("/constraint-workbench/build", data=request_data)
    first_run, first_directory, first_manifest, first_batch = _saved_run(first_response, output_root)
    first_constraints, first_results, first_artifacts = _verified_results(
        first_batch, set(selected["row_ids"]), first_directory,
    )
    # Snapshot every saved file, including intermediate stage RDF, checkpoints,
    # configuration, final RDF, SHACL, batch state and the run manifest.
    previous_files = {
        path.relative_to(first_directory): path.read_bytes()
        for path in first_directory.rglob("*") if path.is_file()
    }
    assert previous_files
    assert all(path.is_relative_to(first_directory.resolve()) for path in first_artifacts)

    # A fresh client must still preserve the first build's persistent history.
    with TestClient(app, follow_redirects=False) as client:
        second_response = client.post("/constraint-workbench/build", data=request_data)
    second_run, second_directory, second_manifest, second_batch = _saved_run(second_response, output_root)
    assert second_run != first_run, (
        "IT-59R2S2 requires a distinct run identity when the same set is built again."
    )
    assert second_directory != first_directory, "A rebuild must not reuse the previous run directory."
    assert not second_directory.is_relative_to(first_directory)
    assert not first_directory.is_relative_to(second_directory)
    assert len(list(output_root.rglob("run.ttl"))) == 2
    for run, manifest in ((first_run, first_manifest), (second_run, second_manifest)):
        assert set(manifest.objects(run, CV.selectedConstraintSet)) == {selected["id"]}
        assert set(manifest.objects(run, CV.sourceVersion)) == {Literal(selected["version"])}
        assert set(manifest.objects(run, CV.sourceWorkbook)) == {selected["workbook"]}
        assert (run, RDF.type, PROV.Activity) in manifest
    second_constraints, second_results, second_artifacts = _verified_results(
        second_batch, set(selected["row_ids"]), second_directory,
    )
    # Governed constraint identities persist; execution-result identities and
    # saved artifacts belong to distinct runs.
    assert second_constraints == first_constraints
    assert first_results.isdisjoint(second_results), "Stage execution identities were reused across runs."
    assert first_artifacts.isdisjoint(second_artifacts), "The two runs share writable artifact files."
    current_files = {
        path.relative_to(first_directory): path.read_bytes()
        for path in first_directory.rglob("*") if path.is_file()
    }
    assert current_files == previous_files, (
        "IT-59R2S2 requires every file in the previous run, including its stage results, "
        "to remain present and byte-for-byte unchanged after rebuilding the same set."
    )
    assert {path: path.read_bytes() for path in immutable_files} == originals
