"""IT-61R1S2: submit one exact textual source-row ID through the pipeline UI.

Proposed UI request: POST /constraint-pipeline/runs with source_workbook_id,
row_scope="one", and unique_id. The run page identifies data-build-run-id.
Proposed saved manifest: CV.ConstraintPipelineRun with CV.sourceWorkbook,
CV.sourceWorksheet, CV.governedUniqueId, CV.rowScope="one", CV.status,
and CV.batchResultsArtifact. No constraint-set registration is required.

A leading-zero ID and its numerically equivalent spelling coexist in the
selected worksheet. Only the exact requested string may pass through workers.
Fixtures/helpers are self-contained; workers execute their real implementation.
"""

from html.parser import HTMLParser
from importlib import import_module
from pathlib import Path
from urllib.parse import urlsplit
from urllib.request import url2pathname
from uuid import uuid4

import pytest
from fastapi.testclient import TestClient
from openpyxl import Workbook
from openpyxl.worksheet.worksheet import Worksheet
from rdflib import Graph, Literal, Namespace, RDF, RDFS, URIRef
from rdflib.namespace import OWL, PROV, SH, XSD

from app.main import app
from app.models.schema_model import (
    ComplexTypeDefinition, ElementDeclaration, ModelGroup, ModelGroupKind,
    QName, SchemaModel,
)
from app.services.construction_exchange import UADEX, read_construction_stage_result
from app.services.governed_constraint_source_reader import EXPECTED_COLUMNS
from app.services.logical_schema_serializer import logical_schema_model_to_graph
from app.services.rdf_projection import UAD_SCHEMA
CV = Namespace("https://dynamicontology.com/uad36/constraint-vocabulary#")
CFG = Namespace("https://dynamicontology.com/uad36/configuration#")
LSM = Namespace("https://dynamicontology.com/uad36/logical-schema#")
WORKERS = (
    ("constraint_normalization_agent", "normalize_governed_constraint_to_file",
     "normalized-constraint-representation", None),
    ("logical_schema_binding_agent", "bind_normalized_constraint_to_logical_schema",
     "logical-schema-binding", "normalized_graph"),
    ("context_resolution_agent", "resolve_bound_constraint_context",
     "context-resolution", "bound_graph"),
    ("constraint_behavior_classification", "classify_governed_constraint_behavior",
     "validation-behavior", "context_graph"),
    ("instance_rdf_binding_agent", "bind_governed_constraint_to_instance_rdf",
     "instance-rdf-binding", "classified_graph"),
    ("shacl_generation_agent", "generate_governed_constraint_shacl",
     "shacl-representation", "bound_graph"),
)


def _models(ns: Namespace) -> tuple[Graph, Graph]:
    xml_namespace = "urn:uad36:test:all-rows-pipeline:xml"

    def qname(name: str) -> QName:
        return QName(xml_namespace, name)

    value = ElementDeclaration(name=qname("AddressLineText"))
    address = ElementDeclaration(name=qname("ADDRESS"), type_name=qname("ADDRESS_TYPE"))
    root = ElementDeclaration(name=qname("PROPERTY"), type_name=qname("PROPERTY_TYPE"))
    address_type = ComplexTypeDefinition(
        name=qname("ADDRESS_TYPE"),
        content=ModelGroup(kind=ModelGroupKind.SEQUENCE, elements=(value,)),
    )
    property_type = ComplexTypeDefinition(
        name=qname("PROPERTY_TYPE"),
        content=ModelGroup(kind=ModelGroupKind.SEQUENCE, elements=(address,)),
    )
    schema = logical_schema_model_to_graph(SchemaModel(
        target_namespace=xml_namespace,
        elements={root.name: root},
        complex_types={item.name: item for item in (address_type, property_type)},
    ))
    projection = Graph()
    for triple in schema:
        projection.add(triple)
    for local_name, predicate, range_iri, property_type_iri in (
        ("PROPERTY", ns.property_predicate, ns.property_class, OWL.ObjectProperty),
        ("ADDRESS", ns.address_predicate, ns.address_class, OWL.ObjectProperty),
        ("AddressLineText", ns.value_predicate, XSD.string, OWL.DatatypeProperty),
    ):
        names = {
            subject for subject in schema.subjects(LSM.local_name, Literal(local_name))
            if (subject, RDF.type, LSM.QName) in schema
        }
        assert len(names) == 1
        resource = next(iter(names))
        assert isinstance(resource, URIRef)
        projection.add((resource, UAD_SCHEMA.sourceQName,
                        Literal(f"{{{xml_namespace}}}{local_name}")))
        projection.add((resource, UAD_SCHEMA.projectsTo, predicate))
        projection.add((predicate, RDF.type, property_type_iri))
        projection.add((predicate, RDFS.range, range_iri))
    projection.add((ns.property_class, RDF.type, OWL.Class))
    projection.add((ns.address_class, RDF.type, OWL.Class))
    return schema, projection


class _BuildPage(HTMLParser):
    def __init__(self) -> None:
        super().__init__(convert_charrefs=True)
        self.forms: list[dict] = []
        self.run_ids: list[str] = []
        self._form: dict | None = None
        self._button_parts: list[str] | None = None

    def handle_starttag(self, tag: str, attrs) -> None:
        attributes = dict(attrs)
        if tag == "form":
            self._form = {"attributes": attributes, "inputs": [], "buttons": []}
            self.forms.append(self._form)
        if tag == "input" and self._form is not None:
            self._form["inputs"].append(attributes)
        if tag == "button" and self._form is not None and attributes.get("type", "submit") == "submit":
            self._button_parts = []
        if attributes.get("data-build-run-id"):
            self.run_ids.append(attributes["data-build-run-id"])

    def handle_data(self, data: str) -> None:
        if self._button_parts is not None:
            self._button_parts.append(data)

    def handle_endtag(self, tag: str) -> None:
        if tag == "button" and self._button_parts is not None:
            assert self._form is not None
            self._form["buttons"].append(" ".join("".join(self._button_parts).split()))
            self._button_parts = None
        if tag == "form":
            self._form = None


def _file_artifact(graph: Graph, subject, predicate, output_root: Path) -> Path:
    values = set(graph.objects(subject, predicate))
    assert len(values) == 1
    uri = next(iter(values))
    assert isinstance(uri, URIRef)
    parsed = urlsplit(str(uri))
    assert parsed.scheme == "file" and parsed.netloc in ("", "localhost")
    path = Path(url2pathname(parsed.path)).resolve()
    assert path.is_file()
    assert path.is_relative_to(output_root.resolve())
    return path


@pytest.mark.parametrize("selected_index", [0, 1], ids=["first-workbook", "second-workbook"])
def test_it_61_r1_s2_processes_only_the_exact_leading_zero_row_id_from_selected_workbook(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch, selected_index: int,
) -> None:
    ns = Namespace(f"urn:uad36:test:it61r1s2:{uuid4().hex}:")
    requirement = "Provide the address line for each property."
    violation = "The property address line is absent."
    source_directory = tmp_path / "sources"
    source_directory.mkdir()
    registrations = []
    for number, version in enumerate(("1.5", "1.6")):
        digits = f"{uuid4().int % (10 ** 8):08d}"
        # These two textual IDs become the same number if coerced numerically.
        ids = ("0" + digits, digits.lstrip("0") or "0")
        workbook_path = source_directory / f"registered-source-{number}.xlsx"
        workbook = Workbook()
        wrong_sheet = workbook.active
        assert isinstance(wrong_sheet, Worksheet)
        wrong_sheet.title = "DoNotUse"
        wrong_sheet.append(["The configured worksheet is GovernedRules."])
        sheet = workbook.create_sheet("GovernedRules")
        sheet.append(list(EXPECTED_COLUMNS))
        for unique_id in ids:
            row = dict.fromkeys(EXPECTED_COLUMNS, "")
            row.update({
                "Unique ID": unique_id,
                "Primary Data Element": "AddressLineText",
                "Message ID": "test-rule-" + uuid4().hex,
                "Message Text": requirement,
                "Rule Logic": violation,
                "Severity": "Fatal",
                "Property Affected": "/PROPERTY/ADDRESS/AddressLineText",
                "Data Point Name / Value": "AddressLineText",
                " xPath": "/PROPERTY/ADDRESS/AddressLineText",
            })
            sheet.append([row[column] for column in EXPECTED_COLUMNS])
        for cells in sheet.iter_rows():
            for cell in cells:
                cell.number_format = "@"
        workbook.save(workbook_path)
        workbook.close()
        registrations.append({
            "id": URIRef(f"{ns}set-{number}"),
            "workbook": URIRef(f"{ns}workbook-{number}"),
            "path": workbook_path,
            "version": version,
            "row_ids": ids,
        })
    selected = registrations[selected_index]
    other = registrations[1 - selected_index]
    schema, projection = _models(ns)
    behaviors = Graph()
    for triple in (
        (ns.behavior, RDF.type, CV.ValidationBehavior),
        (ns.behavior, CV.requirementText, Literal(requirement)),
        (ns.behavior, CV.violationCondition, Literal(violation)),
    ):
        behaviors.add(triple)
    patterns = Graph()
    for triple in (
        (ns.behavior, CV.shaclImplementationPattern, ns.pattern),
        (ns.pattern, RDF.type, CV.ShaclImplementationPattern),
        (ns.pattern, CV.propertyShapeTemplate, ns.template),
        (ns.template, RDF.type, SH.PropertyShape),
        (ns.template, SH.minCount, Literal(1)),
    ):
        patterns.add(triple)
    model_directory = tmp_path / "models"
    model_directory.mkdir()
    model_files = {}
    for name, graph in (
        ("logical-schema", schema), ("projection-model", projection),
        ("established-behaviors", behaviors), ("shacl-patterns", patterns),
    ):
        path = model_directory / f"{name}.ttl"
        graph.serialize(destination=path, format="turtle")
        model_files[name] = path
    output_root = tmp_path / "build-runs"
    configuration = Graph()
    for triple in (
        (ns.configuration, RDF.type, CFG.ApplicationConfiguration),
        (ns.configuration, CFG.hasSourceWorkbooks, ns.workbooks),
        (ns.configuration, CFG.hasConstraintSets, ns.constraint_sets),
        (ns.configuration, CFG.activeComplianceSource, other["workbook"]),
        (ns.workbooks, RDF.type, CFG.GovernedSourceWorkbookCollection),
        (ns.constraint_sets, RDF.type, CFG.GovernedConstraintSetCollection),
        (ns.authority, RDF.type, CFG.GoverningAuthority),
        (ns.authority, RDFS.label, Literal("Registered governing authority")),
        # Proposed build-profile contract: dependencies and destinations are
        # explicit governed configuration, never guessed from directories.
        (ns.configuration, CFG.usesConstraintBuildProfile, ns.build_profile),
        (ns.build_profile, RDF.type, CFG.ConstraintBuildProfile),
        (ns.build_profile, CFG.logicalSchemaFileLocation, Literal(str(model_files["logical-schema"]))),
        (ns.build_profile, CFG.projectionModelFileLocation, Literal(str(model_files["projection-model"]))),
        (ns.build_profile, CFG.establishedBehaviorsFileLocation, Literal(str(model_files["established-behaviors"]))),
        (ns.build_profile, CFG.shaclPatternsFileLocation, Literal(str(model_files["shacl-patterns"]))),
        (ns.build_profile, CFG.buildOutputDirectory, Literal(str(output_root))),
    ):
        configuration.add(triple)
    for registration in registrations:
        for triple in (
            (ns.constraint_sets, CFG.containsConstraintSet, registration["id"]),
            (registration["id"], RDF.type, CFG.GovernedConstraintSet),
            (registration["id"], RDFS.label, Literal(f"Governed release {registration['version']}")),
            (registration["id"], CFG.governingAuthority, ns.authority),
            (registration["id"], CFG.sourceWorkbook, registration["workbook"]),
            (registration["id"], CFG.governedVersion, Literal(registration["version"])),
            (ns.workbooks, CFG.containsWorkbook, registration["workbook"]),
            (registration["workbook"], RDF.type, CFG.GovernedSourceArtifact),
            (registration["workbook"], CFG.sourceFileLocation, Literal(str(registration["path"]))),
            (registration["workbook"], CFG.ruleDataSheetName, Literal("GovernedRules")),
            (registration["workbook"], CFG.sourceFormat, Literal("xlsx")),
        ):
            configuration.add(triple)
    # Workbook-scope construction must not require an IT-59 constraint-set registry.
    configuration.remove((ns.configuration, CFG.hasConstraintSets, None))
    configuration.remove((ns.constraint_sets, None, None))
    for registration in registrations:
        configuration.remove((registration["id"], None, None))
        configuration.add((registration["workbook"], RDFS.label,
                           Literal(f"Governed source workbook {registration['version']}")))
    configuration_path = tmp_path / "configuration.ttl"
    configuration.serialize(destination=configuration_path, format="turtle")
    immutable_files = [configuration_path, *model_files.values(), *[item["path"] for item in registrations]]
    originals = {path: path.read_bytes() for path in immutable_files}
    monkeypatch.setattr(app.state, "configuration_file", configuration_path)
    requested_id = selected["row_ids"][0]
    assert requested_id.startswith("0")
    assert int(requested_id) == int(selected["row_ids"][1])
    calls = []
    pipeline_calls = []
    pipeline = import_module("app.services.constraint_production_pipeline")

    def observe_worker(real, index, stage, input_name):
        def execute(*args, **kwargs):
            if input_name is None:
                unique_id = kwargs["governed_source"]["unique_id"]
            else:
                graph = kwargs[input_name]
                constraint = URIRef(kwargs["constraint_id"])
                ids = set(graph.objects(constraint, CV.governedUniqueId))
                assert ids == {Literal(requested_id)}
                unique_id = str(next(iter(ids)))
            assert unique_id == requested_id, "Construction selected another textual ID or workbook row."
            assert len(calls) == index, "Workers were repeated, skipped or executed out of order."
            calls.append(stage)
            return real(*args, **kwargs)
        return execute

    for index, (module_name, operation_name, stage, input_name) in enumerate(WORKERS):
        module = import_module(f"app.services.{module_name}")
        real = getattr(module, operation_name)
        observed = observe_worker(real, index, stage, input_name)
        monkeypatch.setattr(module, operation_name, observed)
        if getattr(pipeline, operation_name, None) is real:
            monkeypatch.setattr(pipeline, operation_name, observed)
    real_process_one = pipeline.process_governed_source_row

    def observe_process_one(**kwargs):
        pipeline_calls.append(kwargs)
        assert kwargs["unique_id"] == requested_id and isinstance(kwargs["unique_id"], str)
        run_configuration = Graph().parse(kwargs["configuration_path"], format="turtle")
        assert set(run_configuration.objects(ns.configuration, CFG.activeComplianceSource)) == {selected["workbook"]}
        return real_process_one(**kwargs)

    monkeypatch.setattr(pipeline, "process_governed_source_row", observe_process_one)
    routes = import_module("app.api.routes")
    if getattr(routes, "process_governed_source_row", None) is real_process_one:
        monkeypatch.setattr(routes, "process_governed_source_row", observe_process_one)

    def no_all_rows_or_evaluation(*args, **kwargs):
        pytest.fail("A one-row UI request must not build all rows or evaluate an appraisal.")

    real_all = pipeline.process_all_governed_source_rows
    monkeypatch.setattr(pipeline, "process_all_governed_source_rows", no_all_rows_or_evaluation)
    if getattr(routes, "process_all_governed_source_rows", None) is real_all:
        monkeypatch.setattr(routes, "process_all_governed_source_rows", no_all_rows_or_evaluation)
    monkeypatch.setattr(import_module("pyshacl"), "validate", no_all_rows_or_evaluation)
    monkeypatch.setattr(import_module("app.services.validation_finding_projection"),
                        "project_governed_findings", no_all_rows_or_evaluation)

    with TestClient(app, follow_redirects=False) as client:
        response = client.get("/constraint-pipeline", params={
            "source_workbook_id": str(selected["workbook"]),
        })
        assert response.status_code == 200
        page = _BuildPage()
        page.feed(response.text)
        forms = [form for form in page.forms
                 if form["attributes"].get("method", "get").lower() == "post"
                 and form["attributes"].get("action") == "/constraint-pipeline/runs"]
        assert len(forms) == 1
        form = forms[0]
        inputs = [item for item in form["inputs"] if item.get("name") == "unique_id"]
        assert len(inputs) == 1 and inputs[0].get("type", "text") == "text", (
            "IT-61R1S2 requires a text field for an exact governed row identifier, "
            "preserving leading zeroes."
        )
        assert any(item.get("name") == "row_scope" and item.get("value") == "one"
                   for item in form["inputs"]), "Offer single-row construction in the run form."
        assert [item.get("value") for item in form["inputs"]
                if item.get("name") == "source_workbook_id"] == [str(selected["workbook"])]
        response = client.post("/constraint-pipeline/runs", data={
            "source_workbook_id": str(selected["workbook"]),
            "row_scope": "one", "unique_id": requested_id,
        })
    assert response.status_code == 200, (
        f"IT-61R1S2 requires a completed single-row construction request; got HTTP {response.status_code}."
    )
    assert "text/html" in response.headers.get("content-type", "")
    page = _BuildPage()
    page.feed(response.text)
    assert len(page.run_ids) == 1
    assert len(pipeline_calls) == 1
    assert calls == [worker[2] for worker in WORKERS]
    manifests = list(output_root.rglob("run.ttl"))
    assert len(manifests) == 1
    directory = manifests[0].parent.resolve()
    manifest = Graph().parse(manifests[0], format="turtle")
    run = URIRef(page.run_ids[0])
    for triple in (
        (run, RDF.type, CV.ConstraintPipelineRun),
        (run, CV.sourceWorkbook, selected["workbook"]),
        (run, CV.sourceWorksheet, Literal("GovernedRules")),
        (run, CV.governedUniqueId, Literal(requested_id)),
        (run, CV.rowScope, Literal("one")),
        (run, CV.status, Literal("COMPLETE")),
    ):
        assert triple in manifest, "The saved run must retain its exact workbook, worksheet and row selection."
    batch_path = _file_artifact(manifest, run, CV.batchResultsArtifact, directory)
    batch = Graph().parse(batch_path, format="turtle")
    constraints = set(batch.subjects(RDF.type, CV.NormalizedConstraint))
    assert len(constraints) == 1
    constraint = next(iter(constraints))
    assert set(batch.objects(constraint, CV.governedUniqueId)) == {Literal(requested_id)}
    records = [read_construction_stage_result(exchange_graph=batch, result_id=str(node))
               for node in batch.subjects(UADEX.forConstraint, constraint)]
    assert len(records) == len(WORKERS)
    assert all(record.construction_status == "GREEN" and record.available_for_downstream for record in records)
    states = {state for state in batch.subjects(RDF.type, CV.ConstraintProductionState)
              if (state, CV.governedConstraint, constraint) in batch}
    assert len(states) == 1
    state = next(iter(states))
    assert (state, CV.status, Literal("COMPLETE")) in batch
    rdf_path = _file_artifact(batch, state, CV.rdfArtifact, directory)
    shapes_path = _file_artifact(batch, state, CV.shaclArtifact, directory)
    rdf = Graph().parse(rdf_path, format="turtle")
    assert (constraint, RDF.type, CV.NormalizedConstraint) in rdf
    assert list(Graph().parse(shapes_path, format="turtle").subjects(CV.implementsConstraint, constraint))
    sources = set(rdf.objects(constraint, PROV.wasDerivedFrom))
    assert len(sources) == 1 and selected["path"].name in str(next(iter(sources)))
    assert not list(batch.subjects(RDF.type, SH.ValidationResult))
    assert {path: path.read_bytes() for path in immutable_files} == originals
