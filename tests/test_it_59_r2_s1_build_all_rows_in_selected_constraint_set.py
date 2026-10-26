"""IT-59R2S1: build every row in exactly the operator-selected constraint set."""

from __future__ import annotations

from collections import defaultdict
from html.parser import HTMLParser
from importlib import import_module
from pathlib import Path
from uuid import uuid4
from urllib.parse import urlsplit
from urllib.request import url2pathname

import pytest
from fastapi.testclient import TestClient
from app.main import app
from openpyxl import Workbook
from openpyxl.worksheet.worksheet import Worksheet
from rdflib import Graph, Literal, Namespace, RDF, RDFS, URIRef
from rdflib.namespace import OWL, PROV, SH, XSD

from app.models.schema_model import (
    ComplexTypeDefinition, ElementDeclaration, ModelGroup, ModelGroupKind,
    QName, SchemaModel,
)
from app.services.construction_exchange import UADEX, read_construction_stage_result
from app.services.governed_constraint_source_reader import EXPECTED_COLUMNS
from app.services.logical_schema_serializer import logical_schema_model_to_graph
from app.services.rdf_projection import UAD_SCHEMA


CV = Namespace("https://dynamicontology.com/uad36/constraint-vocabulary#")
CFG = Namespace("https://dynamicontology.com/uad36/configuration#")
LSM = Namespace("https://dynamicontology.com/uad36/logical-schema#")
WORKERS = (
    ("constraint_normalization_agent", "normalize_governed_constraint_to_file",
     "normalized-constraint-representation", None),
    ("logical_schema_binding_agent", "bind_normalized_constraint_to_logical_schema",
     "logical-schema-binding", "normalized_graph"),
    ("context_resolution_agent", "resolve_bound_constraint_context",
     "context-resolution", "bound_graph"),
    ("constraint_behavior_classification", "classify_governed_constraint_behavior",
     "validation-behavior", "context_graph"),
    ("instance_rdf_binding_agent", "bind_governed_constraint_to_instance_rdf",
     "instance-rdf-binding", "classified_graph"),
    ("shacl_generation_agent", "generate_governed_constraint_shacl",
     "shacl-representation", "bound_graph"),
)


def _models(ns: Namespace) -> tuple[Graph, Graph]:
    xml_namespace = "urn:uad36:test:all-rows-pipeline:xml"

    def qname(name: str) -> QName:
        return QName(xml_namespace, name)

    value = ElementDeclaration(name=qname("AddressLineText"))
    address = ElementDeclaration(name=qname("ADDRESS"), type_name=qname("ADDRESS_TYPE"))
    root = ElementDeclaration(name=qname("PROPERTY"), type_name=qname("PROPERTY_TYPE"))
    address_type = ComplexTypeDefinition(
        name=qname("ADDRESS_TYPE"),
        content=ModelGroup(kind=ModelGroupKind.SEQUENCE, elements=(value,)),
    )
    property_type = ComplexTypeDefinition(
        name=qname("PROPERTY_TYPE"),
        content=ModelGroup(kind=ModelGroupKind.SEQUENCE, elements=(address,)),
    )
    schema = logical_schema_model_to_graph(SchemaModel(
        target_namespace=xml_namespace,
        elements={root.name: root},
        complex_types={item.name: item for item in (address_type, property_type)},
    ))
    projection = Graph()
    for triple in schema:
        projection.add(triple)
    for local_name, predicate, range_iri, property_type_iri in (
        ("PROPERTY", ns.property_predicate, ns.property_class, OWL.ObjectProperty),
        ("ADDRESS", ns.address_predicate, ns.address_class, OWL.ObjectProperty),
        ("AddressLineText", ns.value_predicate, XSD.string, OWL.DatatypeProperty),
    ):
        names = {
            subject for subject in schema.subjects(LSM.local_name, Literal(local_name))
            if (subject, RDF.type, LSM.QName) in schema
        }
        assert len(names) == 1
        resource = next(iter(names))
        assert isinstance(resource, URIRef)
        projection.add((resource, UAD_SCHEMA.sourceQName,
                        Literal(f"{{{xml_namespace}}}{local_name}")))
        projection.add((resource, UAD_SCHEMA.projectsTo, predicate))
        projection.add((predicate, RDF.type, property_type_iri))
        projection.add((predicate, RDFS.range, range_iri))
    projection.add((ns.property_class, RDF.type, OWL.Class))
    projection.add((ns.address_class, RDF.type, OWL.Class))
    return schema, projection


class _BuildPage(HTMLParser):
    def __init__(self) -> None:
        super().__init__(convert_charrefs=True)
        self.forms: list[dict] = []
        self.run_ids: list[str] = []
        self._form: dict | None = None
        self._button_parts: list[str] | None = None

    def handle_starttag(self, tag: str, attrs) -> None:
        attributes = dict(attrs)
        if tag == "form":
            self._form = {"attributes": attributes, "inputs": [], "buttons": []}
            self.forms.append(self._form)
        if tag == "input" and self._form is not None:
            self._form["inputs"].append(attributes)
        if tag == "button" and self._form is not None and attributes.get("type", "submit") == "submit":
            self._button_parts = []
        if attributes.get("data-build-run-id"):
            self.run_ids.append(attributes["data-build-run-id"])

    def handle_data(self, data: str) -> None:
        if self._button_parts is not None:
            self._button_parts.append(data)

    def handle_endtag(self, tag: str) -> None:
        if tag == "button" and self._button_parts is not None:
            assert self._form is not None
            self._form["buttons"].append(" ".join("".join(self._button_parts).split()))
            self._button_parts = None
        if tag == "form":
            self._form = None


def _file_artifact(graph: Graph, subject, predicate, output_root: Path) -> Path:
    values = set(graph.objects(subject, predicate))
    assert len(values) == 1
    uri = next(iter(values))
    assert isinstance(uri, URIRef)
    parsed = urlsplit(str(uri))
    assert parsed.scheme == "file" and parsed.netloc in ("", "localhost")
    path = Path(url2pathname(parsed.path)).resolve()
    assert path.is_file()
    assert path.is_relative_to(output_root.resolve())
    return path


@pytest.mark.parametrize("selected_index", [0, 1], ids=["first-registered-version", "second-registered-version"])
def test_it_59_r2_s1_builds_all_rows_of_only_the_selected_registered_constraint_set(
    tmp_path: Path,
    monkeypatch: pytest.MonkeyPatch,
    selected_index: int,
) -> None:
    ns = Namespace(f"urn:uad36:test:it59r2s1:{uuid4().hex}:")
    requirement = "Provide the address line for each property."
    violation = "The property address line is absent."
    source_directory = tmp_path / "sources"
    source_directory.mkdir()
    registrations = []
    for number, version in enumerate(("1.5", "1.6")):
        ids = tuple("0" + uuid4().hex for _ in range(2))
        workbook_path = source_directory / f"registered-source-{number}.xlsx"
        workbook = Workbook()
        wrong_sheet = workbook.active
        assert isinstance(wrong_sheet, Worksheet)
        wrong_sheet.title = "DoNotUse"
        wrong_sheet.append(["The configured worksheet is GovernedRules."])
        sheet = workbook.create_sheet("GovernedRules")
        sheet.append(list(EXPECTED_COLUMNS))
        for unique_id in ids:
            row = dict.fromkeys(EXPECTED_COLUMNS, "")
            row.update({
                "Unique ID": unique_id,
                "Primary Data Element": "AddressLineText",
                "Message ID": "test-rule-" + uuid4().hex,
                "Message Text": requirement,
                "Rule Logic": violation,
                "Severity": "Fatal",
                "Property Affected": "/PROPERTY/ADDRESS/AddressLineText",
                "Data Point Name / Value": "AddressLineText",
                " xPath": "/PROPERTY/ADDRESS/AddressLineText",
            })
            sheet.append([row[column] for column in EXPECTED_COLUMNS])
        for cells in sheet.iter_rows():
            for cell in cells:
                cell.number_format = "@"
        workbook.save(workbook_path)
        workbook.close()
        registrations.append({
            "id": URIRef(f"{ns}set-{number}"),
            "workbook": URIRef(f"{ns}workbook-{number}"),
            "path": workbook_path,
            "version": version,
            "row_ids": ids,
        })
    selected = registrations[selected_index]
    other = registrations[1 - selected_index]
    schema, projection = _models(ns)
    behaviors = Graph()
    for triple in (
        (ns.behavior, RDF.type, CV.ValidationBehavior),
        (ns.behavior, CV.requirementText, Literal(requirement)),
        (ns.behavior, CV.violationCondition, Literal(violation)),
    ):
        behaviors.add(triple)
    patterns = Graph()
    for triple in (
        (ns.behavior, CV.shaclImplementationPattern, ns.pattern),
        (ns.pattern, RDF.type, CV.ShaclImplementationPattern),
        (ns.pattern, CV.propertyShapeTemplate, ns.template),
        (ns.template, RDF.type, SH.PropertyShape),
        (ns.template, SH.minCount, Literal(1)),
    ):
        patterns.add(triple)
    model_directory = tmp_path / "models"
    model_directory.mkdir()
    model_files = {}
    for name, graph in (
        ("logical-schema", schema), ("projection-model", projection),
        ("established-behaviors", behaviors), ("shacl-patterns", patterns),
    ):
        path = model_directory / f"{name}.ttl"
        graph.serialize(destination=path, format="turtle")
        model_files[name] = path
    output_root = tmp_path / "build-runs"
    configuration = Graph()
    for triple in (
        (ns.configuration, RDF.type, CFG.ApplicationConfiguration),
        (ns.configuration, CFG.hasSourceWorkbooks, ns.workbooks),
        (ns.configuration, CFG.hasConstraintSets, ns.constraint_sets),
        (ns.configuration, CFG.activeComplianceSource, other["workbook"]),
        (ns.workbooks, RDF.type, CFG.GovernedSourceWorkbookCollection),
        (ns.constraint_sets, RDF.type, CFG.GovernedConstraintSetCollection),
        (ns.authority, RDF.type, CFG.GoverningAuthority),
        (ns.authority, RDFS.label, Literal("Registered governing authority")),
        # Proposed build-profile contract: dependencies and destinations are
        # explicit governed configuration, never guessed from directories.
        (ns.configuration, CFG.usesConstraintBuildProfile, ns.build_profile),
        (ns.build_profile, RDF.type, CFG.ConstraintBuildProfile),
        (ns.build_profile, CFG.logicalSchemaFileLocation, Literal(str(model_files["logical-schema"]))),
        (ns.build_profile, CFG.projectionModelFileLocation, Literal(str(model_files["projection-model"]))),
        (ns.build_profile, CFG.establishedBehaviorsFileLocation, Literal(str(model_files["established-behaviors"]))),
        (ns.build_profile, CFG.shaclPatternsFileLocation, Literal(str(model_files["shacl-patterns"]))),
        (ns.build_profile, CFG.buildOutputDirectory, Literal(str(output_root))),
    ):
        configuration.add(triple)
    for registration in registrations:
        for triple in (
            (ns.constraint_sets, CFG.containsConstraintSet, registration["id"]),
            (registration["id"], RDF.type, CFG.GovernedConstraintSet),
            (registration["id"], RDFS.label, Literal(f"Governed release {registration['version']}")),
            (registration["id"], CFG.governingAuthority, ns.authority),
            (registration["id"], CFG.sourceWorkbook, registration["workbook"]),
            (registration["id"], CFG.governedVersion, Literal(registration["version"])),
            (ns.workbooks, CFG.containsWorkbook, registration["workbook"]),
            (registration["workbook"], RDF.type, CFG.GovernedSourceArtifact),
            (registration["workbook"], CFG.sourceFileLocation, Literal(str(registration["path"]))),
            (registration["workbook"], CFG.ruleDataSheetName, Literal("GovernedRules")),
            (registration["workbook"], CFG.sourceFormat, Literal("xlsx")),
        ):
            configuration.add(triple)
    configuration_path = tmp_path / "configuration.ttl"
    configuration.serialize(destination=configuration_path, format="turtle")
    immutable_files = [configuration_path, *model_files.values(), *[item["path"] for item in registrations]]
    originals = {path: path.read_bytes() for path in immutable_files}
    monkeypatch.setattr(app.state, "configuration_file", configuration_path)
    calls = defaultdict(list)
    batch_calls = []
    pipeline = import_module("app.services.constraint_production_pipeline")

    def observe_worker(real, index, stage, input_name):
        def execute(*args, **kwargs):
            if input_name is None:
                unique_id = kwargs["governed_source"]["unique_id"]
            else:
                graph = kwargs[input_name]
                constraint = URIRef(kwargs["constraint_id"])
                ids = set(graph.objects(constraint, CV.governedUniqueId))
                assert len(ids) == 1
                unique_id = str(next(iter(ids)))
                preceding = {
                    node for node in graph.subjects(UADEX.forConstraint, constraint)
                    if (node, UADEX.constructionStage, Literal(WORKERS[index - 1][2])) in graph
                }
                assert len(preceding) == 1
                record = read_construction_stage_result(
                    exchange_graph=graph, result_id=str(next(iter(preceding))),
                )
                assert record.construction_status == "GREEN" and record.available_for_downstream
            assert unique_id in selected["row_ids"], "The build processed an unselected source row."
            assert index == len(calls[unique_id]), "The build skipped, repeated, or reordered workers."
            calls[unique_id].append(stage)
            return real(*args, **kwargs)
        return execute

    for index, (module_name, operation_name, stage, input_name) in enumerate(WORKERS):
        module = import_module(f"app.services.{module_name}")
        real = getattr(module, operation_name)
        observed = observe_worker(real, index, stage, input_name)
        monkeypatch.setattr(module, operation_name, observed)
        if getattr(pipeline, operation_name, None) is real:
            monkeypatch.setattr(pipeline, operation_name, observed)
    real_process_all = pipeline.process_all_governed_source_rows

    def observe_batch(**kwargs):
        batch_calls.append(kwargs)
        selected_configuration = Graph().parse(kwargs["configuration_path"], format="turtle")
        assert set(selected_configuration.objects(ns.configuration, CFG.activeComplianceSource)) == {selected["workbook"]}
        return real_process_all(**kwargs)

    monkeypatch.setattr(pipeline, "process_all_governed_source_rows", observe_batch)

    def no_evaluation(*args, **kwargs):
        pytest.fail("A constraint-set build must not evaluate an appraisal or project runtime findings.")

    monkeypatch.setattr(import_module("pyshacl"), "validate", no_evaluation)
    monkeypatch.setattr(import_module("app.services.validation_finding_projection"), "project_governed_findings", no_evaluation)
    with TestClient(app, follow_redirects=False) as client:
        selection = client.get("/constraint-workbench", params={
            "authority": str(ns.authority), "constraint_set_id": str(selected["id"]),
        })
        assert selection.status_code == 200
        page = _BuildPage()
        page.feed(selection.text)
        build_forms = [form for form in page.forms
                       if form["attributes"].get("action") == "/constraint-workbench/build"
                       and form["attributes"].get("method", "get").lower() == "post"]
        assert len(build_forms) == 1, (
            "IT-59R2S1 requires a Build action for the operator-selected "
            "registered constraint set, submitting that set to the workbench build handler."
        )
        form = build_forms[0]
        assert "Build" in form["buttons"]
        submitted = [item.get("value", "") for item in form["inputs"] if item.get("name") == "constraint_set_id"]
        assert submitted == [str(selected["id"])], "Build input contains another set or lacks the selected set."
        response = client.post("/constraint-workbench/build", data={
            "authority": str(ns.authority), "constraint_set_id": str(selected["id"]),
        })
    assert response.status_code in (200, 202), f"Build request was not accepted: HTTP {response.status_code}."
    assert "text/html" in response.headers.get("content-type", "")
    result_page = _BuildPage()
    result_page.feed(response.text)
    assert len(result_page.run_ids) == 1, "The workbench must identify the construction run it started."
    assert len(batch_calls) == 1, "A selected-set build must use one IT-58 all-row pipeline request."
    assert set(calls) == set(selected["row_ids"])
    for unique_id in selected["row_ids"]:
        assert calls[unique_id] == [worker[2] for worker in WORKERS]

    # Proposed saved run contract: a run.ttl manifest records the selected
    # set/version and points to the persistent IT-58 batch results.
    manifests = list(output_root.rglob("run.ttl"))
    assert len(manifests) == 1, "The build must persist its run identity and selected source version."
    run_graph = Graph().parse(manifests[0], format="turtle")
    run_ids = set(run_graph.subjects(RDF.type, CV.ConstraintSetBuildRun))
    assert len(run_ids) == 1
    run = next(iter(run_ids))
    assert isinstance(run, URIRef)
    assert str(run) == result_page.run_ids[0]
    assert (run, RDF.type, PROV.Activity) in run_graph
    assert set(run_graph.objects(run, CV.selectedConstraintSet)) == {selected["id"]}
    assert set(run_graph.objects(run, CV.sourceVersion)) == {Literal(selected["version"])}
    assert set(run_graph.objects(run, CV.sourceWorkbook)) == {selected["workbook"]}
    assert (run, PROV.used, selected["id"]) in run_graph
    assert set(run_graph.objects(run, CV.status)) == {Literal("COMPLETE")}
    batch_path = _file_artifact(run_graph, run, CV.batchResultsArtifact, output_root)
    batch = Graph().parse(batch_path, format="turtle")
    constraints = set(batch.subjects(RDF.type, CV.NormalizedConstraint))
    assert len(constraints) == 2
    assert {str(value) for constraint in constraints for value in batch.objects(constraint, CV.governedUniqueId)} == set(selected["row_ids"])
    assert set(run_graph.objects(run, CV.producedConstraint)) == constraints
    for constraint in constraints:
        records = [read_construction_stage_result(exchange_graph=batch, result_id=str(node))
                   for node in batch.subjects(UADEX.forConstraint, constraint)]
        assert len(records) == len(WORKERS)
        assert {record.construction_stage for record in records} == {worker[2] for worker in WORKERS}
        assert all(record.construction_status == "GREEN" and record.available_for_downstream for record in records)
        states = {state for state in batch.subjects(RDF.type, CV.ConstraintProductionState)
                  if (state, CV.governedConstraint, constraint) in batch}
        assert len(states) == 1
        state = next(iter(states))
        assert (state, CV.status, Literal("COMPLETE")) in batch
        rdf_path = _file_artifact(batch, state, CV.rdfArtifact, output_root)
        shacl_path = _file_artifact(batch, state, CV.shaclArtifact, output_root)
        rdf = Graph().parse(rdf_path, format="turtle")
        shapes = Graph().parse(shacl_path, format="turtle")
        assert (constraint, RDF.type, CV.NormalizedConstraint) in rdf
        assert list(shapes.subjects(CV.implementsConstraint, constraint))
        sources = set(rdf.objects(constraint, PROV.wasDerivedFrom))
        assert len(sources) == 1
        assert selected["path"].name in str(next(iter(sources)))
    assert not list(batch.subjects(RDF.type, SH.ValidationResult))
    assert {path: path.read_bytes() for path in immutable_files} == originals


