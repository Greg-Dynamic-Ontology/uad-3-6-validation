"""IT-61R1S4: construct all rows of only the selected governed workbook tab.

Proposed request: POST /constraint-pipeline/runs with source_workbook_id and
row_scope="all"; no individual row IDs are required. One ConstraintPipelineRun
records rowScope="all", the workbook/worksheet and all governedUniqueIds.
The run page and completion summary expose data-run-status,
data-total-row-count and data-completed-row-count, with visible 'X of N' progress.

Real workers execute. At the beginning of each real normalization invocation,
a separate HTTP viewer reads the saved in-progress run. These observations
verify progress during construction, not just a fabricated completion summary.
No worker is mocked as successful. Helpers and fixtures are self-contained.
"""
from html.parser import HTMLParser
from importlib import import_module
from pathlib import Path
from urllib.parse import urlsplit
from urllib.request import url2pathname
from uuid import uuid4

import pytest
from fastapi.testclient import TestClient
from openpyxl import Workbook
from openpyxl.worksheet.worksheet import Worksheet
from rdflib import Graph, Literal, Namespace, RDF, RDFS, URIRef
from rdflib.namespace import OWL, PROV, SH, XSD

from app.main import app
from app.models.schema_model import (
    ComplexTypeDefinition, ElementDeclaration, ModelGroup, ModelGroupKind,
    QName, SchemaModel,
)
from app.services.construction_exchange import UADEX, read_construction_stage_result
from app.services.governed_constraint_source_reader import EXPECTED_COLUMNS
from app.services.logical_schema_serializer import logical_schema_model_to_graph
from app.services.rdf_projection import UAD_SCHEMA
CV = Namespace("https://dynamicontology.com/uad36/constraint-vocabulary#")
CFG = Namespace("https://dynamicontology.com/uad36/configuration#")
LSM = Namespace("https://dynamicontology.com/uad36/logical-schema#")
WORKERS = (
    ("constraint_normalization_agent", "normalize_governed_constraint_to_file",
     "normalized-constraint-representation", None),
    ("logical_schema_binding_agent", "bind_normalized_constraint_to_logical_schema",
     "logical-schema-binding", "normalized_graph"),
    ("context_resolution_agent", "resolve_bound_constraint_context",
     "context-resolution", "bound_graph"),
    ("constraint_behavior_classification", "classify_governed_constraint_behavior",
     "validation-behavior", "context_graph"),
    ("instance_rdf_binding_agent", "bind_governed_constraint_to_instance_rdf",
     "instance-rdf-binding", "classified_graph"),
    ("shacl_generation_agent", "generate_governed_constraint_shacl",
     "shacl-representation", "bound_graph"),
)


def _models(ns: Namespace) -> tuple[Graph, Graph]:
    xml_namespace = "urn:uad36:test:all-rows-pipeline:xml"

    def qname(name: str) -> QName:
        return QName(xml_namespace, name)

    value = ElementDeclaration(name=qname("AddressLineText"))
    address = ElementDeclaration(name=qname("ADDRESS"), type_name=qname("ADDRESS_TYPE"))
    root = ElementDeclaration(name=qname("PROPERTY"), type_name=qname("PROPERTY_TYPE"))
    address_type = ComplexTypeDefinition(
        name=qname("ADDRESS_TYPE"),
        content=ModelGroup(kind=ModelGroupKind.SEQUENCE, elements=(value,)),
    )
    property_type = ComplexTypeDefinition(
        name=qname("PROPERTY_TYPE"),
        content=ModelGroup(kind=ModelGroupKind.SEQUENCE, elements=(address,)),
    )
    schema = logical_schema_model_to_graph(SchemaModel(
        target_namespace=xml_namespace,
        elements={root.name: root},
        complex_types={item.name: item for item in (address_type, property_type)},
    ))
    projection = Graph()
    for triple in schema:
        projection.add(triple)
    for local_name, predicate, range_iri, property_type_iri in (
        ("PROPERTY", ns.property_predicate, ns.property_class, OWL.ObjectProperty),
        ("ADDRESS", ns.address_predicate, ns.address_class, OWL.ObjectProperty),
        ("AddressLineText", ns.value_predicate, XSD.string, OWL.DatatypeProperty),
    ):
        names = {
            subject for subject in schema.subjects(LSM.local_name, Literal(local_name))
            if (subject, RDF.type, LSM.QName) in schema
        }
        assert len(names) == 1
        resource = next(iter(names))
        assert isinstance(resource, URIRef)
        projection.add((resource, UAD_SCHEMA.sourceQName,
                        Literal(f"{{{xml_namespace}}}{local_name}")))
        projection.add((resource, UAD_SCHEMA.projectsTo, predicate))
        projection.add((predicate, RDF.type, property_type_iri))
        projection.add((predicate, RDFS.range, range_iri))
    projection.add((ns.property_class, RDF.type, OWL.Class))
    projection.add((ns.address_class, RDF.type, OWL.Class))
    return schema, projection


class _BuildPage(HTMLParser):
    def __init__(self) -> None:
        super().__init__(convert_charrefs=True)
        self.forms: list[dict] = []
        self.run_ids: list[str] = []
        self._form: dict | None = None
        self._button_parts: list[str] | None = None

    def handle_starttag(self, tag: str, attrs) -> None:
        attributes = dict(attrs)
        if tag == "form":
            self._form = {"attributes": attributes, "inputs": [], "buttons": []}
            self.forms.append(self._form)
        if tag == "input" and self._form is not None:
            self._form["inputs"].append(attributes)
        if tag == "button" and self._form is not None and attributes.get("type", "submit") == "submit":
            self._button_parts = []
        if attributes.get("data-build-run-id"):
            self.run_ids.append(attributes["data-build-run-id"])

    def handle_data(self, data: str) -> None:
        if self._button_parts is not None:
            self._button_parts.append(data)

    def handle_endtag(self, tag: str) -> None:
        if tag == "button" and self._button_parts is not None:
            assert self._form is not None
            self._form["buttons"].append(" ".join("".join(self._button_parts).split()))
            self._button_parts = None
        if tag == "form":
            self._form = None


def _file_artifact(graph: Graph, subject, predicate, output_root: Path) -> Path:
    values = set(graph.objects(subject, predicate))
    assert len(values) == 1
    uri = next(iter(values))
    assert isinstance(uri, URIRef)
    parsed = urlsplit(str(uri))
    assert parsed.scheme == "file" and parsed.netloc in ("", "localhost")
    path = Path(url2pathname(parsed.path)).resolve()
    assert path.is_file()
    assert path.is_relative_to(output_root.resolve())
    return path


class _AllRowsPage(_BuildPage):
    def __init__(self):
        super().__init__()
        self.progress = []
        self.visible = []
        self.hidden_depth = 0

    def handle_starttag(self, tag, attrs):
        super().handle_starttag(tag, attrs)
        attributes = dict(attrs)
        if "data-run-status" in attributes:
            self.progress.append((
                attributes["data-run-status"], attributes.get("data-completed-row-count"),
                attributes.get("data-total-row-count"),
            ))
        if tag in ("script", "style"):
            self.hidden_depth += 1

    def handle_endtag(self, tag):
        super().handle_endtag(tag)
        if tag in ("script", "style"):
            self.hidden_depth -= 1

    def handle_data(self, data):
        super().handle_data(data)
        if not self.hidden_depth:
            self.visible.append(data)


def _assert_progress(response, status, completed, total):
    assert response.status_code == 200
    page = _AllRowsPage()
    page.feed(response.text)
    assert page.progress == [(status, str(completed), str(total))], (
        "IT-61R1S4 requires actual saved-run progress and completion counts in the UI."
    )
    visible = " ".join(" ".join(page.visible).split())
    assert f"{completed} of {total}" in visible, "Show the processed-row count visibly to the operator."
    return page


@pytest.mark.parametrize("selected_index", [0, 1], ids=["first-workbook", "second-workbook"])
def test_it_61_r1_s4_processes_all_selected_worksheet_rows_with_live_progress_and_provenance(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch, selected_index: int,
) -> None:
    ns = Namespace(f"urn:uad36:test:it61r1s4:{uuid4().hex}:")
    requirement = "Provide the address line for each property."
    violation = "The property address line is absent."
    source_directory = tmp_path / "sources"
    source_directory.mkdir()
    registrations = []
    for number, version in enumerate(("1.5", "1.6")):
        digits = f"{uuid4().int % (10 ** 8):08d}"
        # These two textual IDs become the same number if coerced numerically.
        ids = ("0" + digits, digits.lstrip("0") or "0")
        workbook_path = source_directory / f"registered-source-{number}.xlsx"
        workbook = Workbook()
        wrong_sheet = workbook.active
        assert isinstance(wrong_sheet, Worksheet)
        wrong_sheet.title = "DoNotUse"
        wrong_sheet.append(["The configured worksheet is GovernedRules."])
        sheet = workbook.create_sheet("GovernedRules")
        sheet.append(list(EXPECTED_COLUMNS))
        for unique_id in ids:
            row = dict.fromkeys(EXPECTED_COLUMNS, "")
            row.update({
                "Unique ID": unique_id,
                "Primary Data Element": "AddressLineText",
                "Message ID": "test-rule-" + uuid4().hex,
                "Message Text": requirement,
                "Rule Logic": violation,
                "Severity": "Fatal",
                "Property Affected": "/PROPERTY/ADDRESS/AddressLineText",
                "Data Point Name / Value": "AddressLineText",
                " xPath": "/PROPERTY/ADDRESS/AddressLineText",
            })
            sheet.append([row[column] for column in EXPECTED_COLUMNS])
        for cells in sheet.iter_rows():
            for cell in cells:
                cell.number_format = "@"
        workbook.save(workbook_path)
        workbook.close()
        registrations.append({
            "id": URIRef(f"{ns}set-{number}"),
            "workbook": URIRef(f"{ns}workbook-{number}"),
            "path": workbook_path,
            "version": version,
            "row_ids": ids,
        })
    selected = registrations[selected_index]
    other = registrations[1 - selected_index]
    schema, projection = _models(ns)
    behaviors = Graph()
    for triple in (
        (ns.behavior, RDF.type, CV.ValidationBehavior),
        (ns.behavior, CV.requirementText, Literal(requirement)),
        (ns.behavior, CV.violationCondition, Literal(violation)),
    ):
        behaviors.add(triple)
    patterns = Graph()
    for triple in (
        (ns.behavior, CV.shaclImplementationPattern, ns.pattern),
        (ns.pattern, RDF.type, CV.ShaclImplementationPattern),
        (ns.pattern, CV.propertyShapeTemplate, ns.template),
        (ns.template, RDF.type, SH.PropertyShape),
        (ns.template, SH.minCount, Literal(1)),
    ):
        patterns.add(triple)
    model_directory = tmp_path / "models"
    model_directory.mkdir()
    model_files = {}
    for name, graph in (
        ("logical-schema", schema), ("projection-model", projection),
        ("established-behaviors", behaviors), ("shacl-patterns", patterns),
    ):
        path = model_directory / f"{name}.ttl"
        graph.serialize(destination=path, format="turtle")
        model_files[name] = path
    output_root = tmp_path / "build-runs"
    configuration = Graph()
    for triple in (
        (ns.configuration, RDF.type, CFG.ApplicationConfiguration),
        (ns.configuration, CFG.hasSourceWorkbooks, ns.workbooks),
        (ns.configuration, CFG.hasConstraintSets, ns.constraint_sets),
        (ns.configuration, CFG.activeComplianceSource, other["workbook"]),
        (ns.workbooks, RDF.type, CFG.GovernedSourceWorkbookCollection),
        (ns.constraint_sets, RDF.type, CFG.GovernedConstraintSetCollection),
        (ns.authority, RDF.type, CFG.GoverningAuthority),
        (ns.authority, RDFS.label, Literal("Registered governing authority")),
        # Proposed build-profile contract: dependencies and destinations are
        # explicit governed configuration, never guessed from directories.
        (ns.configuration, CFG.usesConstraintBuildProfile, ns.build_profile),
        (ns.build_profile, RDF.type, CFG.ConstraintBuildProfile),
        (ns.build_profile, CFG.logicalSchemaFileLocation, Literal(str(model_files["logical-schema"]))),
        (ns.build_profile, CFG.projectionModelFileLocation, Literal(str(model_files["projection-model"]))),
        (ns.build_profile, CFG.establishedBehaviorsFileLocation, Literal(str(model_files["established-behaviors"]))),
        (ns.build_profile, CFG.shaclPatternsFileLocation, Literal(str(model_files["shacl-patterns"]))),
        (ns.build_profile, CFG.buildOutputDirectory, Literal(str(output_root))),
    ):
        configuration.add(triple)
    for registration in registrations:
        for triple in (
            (ns.constraint_sets, CFG.containsConstraintSet, registration["id"]),
            (registration["id"], RDF.type, CFG.GovernedConstraintSet),
            (registration["id"], RDFS.label, Literal(f"Governed release {registration['version']}")),
            (registration["id"], CFG.governingAuthority, ns.authority),
            (registration["id"], CFG.sourceWorkbook, registration["workbook"]),
            (registration["id"], CFG.governedVersion, Literal(registration["version"])),
            (ns.workbooks, CFG.containsWorkbook, registration["workbook"]),
            (registration["workbook"], RDF.type, CFG.GovernedSourceArtifact),
            (registration["workbook"], CFG.sourceFileLocation, Literal(str(registration["path"]))),
            (registration["workbook"], CFG.ruleDataSheetName, Literal("GovernedRules")),
            (registration["workbook"], CFG.sourceFormat, Literal("xlsx")),
        ):
            configuration.add(triple)
    # Workbook-scope construction must not require an IT-59 constraint-set registry.
    configuration.remove((ns.configuration, CFG.hasConstraintSets, None))
    configuration.remove((ns.constraint_sets, None, None))
    for registration in registrations:
        configuration.remove((registration["id"], None, None))
        configuration.add((registration["workbook"], RDFS.label,
                           Literal(f"Governed source workbook {registration['version']}")))
    configuration_path = tmp_path / "configuration.ttl"
    configuration.serialize(destination=configuration_path, format="turtle")
    immutable_files = [configuration_path, *model_files.values(), *[item["path"] for item in registrations]]
    originals = {path: path.read_bytes() for path in immutable_files}
    monkeypatch.setattr(app.state, "configuration_file", configuration_path)
    requested_ids = tuple(selected["row_ids"])
    calls = {unique_id: [] for unique_id in requested_ids}
    all_requests = []
    progress_observations = []
    pipeline = import_module("app.services.constraint_production_pipeline")

    def observe_worker(real, index, stage, input_name):
        def execute(*args, **kwargs):
            if input_name is None:
                unique_id = kwargs["governed_source"]["unique_id"]
                assert unique_id in requested_ids, "All-row processing selected another workbook."
                manifests = list(output_root.rglob("run.ttl"))
                assert len(manifests) == 1
                directory = manifests[0].parent
                completed = sum(len(values) == len(WORKERS) for values in calls.values())
                with TestClient(app, follow_redirects=False) as monitor:
                    response = monitor.get(f"/constraint-workbench/runs/{directory.name}", params={"playback": "paused"})
                _assert_progress(response, "RUNNING", completed, len(requested_ids))
                progress_observations.append(completed)
            else:
                graph = kwargs[input_name]
                constraint = URIRef(kwargs["constraint_id"])
                ids = set(graph.objects(constraint, CV.governedUniqueId))
                assert len(ids) == 1
                unique_id = str(next(iter(ids)))
            assert unique_id in requested_ids, "The pipeline processed an unselected workbook row."
            assert len(calls[unique_id]) == index, "A row repeated, skipped or reordered production workers."
            calls[unique_id].append(stage)
            return real(*args, **kwargs)
        return execute

    for index, (module_name, operation_name, stage, input_name) in enumerate(WORKERS):
        module = import_module(f"app.services.{module_name}")
        real = getattr(module, operation_name)
        observed = observe_worker(real, index, stage, input_name)
        monkeypatch.setattr(module, operation_name, observed)
        if getattr(pipeline, operation_name, None) is real:
            monkeypatch.setattr(pipeline, operation_name, observed)
    real_all = pipeline.process_all_governed_source_rows

    def observe_all(**kwargs):
        all_requests.append(kwargs)
        run_configuration = Graph().parse(kwargs["configuration_path"], format="turtle")
        assert set(run_configuration.objects(ns.configuration, CFG.activeComplianceSource)) == {selected["workbook"]}
        return real_all(**kwargs)

    monkeypatch.setattr(pipeline, "process_all_governed_source_rows", observe_all)
    routes = import_module("app.api.routes")
    if getattr(routes, "process_all_governed_source_rows", None) is real_all:
        monkeypatch.setattr(routes, "process_all_governed_source_rows", observe_all)

    def no_single_row_or_evaluation(*args, **kwargs):
        pytest.fail("All-row construction must use the all-row pipeline and must not evaluate an appraisal.")

    real_one = pipeline.process_governed_source_row
    monkeypatch.setattr(pipeline, "process_governed_source_row", no_single_row_or_evaluation)
    if getattr(routes, "process_governed_source_row", None) is real_one:
        monkeypatch.setattr(routes, "process_governed_source_row", no_single_row_or_evaluation)
    monkeypatch.setattr(import_module("pyshacl"), "validate", no_single_row_or_evaluation)
    monkeypatch.setattr(import_module("app.services.validation_finding_projection"),
                        "project_governed_findings", no_single_row_or_evaluation)

    with TestClient(app, follow_redirects=False) as client:
        response = client.get("/constraint-pipeline", params={"source_workbook_id": str(selected["workbook"])})
        assert response.status_code == 200
        page = _BuildPage()
        page.feed(response.text)
        forms = [form for form in page.forms
                 if form["attributes"].get("action") == "/constraint-pipeline/runs"
                 and form["attributes"].get("method", "get").lower() == "post"]
        assert len(forms) == 1
        assert any(item.get("name") == "row_scope" and item.get("value") == "all"
                   for item in forms[0]["inputs"]), (
            "IT-61R1S4 requires an all-rows scope for the selected workbook in the construction-run form."
        )
        assert not all_requests
        response = client.post("/constraint-pipeline/runs", data={
            "source_workbook_id": str(selected["workbook"]), "row_scope": "all",
        })
    page = _assert_progress(response, "COMPLETE", len(requested_ids), len(requested_ids))
    assert len(page.run_ids) == 1
    assert len(all_requests) == 1
    assert progress_observations == list(range(len(requested_ids)))
    for unique_id in requested_ids:
        assert calls[unique_id] == [worker[2] for worker in WORKERS]
    manifests = list(output_root.rglob("run.ttl"))
    assert len(manifests) == 1
    directory = manifests[0].parent.resolve()
    manifest = Graph().parse(manifests[0], format="turtle")
    run = URIRef(page.run_ids[0])
    for triple in (
        (run, RDF.type, CV.ConstraintPipelineRun),
        (run, CV.sourceWorkbook, selected["workbook"]),
        (run, CV.sourceWorksheet, Literal("GovernedRules")),
        (run, CV.rowScope, Literal("all")),
        (run, CV.status, Literal("COMPLETE")),
        (run, CV.completedRowCount, Literal(len(requested_ids))),
    ):
        assert triple in manifest
    assert set(manifest.objects(run, CV.governedUniqueId)) == {Literal(value) for value in requested_ids}
    batch_path = _file_artifact(manifest, run, CV.batchResultsArtifact, directory)
    batch = Graph().parse(batch_path, format="turtle")
    constraints = set(batch.subjects(RDF.type, CV.NormalizedConstraint))
    assert len(constraints) == len(requested_ids)
    assert {str(value) for constraint in constraints
            for value in batch.objects(constraint, CV.governedUniqueId)} == set(requested_ids)
    for constraint in constraints:
        ids = set(batch.objects(constraint, CV.governedUniqueId))
        assert len(ids) == 1
        records = [read_construction_stage_result(exchange_graph=batch, result_id=str(node))
                   for node in batch.subjects(UADEX.forConstraint, constraint)]
        assert len(records) == len(WORKERS)
        assert all(record.construction_status == "GREEN" and record.available_for_downstream for record in records)
        states = {state for state in batch.subjects(RDF.type, CV.ConstraintProductionState)
                  if (state, CV.governedConstraint, constraint) in batch}
        assert len(states) == 1
        state = next(iter(states))
        assert (state, CV.status, Literal("COMPLETE")) in batch
        assert set(batch.objects(state, CV.governedUniqueId)) == ids
        assert (state, CV.sourceWorkbook, selected["workbook"]) in batch
        assert (state, CV.sourceWorksheet, Literal("GovernedRules")) in batch
        rdf_path = _file_artifact(batch, state, CV.rdfArtifact, directory)
        shapes_path = _file_artifact(batch, state, CV.shaclArtifact, directory)
        rdf = Graph().parse(rdf_path, format="turtle")
        assert (constraint, RDF.type, CV.NormalizedConstraint) in rdf
        assert list(Graph().parse(shapes_path, format="turtle").subjects(CV.implementsConstraint, constraint))
        sources = set(rdf.objects(constraint, PROV.wasDerivedFrom))
        assert len(sources) == 1 and selected["path"].name in str(next(iter(sources)))
    before_view = {path: path.read_bytes() for path in directory.rglob("*") if path.is_file()}
    with TestClient(app, follow_redirects=False) as client:
        _assert_progress(client.get(f"/constraint-workbench/runs/{directory.name}"),
                         "COMPLETE", len(requested_ids), len(requested_ids))
    assert {path: path.read_bytes() for path in directory.rglob("*") if path.is_file()} == before_view
    assert not list(batch.subjects(RDF.type, SH.ValidationResult))
    assert {path: path.read_bytes() for path in immutable_files} == originals
