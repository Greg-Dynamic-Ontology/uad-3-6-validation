"""IT-58R1S2: construct all configured rows and retain row exceptions."""

from __future__ import annotations

import csv
from collections import defaultdict
from importlib import import_module
from pathlib import Path
from uuid import uuid4

import pytest
from openpyxl import Workbook
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet
from rdflib import Graph, Literal, Namespace, RDF, RDFS, URIRef
from rdflib.namespace import OWL, PROV, SH, XSD

from app.models.schema_model import (
    ComplexTypeDefinition, ElementDeclaration, ModelGroup, ModelGroupKind,
    QName, SchemaModel,
)
from app.services.construction_exchange import UADEX, read_construction_stage_result
from app.services.governed_constraint_source_reader import EXPECTED_COLUMNS
from app.services.logical_schema_serializer import logical_schema_model_to_graph
from app.services.rdf_projection import UAD_SCHEMA


CV = Namespace("https://dynamicontology.com/uad36/constraint-vocabulary#")
CFG = Namespace("https://dynamicontology.com/uad36/configuration#")
LSM = Namespace("https://dynamicontology.com/uad36/logical-schema#")
WORKERS = (
    ("constraint_normalization_agent", "normalize_governed_constraint_to_file",
     "normalized-constraint-representation", None),
    ("logical_schema_binding_agent", "bind_normalized_constraint_to_logical_schema",
     "logical-schema-binding", "normalized_graph"),
    ("context_resolution_agent", "resolve_bound_constraint_context",
     "context-resolution", "bound_graph"),
    ("constraint_behavior_classification", "classify_governed_constraint_behavior",
     "validation-behavior", "context_graph"),
    ("instance_rdf_binding_agent", "bind_governed_constraint_to_instance_rdf",
     "instance-rdf-binding", "classified_graph"),
    ("shacl_generation_agent", "generate_governed_constraint_shacl",
     "shacl-representation", "bound_graph"),
)


def _models(ns: Namespace) -> tuple[Graph, Graph]:
    xml_namespace = "urn:uad36:test:all-rows-pipeline:xml"

    def qname(name: str) -> QName:
        return QName(xml_namespace, name)

    value = ElementDeclaration(name=qname("AddressLineText"))
    address = ElementDeclaration(name=qname("ADDRESS"), type_name=qname("ADDRESS_TYPE"))
    root = ElementDeclaration(name=qname("PROPERTY"), type_name=qname("PROPERTY_TYPE"))
    address_type = ComplexTypeDefinition(
        name=qname("ADDRESS_TYPE"),
        content=ModelGroup(kind=ModelGroupKind.SEQUENCE, elements=(value,)),
    )
    property_type = ComplexTypeDefinition(
        name=qname("PROPERTY_TYPE"),
        content=ModelGroup(kind=ModelGroupKind.SEQUENCE, elements=(address,)),
    )
    schema = logical_schema_model_to_graph(SchemaModel(
        target_namespace=xml_namespace,
        elements={root.name: root},
        complex_types={item.name: item for item in (address_type, property_type)},
    ))
    projection = Graph()
    for triple in schema:
        projection.add(triple)
    for local_name, predicate, range_iri, property_type_iri in (
        ("PROPERTY", ns.property_predicate, ns.property_class, OWL.ObjectProperty),
        ("ADDRESS", ns.address_predicate, ns.address_class, OWL.ObjectProperty),
        ("AddressLineText", ns.value_predicate, XSD.string, OWL.DatatypeProperty),
    ):
        names = {
            subject for subject in schema.subjects(LSM.local_name, Literal(local_name))
            if (subject, RDF.type, LSM.QName) in schema
        }
        assert len(names) == 1
        resource = next(iter(names))
        assert isinstance(resource, URIRef)
        projection.add((resource, UAD_SCHEMA.sourceQName,
                        Literal(f"{{{xml_namespace}}}{local_name}")))
        projection.add((resource, UAD_SCHEMA.projectsTo, predicate))
        projection.add((predicate, RDF.type, property_type_iri))
        projection.add((predicate, RDFS.range, range_iri))
    projection.add((ns.property_class, RDF.type, OWL.Class))
    projection.add((ns.address_class, RDF.type, OWL.Class))
    return schema, projection


@pytest.mark.parametrize("source_format", ["csv", "xlsx"], ids=["configured-csv", "configured-excel"])
@pytest.mark.parametrize("failure_kind", ["blank-required-knowledge", "unknown-behavior"])
def test_it_58_r1_s2_processes_all_rows_and_records_independent_exceptions(
    tmp_path: Path,
    monkeypatch: pytest.MonkeyPatch,
    source_format: str,
    failure_kind: str,
) -> None:
    pipeline = import_module("app.services.constraint_production_pipeline")
    process_all = getattr(pipeline, "process_all_governed_source_rows", None)
    assert callable(process_all), (
        "IT-58R1S2 requires process_all_governed_source_rows() to consider every "
        "configured source row, retain row-level exceptions, and continue "
        "constructing reusable RDF and SHACL for the other eligible rows."
    )

    ns = Namespace(f"urn:uad36:test:it58r1s2:{uuid4().hex}:")
    ids = tuple("0" + uuid4().hex for _ in range(3))
    before_id, failed_id, after_id = ids
    good_ids = {before_id, after_id}
    requirement = "Provide the address line for each property."
    violation = "The property address line is absent."
    rows = []
    for unique_id in ids:
        row = dict.fromkeys(EXPECTED_COLUMNS, "")
        row.update({
            "Unique ID": unique_id,
            "Primary Data Element": "AddressLineText",
            "Message ID": "test-rule-" + uuid4().hex,
            "Message Text": requirement,
            "Rule Logic": violation,
            "Severity": "Fatal",
            "Property Affected": "/PROPERTY/ADDRESS/AddressLineText",
            "Data Point Name / Value": "AddressLineText",
            " xPath": "/PROPERTY/ADDRESS/AddressLineText",
        })
        rows.append(row)
    if failure_kind == "blank-required-knowledge":
        rows[1]["Message Text"] = ""
    else:
        rows[1]["Rule Logic"] = "An unregistered governed violation " + uuid4().hex

    source_path = tmp_path / f"constraints.{source_format}"
    if source_format == "csv":
        with source_path.open("w", encoding="utf-8", newline="") as stream:
            writer = csv.DictWriter(stream, fieldnames=EXPECTED_COLUMNS)
            writer.writeheader()
            writer.writerows(rows)
    else:
        workbook = Workbook()
        wrong_sheet = workbook.active
        assert isinstance(wrong_sheet, Worksheet)
        wrong_sheet.title = "DoNotUse"
        wrong_sheet.append(["This is not the configured rule-data worksheet."])
        sheet = workbook.create_sheet("GovernedRules")
        sheet.append(list(EXPECTED_COLUMNS))
        for row in rows:
            sheet.append([row[column] for column in EXPECTED_COLUMNS])
        for cells in sheet.iter_rows():
            for cell in cells:
                cell.number_format = "@"
        workbook.save(source_path)
        workbook.close()
    source_before = source_path.read_bytes()
    last_column = get_column_letter(len(EXPECTED_COLUMNS))
    sources = {
        unique_id: URIRef(
            f"urn:uad36:governed-source:{source_path.name}#row={number}"
            if source_format == "csv" else
            f"urn:uad36:governed-source:{source_path.name}#GovernedRules!A{number}:{last_column}{number}"
        )
        for number, unique_id in enumerate(ids, start=2)
    }
    configuration = Graph()
    for triple in (
        (ns.configuration, RDF.type, CFG.ApplicationConfiguration),
        (ns.configuration, CFG.hasSourceWorkbooks, ns.collection),
        (ns.configuration, CFG.activeComplianceSource, ns.source),
        (ns.collection, RDF.type, CFG.GovernedSourceWorkbookCollection),
        (ns.collection, CFG.containsWorkbook, ns.source),
        (ns.source, RDF.type, CFG.GovernedSourceArtifact),
        (ns.source, CFG.sourceFileLocation, Literal(str(source_path))),
        (ns.source, CFG.sourceFormat, Literal(source_format)),
    ):
        configuration.add(triple)
    if source_format == "xlsx":
        configuration.add((ns.source, CFG.ruleDataSheetName, Literal("GovernedRules")))
    configuration_path = tmp_path / "configuration.ttl"
    configuration.serialize(destination=configuration_path, format="turtle")
    configuration_before = configuration_path.read_bytes()

    schema, projection = _models(ns)
    behaviors = Graph()
    for triple in (
        (ns.behavior, RDF.type, CV.ValidationBehavior),
        (ns.behavior, CV.requirementText, Literal(requirement)),
        (ns.behavior, CV.violationCondition, Literal(violation)),
    ):
        behaviors.add(triple)
    patterns = Graph()
    for triple in (
        (ns.behavior, CV.shaclImplementationPattern, ns.pattern),
        (ns.pattern, RDF.type, CV.ShaclImplementationPattern),
        (ns.pattern, CV.propertyShapeTemplate, ns.template),
        (ns.template, RDF.type, SH.PropertyShape),
        (ns.template, SH.minCount, Literal(1)),
    ):
        patterns.add(triple)
    model_inputs = (schema, projection, behaviors, patterns)
    models_before = [set(graph) for graph in model_inputs]
    calls = defaultdict(list)

    def observe_worker(real, stage, input_name):
        def observed(*args, **kwargs):
            if input_name is None:
                unique_id = kwargs["governed_source"]["unique_id"]
            else:
                graph = kwargs[input_name]
                identities = set(graph.objects(None, CV.governedUniqueId))
                assert len(identities) == 1, "A worker received mixed row knowledge."
                unique_id = str(next(iter(identities)))
                previous_stage = WORKERS[len(calls[unique_id]) - 1][2]
                results = set(graph.subjects(UADEX.constructionStage, Literal(previous_stage)))
                assert len(results) == 1, "Preceding verified stage result is missing."
                record = read_construction_stage_result(
                    exchange_graph=graph, result_id=str(next(iter(results))),
                )
                assert record.construction_status == "GREEN"
                assert record.available_for_downstream is True
            assert unique_id in ids, "An unconfigured source row was processed."
            index = len(calls[unique_id])
            assert index < len(WORKERS), "A source row was processed more than once."
            assert stage == WORKERS[index][2], "Construction stages ran out of order."
            calls[unique_id].append(stage)
            return real(*args, **kwargs)
        return observed

    for module_name, operation_name, stage, input_name in WORKERS:
        worker_module = import_module(f"app.services.{module_name}")
        real = getattr(worker_module, operation_name)
        observed = observe_worker(real, stage, input_name)
        monkeypatch.setattr(worker_module, operation_name, observed)
        if getattr(pipeline, operation_name, None) is real:
            monkeypatch.setattr(pipeline, operation_name, observed)

    def no_evaluation(*args, **kwargs):
        pytest.fail("All-row construction must not evaluate appraisal XML or generate findings.")

    monkeypatch.setattr(import_module("pyshacl"), "validate", no_evaluation)
    monkeypatch.setattr(
        import_module("app.services.validation_finding_projection"),
        "project_governed_findings", no_evaluation,
    )
    output_directory = tmp_path / "production-artifacts"
    batch_path = tmp_path / "batch-state" / "results.ttl"
    result = process_all(
        configuration_path=configuration_path,
        logical_schema_graph=schema,
        projection_model_graph=projection,
        established_behaviors_graph=behaviors,
        implementation_patterns_graph=patterns,
        output_directory=output_directory,
        batch_results_path=batch_path,
    )

    # Proposed all-row contract: completed results plus explicit row exceptions.
    completed = list(result.completed)
    exceptions = list(result.exceptions)
    assert len(completed) == 2
    assert {item.unique_id for item in completed} == good_ids
    assert len(exceptions) == 1
    failure = exceptions[0]
    assert failure.unique_id == failed_id
    assert failure.exception_reason and str(failure.exception_reason).strip()
    expected_stages = [worker[2] for worker in WORKERS]
    for unique_id in good_ids:
        assert calls[unique_id] == expected_stages
    if failure_kind == "blank-required-knowledge":
        assert not calls[failed_id], "An insufficient source row reached normalization."
    else:
        assert calls[failed_id] == expected_stages[:4]

    batch = Graph().parse(batch_path, format="turtle")
    artifact_paths = set()
    for item in completed:
        assert item.is_complete is True
        constraint = URIRef(item.constraint_id)
        rdf_path, shacl_path = Path(item.rdf_path), Path(item.shacl_path)
        for artifact in (rdf_path, shacl_path):
            assert artifact.is_file()
            assert artifact.resolve().is_relative_to(output_directory.resolve())
            assert artifact.resolve() not in artifact_paths
            artifact_paths.add(artifact.resolve())
        rdf = Graph().parse(rdf_path, format="turtle")
        shapes = Graph().parse(shacl_path, format="turtle")
        assert set(rdf.subjects(RDF.type, CV.NormalizedConstraint)) == {constraint}
        assert set(rdf.objects(constraint, CV.governedUniqueId)) == {Literal(item.unique_id)}
        assert (constraint, PROV.wasDerivedFrom, sources[item.unique_id]) in rdf
        shape_ids = set(shapes.subjects(RDF.type, SH.NodeShape))
        assert len(shape_ids) == 1
        shape = next(iter(shape_ids))
        assert (shape, CV.implementsConstraint, constraint) in shapes
        assert (shape, PROV.wasDerivedFrom, sources[item.unique_id]) in shapes
        assert not list(rdf.subjects(RDF.type, SH.ValidationResult))
        records = [
            read_construction_stage_result(exchange_graph=batch, result_id=str(node))
            for node in batch.subjects(UADEX.forConstraint, constraint)
        ]
        assert len(records) == len(WORKERS)
        assert {record.construction_stage for record in records} == set(expected_stages)
        assert all(record.construction_status == "GREEN" and record.available_for_downstream
                   for record in records)

    # Even a row rejected before normalization must have a persistent,
    # source-traceable exception, rather than silently disappearing.
    exception_id = URIRef(str(failure.result_id))
    assert (exception_id, UADEX.constructionStatus, Literal("EXCEPTION")) in batch
    assert (exception_id, CV.governedUniqueId, Literal(failed_id)) in batch
    assert (exception_id, PROV.wasDerivedFrom, sources[failed_id]) in batch
    assert set(batch.objects(exception_id, UADEX.exceptionReason)) == {
        Literal(failure.exception_reason)
    }
    assert not list(batch.subjects(RDF.type, SH.ValidationResult))
    assert source_path.read_bytes() == source_before
    assert configuration_path.read_bytes() == configuration_before
    assert [set(graph) for graph in model_inputs] == models_before
