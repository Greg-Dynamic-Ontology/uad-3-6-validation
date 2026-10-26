"""IT-61R2S1: UI construction uses the six real workers in governed order.

For single-row, selected-group and all-row UI requests, observe real worker
entry/return. Each downstream input must contain the preceding result's exact
RDF identity, GREEN/available contract and unchanged RDF knowledge. It must
match the persisted stage snapshot whose digest the batch records. IT-58 must
coordinate every worker call; IT-57 and appraisal validation are prohibited.

Only observers and forbidden-runtime guards are installed. No worker, exchange
result or successful artifact is mocked. All fixtures/helpers are local.
"""
from hashlib import sha256
from html.parser import HTMLParser
from importlib import import_module
from pathlib import Path
from urllib.parse import urlsplit
from urllib.request import url2pathname
from uuid import uuid4

import pytest
from fastapi.testclient import TestClient
from openpyxl import Workbook
from openpyxl.worksheet.worksheet import Worksheet
from rdflib import Graph, Literal, Namespace, RDF, RDFS, URIRef
from rdflib.namespace import OWL, PROV, SH, XSD
from rdflib.compare import isomorphic

from app.main import app
from app.models.schema_model import (
    ComplexTypeDefinition, ElementDeclaration, ModelGroup, ModelGroupKind,
    QName, SchemaModel,
)
from app.services.construction_exchange import UADEX, read_construction_stage_result
from app.services.governed_constraint_source_reader import EXPECTED_COLUMNS
from app.services.logical_schema_serializer import logical_schema_model_to_graph
from app.services.rdf_projection import UAD_SCHEMA
CV = Namespace("https://dynamicontology.com/uad36/constraint-vocabulary#")
CFG = Namespace("https://dynamicontology.com/uad36/configuration#")
LSM = Namespace("https://dynamicontology.com/uad36/logical-schema#")
WORKERS = (
    ("constraint_normalization_agent", "normalize_governed_constraint_to_file",
     "normalized-constraint-representation", None),
    ("logical_schema_binding_agent", "bind_normalized_constraint_to_logical_schema",
     "logical-schema-binding", "normalized_graph"),
    ("context_resolution_agent", "resolve_bound_constraint_context",
     "context-resolution", "bound_graph"),
    ("constraint_behavior_classification", "classify_governed_constraint_behavior",
     "validation-behavior", "context_graph"),
    ("instance_rdf_binding_agent", "bind_governed_constraint_to_instance_rdf",
     "instance-rdf-binding", "classified_graph"),
    ("shacl_generation_agent", "generate_governed_constraint_shacl",
     "shacl-representation", "bound_graph"),
)


def _models(ns: Namespace) -> tuple[Graph, Graph]:
    xml_namespace = "urn:uad36:test:all-rows-pipeline:xml"

    def qname(name: str) -> QName:
        return QName(xml_namespace, name)

    value = ElementDeclaration(name=qname("AddressLineText"))
    address = ElementDeclaration(name=qname("ADDRESS"), type_name=qname("ADDRESS_TYPE"))
    root = ElementDeclaration(name=qname("PROPERTY"), type_name=qname("PROPERTY_TYPE"))
    address_type = ComplexTypeDefinition(
        name=qname("ADDRESS_TYPE"),
        content=ModelGroup(kind=ModelGroupKind.SEQUENCE, elements=(value,)),
    )
    property_type = ComplexTypeDefinition(
        name=qname("PROPERTY_TYPE"),
        content=ModelGroup(kind=ModelGroupKind.SEQUENCE, elements=(address,)),
    )
    schema = logical_schema_model_to_graph(SchemaModel(
        target_namespace=xml_namespace,
        elements={root.name: root},
        complex_types={item.name: item for item in (address_type, property_type)},
    ))
    projection = Graph()
    for triple in schema:
        projection.add(triple)
    for local_name, predicate, range_iri, property_type_iri in (
        ("PROPERTY", ns.property_predicate, ns.property_class, OWL.ObjectProperty),
        ("ADDRESS", ns.address_predicate, ns.address_class, OWL.ObjectProperty),
        ("AddressLineText", ns.value_predicate, XSD.string, OWL.DatatypeProperty),
    ):
        names = {
            subject for subject in schema.subjects(LSM.local_name, Literal(local_name))
            if (subject, RDF.type, LSM.QName) in schema
        }
        assert len(names) == 1
        resource = next(iter(names))
        assert isinstance(resource, URIRef)
        projection.add((resource, UAD_SCHEMA.sourceQName,
                        Literal(f"{{{xml_namespace}}}{local_name}")))
        projection.add((resource, UAD_SCHEMA.projectsTo, predicate))
        projection.add((predicate, RDF.type, property_type_iri))
        projection.add((predicate, RDFS.range, range_iri))
    projection.add((ns.property_class, RDF.type, OWL.Class))
    projection.add((ns.address_class, RDF.type, OWL.Class))
    return schema, projection


class _BuildPage(HTMLParser):
    def __init__(self) -> None:
        super().__init__(convert_charrefs=True)
        self.forms: list[dict] = []
        self.run_ids: list[str] = []
        self._form: dict | None = None
        self._button_parts: list[str] | None = None

    def handle_starttag(self, tag: str, attrs) -> None:
        attributes = dict(attrs)
        if tag == "form":
            self._form = {"attributes": attributes, "inputs": [], "buttons": []}
            self.forms.append(self._form)
        if tag == "input" and self._form is not None:
            self._form["inputs"].append(attributes)
        if tag == "button" and self._form is not None and attributes.get("type", "submit") == "submit":
            self._button_parts = []
        if attributes.get("data-build-run-id"):
            self.run_ids.append(attributes["data-build-run-id"])

    def handle_data(self, data: str) -> None:
        if self._button_parts is not None:
            self._button_parts.append(data)

    def handle_endtag(self, tag: str) -> None:
        if tag == "button" and self._button_parts is not None:
            assert self._form is not None
            self._form["buttons"].append(" ".join("".join(self._button_parts).split()))
            self._button_parts = None
        if tag == "form":
            self._form = None


def _file_artifact(graph: Graph, subject, predicate, output_root: Path) -> Path:
    values = set(graph.objects(subject, predicate))
    assert len(values) == 1
    uri = next(iter(values))
    assert isinstance(uri, URIRef)
    parsed = urlsplit(str(uri))
    assert parsed.scheme == "file" and parsed.netloc in ("", "localhost")
    path = Path(url2pathname(parsed.path)).resolve()
    assert path.is_file()
    assert path.is_relative_to(output_root.resolve())
    return path


@pytest.mark.parametrize("row_scope", ["one", "group", "all"], ids=["one-row", "selected-group", "all-rows"])
def test_it_61_r2_s1_executes_real_workers_in_order_through_verified_rdf_handoffs(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch, row_scope: str,
) -> None:
    selected_index = 0
    ns = Namespace(f"urn:uad36:test:it61r2s1:{uuid4().hex}:")
    requirement = "Provide the address line for each property."
    violation = "The property address line is absent."
    source_directory = tmp_path / "sources"
    source_directory.mkdir()
    registrations = []
    for number, version in enumerate(("1.5", "1.6")):
        digits = f"{uuid4().int % (10 ** 8):08d}"
        # These two textual IDs become the same number if coerced numerically.
        ids = ("0" + digits, digits.lstrip("0") or "0")
        workbook_path = source_directory / f"registered-source-{number}.xlsx"
        workbook = Workbook()
        wrong_sheet = workbook.active
        assert isinstance(wrong_sheet, Worksheet)
        wrong_sheet.title = "DoNotUse"
        wrong_sheet.append(["The configured worksheet is GovernedRules."])
        sheet = workbook.create_sheet("GovernedRules")
        sheet.append(list(EXPECTED_COLUMNS))
        for unique_id in ids:
            row = dict.fromkeys(EXPECTED_COLUMNS, "")
            row.update({
                "Unique ID": unique_id,
                "Primary Data Element": "AddressLineText",
                "Message ID": "test-rule-" + uuid4().hex,
                "Message Text": requirement,
                "Rule Logic": violation,
                "Severity": "Fatal",
                "Property Affected": "/PROPERTY/ADDRESS/AddressLineText",
                "Data Point Name / Value": "AddressLineText",
                " xPath": "/PROPERTY/ADDRESS/AddressLineText",
            })
            sheet.append([row[column] for column in EXPECTED_COLUMNS])
        for cells in sheet.iter_rows():
            for cell in cells:
                cell.number_format = "@"
        workbook.save(workbook_path)
        workbook.close()
        registrations.append({
            "id": URIRef(f"{ns}set-{number}"),
            "workbook": URIRef(f"{ns}workbook-{number}"),
            "path": workbook_path,
            "version": version,
            "row_ids": ids,
        })
    selected = registrations[selected_index]
    other = registrations[1 - selected_index]
    schema, projection = _models(ns)
    behaviors = Graph()
    for triple in (
        (ns.behavior, RDF.type, CV.ValidationBehavior),
        (ns.behavior, CV.requirementText, Literal(requirement)),
        (ns.behavior, CV.violationCondition, Literal(violation)),
    ):
        behaviors.add(triple)
    patterns = Graph()
    for triple in (
        (ns.behavior, CV.shaclImplementationPattern, ns.pattern),
        (ns.pattern, RDF.type, CV.ShaclImplementationPattern),
        (ns.pattern, CV.propertyShapeTemplate, ns.template),
        (ns.template, RDF.type, SH.PropertyShape),
        (ns.template, SH.minCount, Literal(1)),
    ):
        patterns.add(triple)
    model_directory = tmp_path / "models"
    model_directory.mkdir()
    model_files = {}
    for name, graph in (
        ("logical-schema", schema), ("projection-model", projection),
        ("established-behaviors", behaviors), ("shacl-patterns", patterns),
    ):
        path = model_directory / f"{name}.ttl"
        graph.serialize(destination=path, format="turtle")
        model_files[name] = path
    output_root = tmp_path / "build-runs"
    configuration = Graph()
    for triple in (
        (ns.configuration, RDF.type, CFG.ApplicationConfiguration),
        (ns.configuration, CFG.hasSourceWorkbooks, ns.workbooks),
        (ns.configuration, CFG.hasConstraintSets, ns.constraint_sets),
        (ns.configuration, CFG.activeComplianceSource, other["workbook"]),
        (ns.workbooks, RDF.type, CFG.GovernedSourceWorkbookCollection),
        (ns.constraint_sets, RDF.type, CFG.GovernedConstraintSetCollection),
        (ns.authority, RDF.type, CFG.GoverningAuthority),
        (ns.authority, RDFS.label, Literal("Registered governing authority")),
        # Proposed build-profile contract: dependencies and destinations are
        # explicit governed configuration, never guessed from directories.
        (ns.configuration, CFG.usesConstraintBuildProfile, ns.build_profile),
        (ns.build_profile, RDF.type, CFG.ConstraintBuildProfile),
        (ns.build_profile, CFG.logicalSchemaFileLocation, Literal(str(model_files["logical-schema"]))),
        (ns.build_profile, CFG.projectionModelFileLocation, Literal(str(model_files["projection-model"]))),
        (ns.build_profile, CFG.establishedBehaviorsFileLocation, Literal(str(model_files["established-behaviors"]))),
        (ns.build_profile, CFG.shaclPatternsFileLocation, Literal(str(model_files["shacl-patterns"]))),
        (ns.build_profile, CFG.buildOutputDirectory, Literal(str(output_root))),
    ):
        configuration.add(triple)
    for registration in registrations:
        for triple in (
            (ns.constraint_sets, CFG.containsConstraintSet, registration["id"]),
            (registration["id"], RDF.type, CFG.GovernedConstraintSet),
            (registration["id"], RDFS.label, Literal(f"Governed release {registration['version']}")),
            (registration["id"], CFG.governingAuthority, ns.authority),
            (registration["id"], CFG.sourceWorkbook, registration["workbook"]),
            (registration["id"], CFG.governedVersion, Literal(registration["version"])),
            (ns.workbooks, CFG.containsWorkbook, registration["workbook"]),
            (registration["workbook"], RDF.type, CFG.GovernedSourceArtifact),
            (registration["workbook"], CFG.sourceFileLocation, Literal(str(registration["path"]))),
            (registration["workbook"], CFG.ruleDataSheetName, Literal("GovernedRules")),
            (registration["workbook"], CFG.sourceFormat, Literal("xlsx")),
        ):
            configuration.add(triple)
    # Workbook-scope construction must not require an IT-59 constraint-set registry.
    configuration.remove((ns.configuration, CFG.hasConstraintSets, None))
    configuration.remove((ns.constraint_sets, None, None))
    for registration in registrations:
        configuration.remove((registration["id"], None, None))
        configuration.add((registration["workbook"], RDFS.label,
                           Literal(f"Governed source workbook {registration['version']}")))
    configuration_path = tmp_path / "configuration.ttl"
    configuration.serialize(destination=configuration_path, format="turtle")
    immutable_files = [configuration_path, *model_files.values(), *[item["path"] for item in registrations]]
    originals = {path: path.read_bytes() for path in immutable_files}
    monkeypatch.setattr(app.state, "configuration_file", configuration_path)
    requested_ids = tuple(selected["row_ids"][:1] if row_scope == "one" else selected["row_ids"])
    events = {value: [] for value in requested_ids}
    returned_knowledge = {}
    handoffs = []
    coordinator_requests = []
    active_requests = []
    pipeline = import_module("app.services.constraint_production_pipeline")
    routes = import_module("app.api.routes")
    workbench = import_module("app.services.constraint_pipeline_workbench")

    def copy_graph(graph):
        copied = Graph()
        for triple in graph:
            copied.add(triple)
        return copied

    def observe_coordinator(real, name):
        def coordinate(**kwargs):
            coordinator_requests.append((name, kwargs))
            active_requests.append(kwargs)
            try:
                return real(**kwargs)
            finally:
                active_requests.pop()
        return coordinate

    for operation_name in ("process_governed_source_row", "process_all_governed_source_rows"):
        real = getattr(pipeline, operation_name)
        observed = observe_coordinator(real, operation_name)
        monkeypatch.setattr(pipeline, operation_name, observed)
        for consumer in (routes, workbench):
            if getattr(consumer, operation_name, None) is real:
                monkeypatch.setattr(consumer, operation_name, observed)

    def observe_worker(real, index, stage, input_name):
        def execute(*args, **kwargs):
            assert len(active_requests) == 1, "IT-58 must coordinate every worker invocation."
            if index == 0:
                unique_id = kwargs["governed_source"]["unique_id"]
                assert isinstance(unique_id, str)
            else:
                incoming = kwargs[input_name]
                assert isinstance(incoming, Graph), "Handoffs must carry governed RDF, not worker-specific dictionaries."
                constraint = URIRef(kwargs["constraint_id"])
                ids = set(incoming.objects(constraint, CV.governedUniqueId))
                assert len(ids) == 1
                unique_id = str(next(iter(ids)))
                assert unique_id in requested_ids
                assert events[unique_id] == [(item[2], event) for item in WORKERS[:index]
                                              for event in ("enter", "return")], (
                    "A downstream worker ran before all preceding workers returned in governed order."
                )
                upstream_stage = WORKERS[index - 1][2]
                result_nodes = {
                    node for node in incoming.subjects(UADEX.forConstraint, constraint)
                    if (node, UADEX.constructionStage, Literal(upstream_stage)) in incoming
                }
                assert len(result_nodes) == 1
                result_id = next(iter(result_nodes))
                record = read_construction_stage_result(exchange_graph=incoming, result_id=str(result_id))
                assert record.construction_status == "GREEN" and record.available_for_downstream, (
                    "A downstream worker must consume only verified GREEN upstream output."
                )
                assert set(returned_knowledge[unique_id, index - 1]).issubset(set(incoming)), (
                    "The preceding worker's produced RDF was lost or changed during handoff."
                )
                locations = set(incoming.objects(result_id, UADEX.rdfSnapshotPath))
                assert len(locations) == 1
                path = Path(str(next(iter(locations))))
                saved_bytes = path.read_bytes()
                saved = Graph().parse(data=saved_bytes, format="turtle")
                assert isomorphic(saved, incoming), "The next worker received knowledge different from the verified snapshot."
                batch = Graph().parse(active_requests[-1]["batch_results_path"], format="turtle")
                assert (result_id, RDF.type, UADEX.ConstructionStageResult) in batch
                assert set(batch.objects(result_id, UADEX.snapshotSha256)) == {Literal(sha256(saved_bytes).hexdigest())}
                persisted = read_construction_stage_result(exchange_graph=batch, result_id=str(result_id))
                assert persisted == record, "The handoff replaced or changed the saved exchange result identity."
                prior_stages = {str(value) for node in incoming.subjects(UADEX.forConstraint, constraint)
                                for value in incoming.objects(node, UADEX.constructionStage)}
                assert prior_stages == {item[2] for item in WORKERS[:index]}
                handoffs.append((unique_id, upstream_stage, stage, result_id))
            assert unique_id in requested_ids
            assert len(events[unique_id]) == index * 2, "Worker execution was repeated or out of order."
            events[unique_id].append((stage, "enter"))
            result = real(*args, **kwargs)
            if index == 0:
                graph = Graph().parse(kwargs["output_path"], format="turtle")
            else:
                graph = kwargs["output_graph"]
                assert set(incoming).issubset(set(graph)), "A completed worker discarded verified upstream knowledge."
            returned_knowledge[unique_id, index] = copy_graph(graph)
            events[unique_id].append((stage, "return"))
            return result
        return execute

    for index, (module_name, operation_name, stage, input_name) in enumerate(WORKERS):
        module = import_module(f"app.services.{module_name}")
        real = getattr(module, operation_name)
        observed = observe_worker(real, index, stage, input_name)
        monkeypatch.setattr(module, operation_name, observed)
        if getattr(pipeline, operation_name, None) is real:
            monkeypatch.setattr(pipeline, operation_name, observed)

    def no_appraisal_execution(*args, **kwargs):
        pytest.fail("IT-61R2S1 construction must not evaluate appraisal XML or execute IT-57 runtime finding projection.")

    pyshacl = import_module("pyshacl")
    projection_service = import_module("app.services.validation_finding_projection")
    for module, name in ((pyshacl, "validate"), (projection_service, "project_governed_findings")):
        real = getattr(module, name)
        monkeypatch.setattr(module, name, no_appraisal_execution)
        for consumer in (routes, workbench, pipeline):
            if getattr(consumer, name, None) is real:
                monkeypatch.setattr(consumer, name, no_appraisal_execution)
    monkeypatch.setattr(import_module("app.services.validation").validation_service,
                        "validate", no_appraisal_execution)

    request_data = {"source_workbook_id": str(selected["workbook"]), "row_scope": row_scope}
    if row_scope == "one":
        request_data["unique_id"] = requested_ids[0]
    elif row_scope == "group":
        request_data["unique_ids"] = list(requested_ids)
    with TestClient(app, follow_redirects=False) as client:
        response = client.post("/constraint-pipeline/runs", data=request_data)
    assert response.status_code == 200
    page = _BuildPage()
    page.feed(response.text)
    assert len(page.run_ids) == 1
    expected_events = [(worker[2], event) for worker in WORKERS for event in ("enter", "return")]
    for unique_id in requested_ids:
        assert events[unique_id] == expected_events
    assert len(handoffs) == len(requested_ids) * (len(WORKERS) - 1)
    expected_operation = "process_all_governed_source_rows" if row_scope == "all" else "process_governed_source_row"
    assert [name for name, _ in coordinator_requests] == [expected_operation] * (1 if row_scope == "all" else len(requested_ids))
    assert not active_requests
    manifests = list(output_root.rglob("run.ttl"))
    assert len(manifests) == 1
    manifest = Graph().parse(manifests[0], format="turtle")
    run = URIRef(page.run_ids[0])
    assert (run, CV.status, Literal("COMPLETE")) in manifest
    batch_path = _file_artifact(manifest, run, CV.batchResultsArtifact, manifests[0].parent)
    batch = Graph().parse(batch_path, format="turtle")
    constraints = set(batch.subjects(RDF.type, CV.NormalizedConstraint))
    assert len(constraints) == len(requested_ids)
    for constraint in constraints:
        records = [read_construction_stage_result(exchange_graph=batch, result_id=str(node))
                   for node in batch.subjects(UADEX.forConstraint, constraint)]
        assert len(records) == len(WORKERS)
        assert {record.construction_stage for record in records} == {worker[2] for worker in WORKERS}
        assert all(record.construction_status == "GREEN" and record.available_for_downstream for record in records)
    assert not list(batch.subjects(RDF.type, SH.ValidationResult))
    assert {path: path.read_bytes() for path in immutable_files} == originals
