"""IT-58R1S1: process exactly one configured source row through real workers."""

from __future__ import annotations

import csv
from importlib import import_module
from importlib.util import find_spec
from pathlib import Path
from uuid import uuid4

import pytest
from openpyxl import Workbook
from openpyxl.utils import get_column_letter
from rdflib import Graph, Literal, Namespace, RDF, RDFS, URIRef
from rdflib.namespace import OWL, PROV, SH, XSD

from app.models.schema_model import (
    ComplexTypeDefinition, ElementDeclaration, ModelGroup, ModelGroupKind, QName, SchemaModel,
)
from app.services.construction_exchange import UADEX, read_construction_stage_result
from app.services.governed_constraint_source_reader import (
    EXPECTED_COLUMNS, read_governed_constraint_row_from_csv, read_governed_constraint_row,
)
from app.services.logical_schema_serializer import logical_schema_model_to_graph
from app.services.rdf_projection import UAD_SCHEMA


CV = Namespace("https://dynamicontology.com/uad36/constraint-vocabulary#")
LSM = Namespace("https://dynamicontology.com/uad36/logical-schema#")
PIPELINE_MODULE = "app.services.constraint_production_pipeline"
WORKERS = (
    ("constraint_normalization_agent", "normalize_governed_constraint_to_file",
     "normalized-constraint-representation", None),
    ("logical_schema_binding_agent", "bind_normalized_constraint_to_logical_schema",
     "logical-schema-binding", "normalized_graph"),
    ("context_resolution_agent", "resolve_bound_constraint_context",
     "context-resolution", "bound_graph"),
    ("constraint_behavior_classification", "classify_governed_constraint_behavior",
     "validation-behavior", "context_graph"),
    ("instance_rdf_binding_agent", "bind_governed_constraint_to_instance_rdf",
     "instance-rdf-binding", "classified_graph"),
    ("shacl_generation_agent", "generate_governed_constraint_shacl",
     "shacl-representation", "bound_graph"),
)


def _models(ns: Namespace) -> tuple[Graph, Graph]:
    xml_namespace = "urn:uad36:test:single-row-pipeline:xml"
    def qname(name: str) -> QName:
        return QName(xml_namespace, name)
    value = ElementDeclaration(name=qname("AddressLineText"))
    address = ElementDeclaration(name=qname("ADDRESS"), type_name=qname("ADDRESS_TYPE"))
    root = ElementDeclaration(name=qname("PROPERTY"), type_name=qname("PROPERTY_TYPE"))
    address_type = ComplexTypeDefinition(
        name=qname("ADDRESS_TYPE"),
        content=ModelGroup(kind=ModelGroupKind.SEQUENCE, elements=(value,)),
    )
    property_type = ComplexTypeDefinition(
        name=qname("PROPERTY_TYPE"),
        content=ModelGroup(kind=ModelGroupKind.SEQUENCE, elements=(address,)),
    )
    schema = logical_schema_model_to_graph(SchemaModel(
        target_namespace=xml_namespace,
        elements={root.name: root},
        complex_types={item.name: item for item in (address_type, property_type)},
    ))
    projection = Graph()
    for triple in schema:
        projection.add(triple)
    for local_name, predicate, range_iri, property_type_iri in (
        ("PROPERTY", ns.property_predicate, ns.property_class, OWL.ObjectProperty),
        ("ADDRESS", ns.address_predicate, ns.address_class, OWL.ObjectProperty),
        ("AddressLineText", ns.value_predicate, XSD.string, OWL.DatatypeProperty),
    ):
        names = {
            subject for subject in schema.subjects(LSM.local_name, Literal(local_name))
            if (subject, RDF.type, LSM.QName) in schema
        }
        assert len(names) == 1
        resource = next(iter(names))
        projection.add((resource, UAD_SCHEMA.sourceQName, Literal(f"{{{xml_namespace}}}{local_name}")))
        projection.add((resource, UAD_SCHEMA.projectsTo, predicate))
        projection.add((predicate, RDF.type, property_type_iri))
        projection.add((predicate, RDFS.range, range_iri))
    projection.add((ns.property_class, RDF.type, OWL.Class))
    projection.add((ns.address_class, RDF.type, OWL.Class))
    return schema, projection


@pytest.mark.parametrize("source_format", ["csv", "xlsx"], ids=["configured-csv", "configured-excel"])
def test_it_58_r1_s1_processes_only_the_requested_governed_source_row(
    tmp_path: Path,
    monkeypatch: pytest.MonkeyPatch,
    source_format: str,
) -> None:
    ns = Namespace(f"urn:uad36:test:it58r1s1:{uuid4().hex}:")
    selected_id = "0" + uuid4().hex
    other_id = "0" + uuid4().hex
    requirement = "Provide the address line for each property."
    violation = "The property address line is absent."
    selected = {column: "" for column in EXPECTED_COLUMNS}
    selected.update({
        "Unique ID": selected_id,
        "Primary Data Element": "AddressLineText",
        "Message ID": "test-rule-" + uuid4().hex,
        "Message Text": requirement,
        "Rule Logic": violation,
        "Severity": "Fatal",
        "Property Affected": "/PROPERTY/ADDRESS/AddressLineText",
        "Data Point Name / Value": "AddressLineText",
        " xPath": "/PROPERTY/ADDRESS/AddressLineText",
    })
    other = dict(selected)
    other.update({"Unique ID": other_id, "Message Text": ""})
    source_directory = tmp_path / "source-fixture"
    source_directory.mkdir()
    if source_format == "csv":
        source_path = source_directory / "data-constraints.csv"
        with source_path.open("w", encoding="utf-8", newline="") as stream:
            writer = csv.DictWriter(stream, fieldnames=EXPECTED_COLUMNS)
            writer.writeheader()
            writer.writerows([other, selected])
        selected_source = read_governed_constraint_row_from_csv(source_path, selected_id)
    else:
        source_path = source_directory / "constraints.xlsx"
        workbook = Workbook()
        wrong_sheet = workbook.active
        wrong_sheet.title = "DoNotUse"
        wrong_sheet.append(list(EXPECTED_COLUMNS))
        wrong_sheet.append([other[column] for column in EXPECTED_COLUMNS])
        worksheet = workbook.create_sheet("GovernedRules")
        worksheet.append(list(EXPECTED_COLUMNS))
        worksheet.append([other[column] for column in EXPECTED_COLUMNS])
        worksheet.append([selected[column] for column in EXPECTED_COLUMNS])
        for sheet in workbook:
            for row in sheet.iter_rows():
                for cell in row:
                    cell.number_format = "@"
        workbook.save(source_path)
        workbook.close()
        last_column = get_column_letter(len(EXPECTED_COLUMNS))
        selected_source = read_governed_constraint_row(
            workbook_path=source_path,
            worksheet_name="GovernedRules",
            cell_range=f"A3:{last_column}3",
        )
    source_before = source_path.read_bytes()
    assert selected_source["unique_id"] == selected_id

    # Configuration, not a filename glob or the first workbook sheet, selects
    # the source. Reuse the existing source-collection relationship vocabulary.
    cfg = Namespace("https://dynamicontology.com/uad36/configuration#")
    configuration = Graph()
    for triple in (
        (ns.configuration, RDF.type, cfg.ApplicationConfiguration),
        (ns.configuration, cfg.hasSourceWorkbooks, ns.source_collection),
        (ns.configuration, cfg.activeComplianceSource, ns.active_source),
        (ns.source_collection, RDF.type, cfg.GovernedSourceWorkbookCollection),
        (ns.source_collection, cfg.containsWorkbook, ns.active_source),
        (ns.active_source, RDF.type, cfg.GovernedSourceArtifact),
        (ns.active_source, cfg.sourceFileLocation, Literal(str(source_path))),
        (ns.active_source, cfg.sourceFormat, Literal(source_format)),
    ):
        configuration.add(triple)
    if source_format == "xlsx":
        configuration.add((ns.active_source, cfg.ruleDataSheetName, Literal("GovernedRules")))
    configuration_path = tmp_path / "configuration.ttl"
    configuration.serialize(destination=configuration_path, format="turtle")
    configuration_before = configuration_path.read_bytes()
    schema, projection = _models(ns)
    behaviors = Graph()
    behaviors.add((ns.behavior, RDF.type, CV.ValidationBehavior))
    behaviors.add((ns.behavior, CV.requirementText, Literal(requirement)))
    behaviors.add((ns.behavior, CV.violationCondition, Literal(violation)))
    patterns = Graph()
    patterns.add((ns.behavior, CV.shaclImplementationPattern, ns.pattern))
    patterns.add((ns.pattern, RDF.type, CV.ShaclImplementationPattern))
    patterns.add((ns.pattern, CV.propertyShapeTemplate, ns.template))
    patterns.add((ns.template, RDF.type, SH.PropertyShape))
    patterns.add((ns.template, SH.minCount, Literal(1)))
    inputs_before = [set(graph) for graph in (schema, projection, behaviors, patterns)]

    assert find_spec(PIPELINE_MODULE) is not None, (
        "IT-58R1S1 requires app.services.constraint_production_pipeline to "
        "select one configured CSV or Excel row, orchestrate real workers through verified "
        "RDF handoffs, and persist reusable artifacts without appraisal evaluation."
    )
    pipeline = import_module(PIPELINE_MODULE)
    process = getattr(pipeline, "process_governed_source_row", None)
    assert callable(process), "IT-58R1S1 requires process_governed_source_row()."
    calls = []

    def observe_worker(real, stage, input_name):
        def observed(*args, **kwargs):
            expected_index = len(calls)
            assert expected_index < len(WORKERS), "An unselected row or stage was processed."
            assert stage == WORKERS[expected_index][2], "Workers executed out of governed order."
            if input_name is None:
                assert kwargs["governed_source"]["unique_id"] == selected_id
            else:
                graph = kwargs[input_name]
                previous_stage = WORKERS[expected_index - 1][2]
                previous_results = set(graph.subjects(UADEX.constructionStage, Literal(previous_stage)))
                assert previous_results, "Downstream input lacks its preceding stage result."
                for result_id in previous_results:
                    record = read_construction_stage_result(exchange_graph=graph, result_id=str(result_id))
                    assert record.construction_status == "GREEN"
                    assert record.available_for_downstream is True
                assert set(graph.objects(None, CV.governedUniqueId)) == {Literal(selected_id)}
            calls.append(stage)
            return real(*args, **kwargs)
        return observed

    for module_name, operation_name, stage, input_name in WORKERS:
        worker_module = import_module(f"app.services.{module_name}")
        real = getattr(worker_module, operation_name)
        observed = observe_worker(real, stage, input_name)
        monkeypatch.setattr(worker_module, operation_name, observed)
        # Support an orchestrator that imports the public worker directly.
        if getattr(pipeline, operation_name, None) is real:
            monkeypatch.setattr(pipeline, operation_name, observed)

    def no_evaluation(*args, **kwargs):
        pytest.fail("Constraint production must not evaluate an appraisal or produce runtime findings.")
    monkeypatch.setattr(import_module("pyshacl"), "validate", no_evaluation)
    monkeypatch.setattr(
        import_module("app.services.validation_finding_projection"),
        "project_governed_findings", no_evaluation,
    )
    output_directory = tmp_path / "production-artifacts"
    batch_path = tmp_path / "batch-state" / "results.ttl"
    result = process(
        configuration_path=configuration_path,
        unique_id=selected_id,
        logical_schema_graph=schema,
        projection_model_graph=projection,
        established_behaviors_graph=behaviors,
        implementation_patterns_graph=patterns,
        output_directory=output_directory,
        batch_results_path=batch_path,
    )

    assert calls == [worker[2] for worker in WORKERS]
    assert result.is_complete is True
    assert source_path.read_bytes() == source_before
    assert configuration_path.read_bytes() == configuration_before
    assert [set(graph) for graph in (schema, projection, behaviors, patterns)] == inputs_before
    rdf_path, shacl_path = Path(result.rdf_path), Path(result.shacl_path)
    for artifact in (rdf_path, shacl_path):
        assert artifact.is_file()
        assert artifact.resolve().is_relative_to(output_directory.resolve())
    assert rdf_path != shacl_path
    rdf = Graph().parse(rdf_path, format="turtle")
    shapes = Graph().parse(shacl_path, format="turtle")
    constraint = URIRef(result.constraint_id)
    assert set(rdf.subjects(RDF.type, CV.NormalizedConstraint)) == {constraint}
    assert set(rdf.objects(constraint, CV.governedUniqueId)) == {Literal(selected_id)}
    assert Literal(other_id) not in set(rdf.objects(None, CV.governedUniqueId))
    source = URIRef(selected_source["source_constraint_uri"])
    assert (constraint, PROV.wasDerivedFrom, source) in rdf
    shape_ids = set(shapes.subjects(RDF.type, SH.NodeShape))
    assert len(shape_ids) == 1
    shape = next(iter(shape_ids))
    assert (shape, CV.implementsConstraint, constraint) in shapes
    assert (shape, PROV.wasDerivedFrom, source) in shapes
    assert not list(rdf.subjects(RDF.type, SH.ValidationResult))
    assert batch_path.is_file()
    batch = Graph().parse(batch_path, format="turtle")
    stage_results = set(batch.subjects(UADEX.forConstraint, constraint))
    stage_records = [read_construction_stage_result(exchange_graph=batch, result_id=str(item)) for item in stage_results]
    assert {record.construction_stage for record in stage_records} == {worker[2] for worker in WORKERS}
    assert len(stage_records) == len(WORKERS)
    assert all(record.construction_status == "GREEN" and record.available_for_downstream for record in stage_records)
