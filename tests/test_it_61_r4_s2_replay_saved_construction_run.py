"""IT-61R4S2: replay a real saved UI run without executing construction again.

After real construction, block coordinators, workers and workbook readers.
Open the saved timeline, choose its first result and activate Resume. Follow
each actual automatic playback URL through the last saved stage result.
Compare every frame with preserved Turtle/RDF and its own generating activity.
All helpers and temporary source/model fixtures are self-contained.
"""
from hashlib import sha256
from html.parser import HTMLParser
from importlib import import_module
from pathlib import Path
from urllib.parse import parse_qs, urlsplit
from urllib.request import url2pathname
from uuid import uuid4

import pytest
from fastapi.testclient import TestClient
from openpyxl import Workbook
from openpyxl.worksheet.worksheet import Worksheet
from rdflib import Graph, Literal, Namespace, RDF, RDFS, URIRef
from rdflib.namespace import OWL, PROV, SH, XSD
from rdflib.compare import isomorphic
from rdflib.util import from_n3

from app.main import app
from app.models.schema_model import (
    ComplexTypeDefinition, ElementDeclaration, ModelGroup, ModelGroupKind,
    QName, SchemaModel,
)
from app.services.construction_exchange import UADEX, read_construction_stage_result
from app.services.governed_constraint_source_reader import EXPECTED_COLUMNS
from app.services.logical_schema_serializer import logical_schema_model_to_graph
from app.services.rdf_projection import UAD_SCHEMA
CV = Namespace("https://dynamicontology.com/uad36/constraint-vocabulary#")
CFG = Namespace("https://dynamicontology.com/uad36/configuration#")
LSM = Namespace("https://dynamicontology.com/uad36/logical-schema#")
WORKERS = (
    ("constraint_normalization_agent", "normalize_governed_constraint_to_file",
     "normalized-constraint-representation", None),
    ("logical_schema_binding_agent", "bind_normalized_constraint_to_logical_schema",
     "logical-schema-binding", "normalized_graph"),
    ("context_resolution_agent", "resolve_bound_constraint_context",
     "context-resolution", "bound_graph"),
    ("constraint_behavior_classification", "classify_governed_constraint_behavior",
     "validation-behavior", "context_graph"),
    ("instance_rdf_binding_agent", "bind_governed_constraint_to_instance_rdf",
     "instance-rdf-binding", "classified_graph"),
    ("shacl_generation_agent", "generate_governed_constraint_shacl",
     "shacl-representation", "bound_graph"),
)


def _models(ns: Namespace) -> tuple[Graph, Graph]:
    xml_namespace = "urn:uad36:test:all-rows-pipeline:xml"

    def qname(name: str) -> QName:
        return QName(xml_namespace, name)

    value = ElementDeclaration(name=qname("AddressLineText"))
    address = ElementDeclaration(name=qname("ADDRESS"), type_name=qname("ADDRESS_TYPE"))
    root = ElementDeclaration(name=qname("PROPERTY"), type_name=qname("PROPERTY_TYPE"))
    address_type = ComplexTypeDefinition(
        name=qname("ADDRESS_TYPE"),
        content=ModelGroup(kind=ModelGroupKind.SEQUENCE, elements=(value,)),
    )
    property_type = ComplexTypeDefinition(
        name=qname("PROPERTY_TYPE"),
        content=ModelGroup(kind=ModelGroupKind.SEQUENCE, elements=(address,)),
    )
    schema = logical_schema_model_to_graph(SchemaModel(
        target_namespace=xml_namespace,
        elements={root.name: root},
        complex_types={item.name: item for item in (address_type, property_type)},
    ))
    projection = Graph()
    for triple in schema:
        projection.add(triple)
    for local_name, predicate, range_iri, property_type_iri in (
        ("PROPERTY", ns.property_predicate, ns.property_class, OWL.ObjectProperty),
        ("ADDRESS", ns.address_predicate, ns.address_class, OWL.ObjectProperty),
        ("AddressLineText", ns.value_predicate, XSD.string, OWL.DatatypeProperty),
    ):
        names = {
            subject for subject in schema.subjects(LSM.local_name, Literal(local_name))
            if (subject, RDF.type, LSM.QName) in schema
        }
        assert len(names) == 1
        resource = next(iter(names))
        assert isinstance(resource, URIRef)
        projection.add((resource, UAD_SCHEMA.sourceQName,
                        Literal(f"{{{xml_namespace}}}{local_name}")))
        projection.add((resource, UAD_SCHEMA.projectsTo, predicate))
        projection.add((predicate, RDF.type, property_type_iri))
        projection.add((predicate, RDFS.range, range_iri))
    projection.add((ns.property_class, RDF.type, OWL.Class))
    projection.add((ns.address_class, RDF.type, OWL.Class))
    return schema, projection


class _BuildPage(HTMLParser):
    def __init__(self) -> None:
        super().__init__(convert_charrefs=True)
        self.forms: list[dict] = []
        self.run_ids: list[str] = []
        self._form: dict | None = None
        self._button_parts: list[str] | None = None

    def handle_starttag(self, tag: str, attrs) -> None:
        attributes = dict(attrs)
        if tag == "form":
            self._form = {"attributes": attributes, "inputs": [], "buttons": []}
            self.forms.append(self._form)
        if tag == "input" and self._form is not None:
            self._form["inputs"].append(attributes)
        if tag == "button" and self._form is not None and attributes.get("type", "submit") == "submit":
            self._button_parts = []
        if attributes.get("data-build-run-id"):
            self.run_ids.append(attributes["data-build-run-id"])

    def handle_data(self, data: str) -> None:
        if self._button_parts is not None:
            self._button_parts.append(data)

    def handle_endtag(self, tag: str) -> None:
        if tag == "button" and self._button_parts is not None:
            assert self._form is not None
            self._form["buttons"].append(" ".join("".join(self._button_parts).split()))
            self._button_parts = None
        if tag == "form":
            self._form = None


def _file_artifact(graph: Graph, subject, predicate, output_root: Path) -> Path:
    values = set(graph.objects(subject, predicate))
    assert len(values) == 1
    uri = next(iter(values))
    assert isinstance(uri, URIRef)
    parsed = urlsplit(str(uri))
    assert parsed.scheme == "file" and parsed.netloc in ("", "localhost")
    path = Path(url2pathname(parsed.path)).resolve()
    assert path.is_file()
    assert path.is_relative_to(output_root.resolve())
    return path


class _SavedRunPage(_BuildPage):
    def __init__(self):
        super().__init__()
        self.links = []
        self.stage_links = []
        self.selected = []
        self.rdf_blocks = []
        self.edges = []
        self._rdf_parts = None
        self._svg_depth = 0

    def handle_starttag(self, tag, attrs):
        super().handle_starttag(tag, attrs)
        attributes = dict(attrs)
        if tag == "a" and attributes.get("href"):
            self.links.append(attributes["href"])
            if attributes.get("data-visualization-action") == "select":
                self.stage_links.append(attributes["href"])
        if "data-selected-stage-result-id" in attributes:
            self.selected.append(attributes["data-selected-stage-result-id"])
        if tag == "pre" and attributes.get("id") == "stage-rdf":
            self._rdf_parts = []
        if tag == "svg":
            self._svg_depth += 1
        keys = ("data-rdf-subject-n3", "data-rdf-predicate-n3", "data-rdf-object-n3")
        if self._svg_depth and all(key in attributes for key in keys):
            self.edges.append(tuple(attributes[key] for key in keys))

    def handle_data(self, data):
        super().handle_data(data)
        if self._rdf_parts is not None:
            self._rdf_parts.append(data)

    def handle_endtag(self, tag):
        super().handle_endtag(tag)
        if tag == "pre" and self._rdf_parts is not None:
            self.rdf_blocks.append(''.join(self._rdf_parts))
            self._rdf_parts = None
        if tag == "svg":
            self._svg_depth -= 1


def _files(directory):
    return {path.relative_to(directory): path.read_bytes()
            for path in directory.rglob("*") if path.is_file()}


class _ReplayPage(_SavedRunPage):
    def __init__(self):
        super().__init__()
        self.controls = {}
        self.refresh_targets = []
        self.refresh_count = 0
        self.provenance = []
        self.playback_states = []

    def handle_starttag(self, tag, attrs):
        super().handle_starttag(tag, attrs)
        attributes = dict(attrs)
        if tag == "a" and attributes.get("data-visualization-action") in ("resume", "pause"):
            self.controls.setdefault(attributes["data-visualization-action"], []).append(attributes.get("href", ""))
        if tag == "details" and attributes.get("id") == "stage-provenance":
            self.provenance.append(attributes)
        if "data-playback-state" in attributes:
            self.playback_states.append(attributes["data-playback-state"])
        if tag == "meta" and attributes.get("http-equiv", "").lower() == "refresh":
            self.refresh_count += 1
            parts = attributes.get("content", "").split(";", 1)
            if len(parts) == 2:
                key, separator, value = parts[1].strip().partition("=")
                if separator and key.lower() == "url":
                    self.refresh_targets.append(value.strip().strip("\"'"))


@pytest.mark.parametrize("row_scope", ["one", "group", "all"], ids=["one-row", "selected-group", "all-rows"])
def test_it_61_r4_s2_replays_saved_results_with_original_stage_activity_and_no_construction(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch, row_scope: str,
) -> None:
    selected_index = 0
    ns = Namespace(f"urn:uad36:test:it61r4s2:{uuid4().hex}:")
    requirement = "Provide the address line for each property."
    violation = "The property address line is absent."
    source_directory = tmp_path / "sources"
    source_directory.mkdir()
    registrations = []
    for number, version in enumerate(("1.5", "1.6")):
        digits = f"{uuid4().int % (10 ** 8):08d}"
        # These two textual IDs become the same number if coerced numerically.
        ids = ("0" + digits, digits.lstrip("0") or "0")
        workbook_path = source_directory / f"registered-source-{number}.xlsx"
        workbook = Workbook()
        wrong_sheet = workbook.active
        assert isinstance(wrong_sheet, Worksheet)
        wrong_sheet.title = "DoNotUse"
        wrong_sheet.append(["The configured worksheet is GovernedRules."])
        sheet = workbook.create_sheet("GovernedRules")
        sheet.append(list(EXPECTED_COLUMNS))
        for unique_id in ids:
            row = dict.fromkeys(EXPECTED_COLUMNS, "")
            row.update({
                "Unique ID": unique_id,
                "Primary Data Element": "AddressLineText",
                "Message ID": "test-rule-" + uuid4().hex,
                "Message Text": requirement,
                "Rule Logic": violation,
                "Severity": "Fatal",
                "Property Affected": "/PROPERTY/ADDRESS/AddressLineText",
                "Data Point Name / Value": "AddressLineText",
                " xPath": "/PROPERTY/ADDRESS/AddressLineText",
            })
            sheet.append([row[column] for column in EXPECTED_COLUMNS])
        for cells in sheet.iter_rows():
            for cell in cells:
                cell.number_format = "@"
        workbook.save(workbook_path)
        workbook.close()
        registrations.append({
            "id": URIRef(f"{ns}set-{number}"),
            "workbook": URIRef(f"{ns}workbook-{number}"),
            "path": workbook_path,
            "version": version,
            "row_ids": ids,
        })
    selected = registrations[selected_index]
    other = registrations[1 - selected_index]
    schema, projection = _models(ns)
    behaviors = Graph()
    for triple in (
        (ns.behavior, RDF.type, CV.ValidationBehavior),
        (ns.behavior, CV.requirementText, Literal(requirement)),
        (ns.behavior, CV.violationCondition, Literal(violation)),
    ):
        behaviors.add(triple)
    patterns = Graph()
    for triple in (
        (ns.behavior, CV.shaclImplementationPattern, ns.pattern),
        (ns.pattern, RDF.type, CV.ShaclImplementationPattern),
        (ns.pattern, CV.propertyShapeTemplate, ns.template),
        (ns.template, RDF.type, SH.PropertyShape),
        (ns.template, SH.minCount, Literal(1)),
    ):
        patterns.add(triple)
    model_directory = tmp_path / "models"
    model_directory.mkdir()
    model_files = {}
    for name, graph in (
        ("logical-schema", schema), ("projection-model", projection),
        ("established-behaviors", behaviors), ("shacl-patterns", patterns),
    ):
        path = model_directory / f"{name}.ttl"
        graph.serialize(destination=path, format="turtle")
        model_files[name] = path
    output_root = tmp_path / "build-runs"
    configuration = Graph()
    for triple in (
        (ns.configuration, RDF.type, CFG.ApplicationConfiguration),
        (ns.configuration, CFG.hasSourceWorkbooks, ns.workbooks),
        (ns.configuration, CFG.hasConstraintSets, ns.constraint_sets),
        (ns.configuration, CFG.activeComplianceSource, other["workbook"]),
        (ns.workbooks, RDF.type, CFG.GovernedSourceWorkbookCollection),
        (ns.constraint_sets, RDF.type, CFG.GovernedConstraintSetCollection),
        (ns.authority, RDF.type, CFG.GoverningAuthority),
        (ns.authority, RDFS.label, Literal("Registered governing authority")),
        # Proposed build-profile contract: dependencies and destinations are
        # explicit governed configuration, never guessed from directories.
        (ns.configuration, CFG.usesConstraintBuildProfile, ns.build_profile),
        (ns.build_profile, RDF.type, CFG.ConstraintBuildProfile),
        (ns.build_profile, CFG.logicalSchemaFileLocation, Literal(str(model_files["logical-schema"]))),
        (ns.build_profile, CFG.projectionModelFileLocation, Literal(str(model_files["projection-model"]))),
        (ns.build_profile, CFG.establishedBehaviorsFileLocation, Literal(str(model_files["established-behaviors"]))),
        (ns.build_profile, CFG.shaclPatternsFileLocation, Literal(str(model_files["shacl-patterns"]))),
        (ns.build_profile, CFG.buildOutputDirectory, Literal(str(output_root))),
    ):
        configuration.add(triple)
    for registration in registrations:
        for triple in (
            (ns.constraint_sets, CFG.containsConstraintSet, registration["id"]),
            (registration["id"], RDF.type, CFG.GovernedConstraintSet),
            (registration["id"], RDFS.label, Literal(f"Governed release {registration['version']}")),
            (registration["id"], CFG.governingAuthority, ns.authority),
            (registration["id"], CFG.sourceWorkbook, registration["workbook"]),
            (registration["id"], CFG.governedVersion, Literal(registration["version"])),
            (ns.workbooks, CFG.containsWorkbook, registration["workbook"]),
            (registration["workbook"], RDF.type, CFG.GovernedSourceArtifact),
            (registration["workbook"], CFG.sourceFileLocation, Literal(str(registration["path"]))),
            (registration["workbook"], CFG.ruleDataSheetName, Literal("GovernedRules")),
            (registration["workbook"], CFG.sourceFormat, Literal("xlsx")),
        ):
            configuration.add(triple)
    # Workbook-scope construction must not require an IT-59 constraint-set registry.
    configuration.remove((ns.configuration, CFG.hasConstraintSets, None))
    configuration.remove((ns.constraint_sets, None, None))
    for registration in registrations:
        configuration.remove((registration["id"], None, None))
        configuration.add((registration["workbook"], RDFS.label,
                           Literal(f"Governed source workbook {registration['version']}")))
    configuration_path = tmp_path / "configuration.ttl"
    configuration.serialize(destination=configuration_path, format="turtle")
    immutable_files = [configuration_path, *model_files.values(), *[item["path"] for item in registrations]]
    originals_before_build = {path: path.read_bytes() for path in immutable_files}
    monkeypatch.setattr(app.state, "configuration_file", configuration_path)
    requested_ids = tuple(selected["row_ids"][:1] if row_scope == "one" else selected["row_ids"])
    data = {"source_workbook_id": str(selected["workbook"]), "row_scope": row_scope}
    if row_scope == "one":
        data["unique_id"] = requested_ids[0]
    elif row_scope == "group":
        data["unique_ids"] = list(requested_ids)
    with TestClient(app, follow_redirects=False) as client:
        response = client.post("/constraint-pipeline/runs", data=data)
    assert response.status_code == 200
    completion = _SavedRunPage()
    completion.feed(response.text)
    assert len(completion.run_ids) == 1
    run = URIRef(completion.run_ids[0])
    manifests = list(output_root.rglob("run.ttl"))
    assert len(manifests) == 1
    directory = manifests[0].parent.resolve()
    manifest = Graph().parse(manifests[0], format="turtle")
    assert (run, CV.status, Literal("COMPLETE")) in manifest
    expected = {}
    for result in manifest.objects(run, CV.preservedStageResult):
        ordinals = set(manifest.objects(result, CV.stageOrder))
        locations = set(manifest.objects(result, UADEX.rdfSnapshotPath))
        assert len(ordinals) == len(locations) == 1
        path = Path(str(next(iter(locations)))).resolve()
        assert path.is_relative_to(directory) and path.is_file()
        raw = path.read_bytes()
        graph = Graph().parse(data=raw, format="turtle")
        activities = set(graph.objects(result, PROV.wasGeneratedBy))
        assert len(activities) == 1
        activity = next(iter(activities))
        assert (activity, RDF.type, PROV.Activity) in graph
        assert (result, PROV.wasGeneratedBy, activity) in manifest
        expected[str(result)] = (int(next(iter(ordinals))), raw.decode("utf-8"), graph, str(activity))
    assert len(expected) == len(requested_ids) * len(WORKERS)
    ordered_ids = sorted(expected, key=lambda value: (expected[value][0], value))
    originals = _files(directory)
    attempted = []
    routes = import_module("app.api.routes")
    pipeline = import_module("app.services.constraint_production_pipeline")
    workbench = import_module("app.services.constraint_pipeline_workbench")

    def forbid(name):
        def rejected(*args, **kwargs):
            attempted.append(name)
            pytest.fail(f"IT-61R4S2 saved-run replay must not execute {name}.")
        return rejected

    operations = (
        ("constraint_production_pipeline", "process_governed_source_row"),
        ("constraint_production_pipeline", "process_all_governed_source_rows"),
        ("constraint_production_pipeline", "resume_governed_source_row"),
        ("constraint_production_pipeline", "_construct_governed_source"),
        ("constraint_pipeline_workbench", "process_one_workbook_row"),
        ("constraint_pipeline_workbench", "process_selected_workbook_rows"),
        ("constraint_pipeline_workbench", "process_all_workbook_rows"),
        ("constraint_set_build_workbench", "build_registered_constraint_set"),
        *[(module, operation) for module, operation, _, _ in WORKERS],
    )
    for module_name, operation_name in operations:
        module = import_module(f"app.services.{module_name}")
        real = getattr(module, operation_name)
        blocked = forbid(f"{module_name}.{operation_name}")
        monkeypatch.setattr(module, operation_name, blocked)
        for consumer in (routes, pipeline, workbench):
            if getattr(consumer, operation_name, None) is real:
                monkeypatch.setattr(consumer, operation_name, blocked)
    for module in (pipeline, workbench):
        monkeypatch.setattr(module, "load_workbook", forbid("load_workbook"))
    review_url = f"/constraint-workbench/runs/{directory.name}"

    def check_link(href, result_id, playback):
        parsed = urlsplit(href)
        assert not parsed.scheme and not parsed.netloc and parsed.path == review_url
        query = parse_qs(parsed.query)
        assert query.get("result_id") == [result_id]
        assert query.get("playback") == [playback]
        return href

    def assert_frame(response, result_id, playback):
        assert response.status_code == 200
        page = _ReplayPage()
        page.feed(response.text)
        assert page.run_ids == [str(run)]
        assert page.selected == [result_id]
        assert page.playback_states == [playback]
        _, turtle, graph, activity_id = expected[result_id]
        assert page.rdf_blocks == [turtle]
        displayed = Graph()
        assert page.edges
        for edge in page.edges:
            displayed.add(tuple(from_n3(term) for term in edge))
        assert isomorphic(displayed, graph)
        assert len(page.provenance) == 1
        assert page.provenance[0].get("data-stage-result-id") == result_id
        assert page.provenance[0].get("data-stage-activity-id") == activity_id
        assert _files(directory) == originals
        return page

    seen = []
    with TestClient(app, follow_redirects=False) as later_client:
        response = later_client.get(review_url, params={"playback": "paused"})
        assert response.status_code == 200
        page = _ReplayPage()
        page.feed(response.text)
        choices = [href for href in page.stage_links
                   if parse_qs(urlsplit(href).query).get("result_id") == [ordered_ids[0]]]
        assert len(choices) == 1
        paused = assert_frame(later_client.get(check_link(choices[0], ordered_ids[0], "paused")),
                              ordered_ids[0], "paused")
        resume = paused.controls.get("resume", [])
        assert len(resume) == 1
        href = check_link(resume[0], ordered_ids[0], "playing")
        for index, result_id in enumerate(ordered_ids):
            frame = assert_frame(later_client.get(href), result_id, "playing")
            seen.append(result_id)
            if index + 1 < len(ordered_ids):
                assert frame.refresh_count == 1 and len(frame.refresh_targets) == 1
                href = check_link(frame.refresh_targets[0], ordered_ids[index + 1], "playing")
            else:
                assert frame.refresh_count == 0, "Saved playback must stop at the final preserved result."
    assert seen == ordered_ids
    assert not attempted
    assert _files(directory) == originals
    assert {path: path.read_bytes() for path in immutable_files} == originals_before_build
