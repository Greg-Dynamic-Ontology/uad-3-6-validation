"""IT-59R5S2: a later real worker exception retains earlier verified RDF.

Build a configured two-row workbook through the actual workbench. Missing
behavior, projection or implementation-pattern knowledge triggers a contracted
exception. Every preceding snapshot must remain indexed with original identity,
order and digest. The run must retain each failure's stage and source row, with
no successful failed-stage result or downstream result/artifact invented.

Fixtures/helpers are local to this file. No production operation is mocked.
"""
from hashlib import sha256
from html.parser import HTMLParser
from pathlib import Path
from urllib.parse import urlsplit
from urllib.request import url2pathname
from uuid import uuid4

import pytest
from fastapi.testclient import TestClient
from openpyxl import Workbook
from openpyxl.worksheet.worksheet import Worksheet
from rdflib import Graph, Literal, Namespace, RDF, RDFS, URIRef
from rdflib.namespace import OWL, PROV, SH, XSD

from app.main import app
from app.models.schema_model import (
    ComplexTypeDefinition, ElementDeclaration, ModelGroup, ModelGroupKind,
    QName, SchemaModel,
)
from app.services.construction_exchange import UADEX, read_construction_stage_result
from app.services.governed_constraint_source_reader import EXPECTED_COLUMNS
from app.services.logical_schema_serializer import logical_schema_model_to_graph
from app.services.rdf_projection import UAD_SCHEMA

CV = Namespace("https://dynamicontology.com/uad36/constraint-vocabulary#")
CFG = Namespace("https://dynamicontology.com/uad36/configuration#")
LSM = Namespace("https://dynamicontology.com/uad36/logical-schema#")
WORKERS = (
    ("constraint_normalization_agent", "normalize_governed_constraint_to_file",
     "normalized-constraint-representation", None),
    ("logical_schema_binding_agent", "bind_normalized_constraint_to_logical_schema",
     "logical-schema-binding", "normalized_graph"),
    ("context_resolution_agent", "resolve_bound_constraint_context",
     "context-resolution", "bound_graph"),
    ("constraint_behavior_classification", "classify_governed_constraint_behavior",
     "validation-behavior", "context_graph"),
    ("instance_rdf_binding_agent", "bind_governed_constraint_to_instance_rdf",
     "instance-rdf-binding", "classified_graph"),
    ("shacl_generation_agent", "generate_governed_constraint_shacl",
     "shacl-representation", "bound_graph"),
)


def _models(ns: Namespace) -> tuple[Graph, Graph]:
    xml_namespace = "urn:uad36:test:all-rows-pipeline:xml"

    def qname(name: str) -> QName:
        return QName(xml_namespace, name)

    value = ElementDeclaration(name=qname("AddressLineText"))
    address = ElementDeclaration(name=qname("ADDRESS"), type_name=qname("ADDRESS_TYPE"))
    root = ElementDeclaration(name=qname("PROPERTY"), type_name=qname("PROPERTY_TYPE"))
    address_type = ComplexTypeDefinition(
        name=qname("ADDRESS_TYPE"),
        content=ModelGroup(kind=ModelGroupKind.SEQUENCE, elements=(value,)),
    )
    property_type = ComplexTypeDefinition(
        name=qname("PROPERTY_TYPE"),
        content=ModelGroup(kind=ModelGroupKind.SEQUENCE, elements=(address,)),
    )
    schema = logical_schema_model_to_graph(SchemaModel(
        target_namespace=xml_namespace,
        elements={root.name: root},
        complex_types={item.name: item for item in (address_type, property_type)},
    ))
    projection = Graph()
    for triple in schema:
        projection.add(triple)
    for local_name, predicate, range_iri, property_type_iri in (
        ("PROPERTY", ns.property_predicate, ns.property_class, OWL.ObjectProperty),
        ("ADDRESS", ns.address_predicate, ns.address_class, OWL.ObjectProperty),
        ("AddressLineText", ns.value_predicate, XSD.string, OWL.DatatypeProperty),
    ):
        names = {
            subject for subject in schema.subjects(LSM.local_name, Literal(local_name))
            if (subject, RDF.type, LSM.QName) in schema
        }
        assert len(names) == 1
        resource = next(iter(names))
        assert isinstance(resource, URIRef)
        projection.add((resource, UAD_SCHEMA.sourceQName,
                        Literal(f"{{{xml_namespace}}}{local_name}")))
        projection.add((resource, UAD_SCHEMA.projectsTo, predicate))
        projection.add((predicate, RDF.type, property_type_iri))
        projection.add((predicate, RDFS.range, range_iri))
    projection.add((ns.property_class, RDF.type, OWL.Class))
    projection.add((ns.address_class, RDF.type, OWL.Class))
    return schema, projection


class _BuildPage(HTMLParser):
    def __init__(self) -> None:
        super().__init__(convert_charrefs=True)
        self.forms: list[dict] = []
        self.run_ids: list[str] = []
        self._form: dict | None = None
        self._button_parts: list[str] | None = None

    def handle_starttag(self, tag: str, attrs) -> None:
        attributes = dict(attrs)
        if tag == "form":
            self._form = {"attributes": attributes, "inputs": [], "buttons": []}
            self.forms.append(self._form)
        if tag == "input" and self._form is not None:
            self._form["inputs"].append(attributes)
        if tag == "button" and self._form is not None and attributes.get("type", "submit") == "submit":
            self._button_parts = []
        if attributes.get("data-build-run-id"):
            self.run_ids.append(attributes["data-build-run-id"])

    def handle_data(self, data: str) -> None:
        if self._button_parts is not None:
            self._button_parts.append(data)

    def handle_endtag(self, tag: str) -> None:
        if tag == "button" and self._button_parts is not None:
            assert self._form is not None
            self._form["buttons"].append(" ".join("".join(self._button_parts).split()))
            self._button_parts = None
        if tag == "form":
            self._form = None


def _file_artifact(graph: Graph, subject, predicate, output_root: Path) -> Path:
    values = set(graph.objects(subject, predicate))
    assert len(values) == 1
    uri = next(iter(values))
    assert isinstance(uri, URIRef)
    parsed = urlsplit(str(uri))
    assert parsed.scheme == "file" and parsed.netloc in ("", "localhost")
    path = Path(url2pathname(parsed.path)).resolve()
    assert path.is_file()
    assert path.is_relative_to(output_root.resolve())
    return path


@pytest.mark.parametrize("failed_stage", [
    "validation-behavior", "instance-rdf-binding", "shacl-representation",
], ids=["unknown-behavior", "unresolved-instance-binding", "missing-shacl-pattern"])
def test_it_59_r5_s2_preserves_verified_stages_and_row_failure_without_downstream_results(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch, failed_stage: str,
) -> None:
    selected_index = 0
    ns = Namespace(f"urn:uad36:test:it59r5s2:{uuid4().hex}:")
    requirement = "Provide the address line for each property."
    violation = "The property address line is absent."
    source_directory = tmp_path / "sources"
    source_directory.mkdir()
    registrations = []
    for number, version in enumerate(("1.5", "1.6")):
        ids = tuple("0" + uuid4().hex for _ in range(2))
        workbook_path = source_directory / f"registered-source-{number}.xlsx"
        workbook = Workbook()
        wrong_sheet = workbook.active
        assert isinstance(wrong_sheet, Worksheet)
        wrong_sheet.title = "DoNotUse"
        wrong_sheet.append(["The configured worksheet is GovernedRules."])
        sheet = workbook.create_sheet("GovernedRules")
        sheet.append(list(EXPECTED_COLUMNS))
        for unique_id in ids:
            row = dict.fromkeys(EXPECTED_COLUMNS, "")
            row.update({
                "Unique ID": unique_id,
                "Primary Data Element": "AddressLineText",
                "Message ID": "test-rule-" + uuid4().hex,
                "Message Text": requirement,
                "Rule Logic": violation,
                "Severity": "Fatal",
                "Property Affected": "/PROPERTY/ADDRESS/AddressLineText",
                "Data Point Name / Value": "AddressLineText",
                " xPath": "/PROPERTY/ADDRESS/AddressLineText",
            })
            sheet.append([row[column] for column in EXPECTED_COLUMNS])
        for cells in sheet.iter_rows():
            for cell in cells:
                cell.number_format = "@"
        workbook.save(workbook_path)
        workbook.close()
        registrations.append({
            "id": URIRef(f"{ns}set-{number}"),
            "workbook": URIRef(f"{ns}workbook-{number}"),
            "path": workbook_path,
            "version": version,
            "row_ids": ids,
        })
    selected = registrations[selected_index]
    other = registrations[1 - selected_index]
    schema, projection = _models(ns)
    behaviors = Graph()
    for triple in (
        (ns.behavior, RDF.type, CV.ValidationBehavior),
        (ns.behavior, CV.requirementText, Literal(requirement)),
        (ns.behavior, CV.violationCondition, Literal(violation)),
    ):
        behaviors.add(triple)
    patterns = Graph()
    for triple in (
        (ns.behavior, CV.shaclImplementationPattern, ns.pattern),
        (ns.pattern, RDF.type, CV.ShaclImplementationPattern),
        (ns.pattern, CV.propertyShapeTemplate, ns.template),
        (ns.template, RDF.type, SH.PropertyShape),
        (ns.template, SH.minCount, Literal(1)),
    ):
        patterns.add(triple)
    # Supply genuine insufficient governed knowledge at the selected stage.
    # All preceding workers run normally; no worker is mocked as successful.
    if failed_stage == "validation-behavior":
        behaviors = Graph()
    elif failed_stage == "instance-rdf-binding":
        projection.remove((None, UAD_SCHEMA.projectsTo, ns.value_predicate))
    elif failed_stage == "shacl-representation":
        patterns = Graph()
    else:
        raise AssertionError("Unexpected fixture stage.")
    model_directory = tmp_path / "models"
    model_directory.mkdir()
    model_files = {}
    for name, graph in (
        ("logical-schema", schema), ("projection-model", projection),
        ("established-behaviors", behaviors), ("shacl-patterns", patterns),
    ):
        path = model_directory / f"{name}.ttl"
        graph.serialize(destination=path, format="turtle")
        model_files[name] = path
    output_root = tmp_path / "build-runs"
    configuration = Graph()
    for triple in (
        (ns.configuration, RDF.type, CFG.ApplicationConfiguration),
        (ns.configuration, CFG.hasSourceWorkbooks, ns.workbooks),
        (ns.configuration, CFG.hasConstraintSets, ns.constraint_sets),
        (ns.configuration, CFG.activeComplianceSource, other["workbook"]),
        (ns.workbooks, RDF.type, CFG.GovernedSourceWorkbookCollection),
        (ns.constraint_sets, RDF.type, CFG.GovernedConstraintSetCollection),
        (ns.authority, RDF.type, CFG.GoverningAuthority),
        (ns.authority, RDFS.label, Literal("Registered governing authority")),
        # Proposed build-profile contract: dependencies and destinations are
        # explicit governed configuration, never guessed from directories.
        (ns.configuration, CFG.usesConstraintBuildProfile, ns.build_profile),
        (ns.build_profile, RDF.type, CFG.ConstraintBuildProfile),
        (ns.build_profile, CFG.logicalSchemaFileLocation, Literal(str(model_files["logical-schema"]))),
        (ns.build_profile, CFG.projectionModelFileLocation, Literal(str(model_files["projection-model"]))),
        (ns.build_profile, CFG.establishedBehaviorsFileLocation, Literal(str(model_files["established-behaviors"]))),
        (ns.build_profile, CFG.shaclPatternsFileLocation, Literal(str(model_files["shacl-patterns"]))),
        (ns.build_profile, CFG.buildOutputDirectory, Literal(str(output_root))),
    ):
        configuration.add(triple)
    for registration in registrations:
        for triple in (
            (ns.constraint_sets, CFG.containsConstraintSet, registration["id"]),
            (registration["id"], RDF.type, CFG.GovernedConstraintSet),
            (registration["id"], RDFS.label, Literal(f"Governed release {registration['version']}")),
            (registration["id"], CFG.governingAuthority, ns.authority),
            (registration["id"], CFG.sourceWorkbook, registration["workbook"]),
            (registration["id"], CFG.governedVersion, Literal(registration["version"])),
            (ns.workbooks, CFG.containsWorkbook, registration["workbook"]),
            (registration["workbook"], RDF.type, CFG.GovernedSourceArtifact),
            (registration["workbook"], CFG.sourceFileLocation, Literal(str(registration["path"]))),
            (registration["workbook"], CFG.ruleDataSheetName, Literal("GovernedRules")),
            (registration["workbook"], CFG.sourceFormat, Literal("xlsx")),
        ):
            configuration.add(triple)
    configuration_path = tmp_path / "configuration.ttl"
    configuration.serialize(destination=configuration_path, format="turtle")
    immutable_files = [configuration_path, *model_files.values(), *[item["path"] for item in registrations]]
    originals = {path: path.read_bytes() for path in immutable_files}
    monkeypatch.setattr(app.state, "configuration_file", configuration_path)
    with TestClient(app, follow_redirects=False) as client:
        response = client.post("/constraint-workbench/build", data={
            "authority": str(ns.authority), "constraint_set_id": str(selected["id"]),
        })
    assert response.status_code == 200, "Row exceptions should return a saved build summary."
    page = _BuildPage()
    page.feed(response.text)
    assert len(page.run_ids) == 1
    run = URIRef(page.run_ids[0])
    manifests = list(output_root.rglob("run.ttl"))
    assert len(manifests) == 1
    directory = manifests[0].parent.resolve()
    manifest = Graph().parse(manifests[0], format="turtle")
    assert (run, CV.status, Literal("COMPLETE_WITH_EXCEPTIONS")) in manifest
    assert (run, CV.completedRowCount, Literal(0)) in manifest
    assert (run, CV.exceptionRowCount, Literal(len(selected["row_ids"]))) in manifest
    assert not list(manifest.objects(run, CV.producedConstraint))
    assert (run, CV.selectedConstraintSet, selected["id"]) in manifest
    assert (run, CV.sourceWorkbook, selected["workbook"]) in manifest
    batch_path = _file_artifact(manifest, run, CV.batchResultsArtifact, directory)
    batch = Graph().parse(batch_path, format="turtle")
    stages = [worker[2] for worker in WORKERS]
    failed_index = stages.index(failed_stage)
    earlier_stages = stages[:failed_index]
    constraints = set(batch.subjects(RDF.type, CV.NormalizedConstraint))
    assert len(constraints) == len(selected["row_ids"])
    assert {str(value) for constraint in constraints
            for value in batch.objects(constraint, CV.governedUniqueId)} == set(selected["row_ids"])
    expected_preserved = set()
    expected_exceptions = set()
    seen_paths = set()
    for constraint in constraints:
        ids = set(batch.objects(constraint, CV.governedUniqueId))
        assert len(ids) == 1
        row_id = next(iter(ids))
        assert str(row_id) in selected["row_ids"]
        sources = set(batch.objects(constraint, PROV.wasDerivedFrom))
        assert len(sources) == 1
        source = next(iter(sources))
        assert selected["path"].name in str(source)
        result_nodes = set(batch.subjects(UADEX.forConstraint, constraint))
        records = [read_construction_stage_result(exchange_graph=batch, result_id=str(node))
                   for node in result_nodes]
        assert len(records) == len(earlier_stages) + 1
        assert {record.construction_stage for record in records} == set(earlier_stages + [failed_stage]), (
            "No downstream stage may have a result after the contracted worker exception."
        )
        for ordinal, stage in enumerate(earlier_stages, start=1):
            matches = [record for record in records if record.construction_stage == stage]
            assert len(matches) == 1
            record = matches[0]
            assert record.construction_status == "GREEN" and record.available_for_downstream
            result = record.result_id
            expected_preserved.add(result)
            assert (run, CV.preservedStageResult, result) in manifest, (
                "IT-59R5S2 requires partial-run bundles to retain every completed upstream result."
            )
            assert set(manifest.objects(result, CV.stageOrder)) == {Literal(ordinal)}
            locations = set(batch.objects(result, UADEX.rdfSnapshotPath))
            assert len(locations) == 1
            assert set(manifest.objects(result, UADEX.rdfSnapshotPath)) == locations
            path = Path(str(next(iter(locations)))).resolve()
            assert path.is_file() and path.is_relative_to(directory)
            assert path not in seen_paths
            seen_paths.add(path)
            data = path.read_bytes()
            digest = Literal(sha256(data).hexdigest())
            assert set(batch.objects(result, UADEX.snapshotSha256)) == {digest}
            assert set(manifest.objects(result, UADEX.snapshotSha256)) == {digest}
            snapshot = Graph().parse(data=data, format="turtle")
            assert (result, RDF.type, UADEX.ConstructionStageResult) in snapshot
            assert (result, UADEX.constructionStage, Literal(stage)) in snapshot
            assert (result, UADEX.constructionStatus, Literal("GREEN")) in snapshot
            assert (constraint, PROV.wasDerivedFrom, source) in snapshot
            assert (constraint, CV.governedUniqueId, row_id) in snapshot
        failures = [record for record in records if record.construction_stage == failed_stage]
        assert len(failures) == 1
        failure = failures[0]
        expected_exceptions.add(failure.result_id)
        assert failure.construction_status == "EXCEPTION"
        assert not failure.available_for_downstream
        assert failure.exception_reason and failure.exception_reason.strip()
        assert (run, CV.productionException, failure.result_id) in manifest
        assert (failure.result_id, CV.governedUniqueId, row_id) in batch
        assert (failure.result_id, PROV.wasDerivedFrom, source) in batch
        assert not list(batch.subjects(CV.implementsConstraint, constraint)), (
            "A failed row must not acquire invented executable SHACL."
        )
        complete_states = {state for state in batch.subjects(CV.governedConstraint, constraint)
                           if (state, CV.status, Literal("COMPLETE")) in batch}
        assert not complete_states
    assert set(manifest.objects(run, CV.preservedStageResult)) == expected_preserved
    assert set(manifest.objects(run, CV.productionException)) == expected_exceptions
    assert not list(directory.rglob("shapes.ttl"))
    assert not list(directory.rglob("constraint.ttl"))
    assert {path: path.read_bytes() for path in immutable_files} == originals
